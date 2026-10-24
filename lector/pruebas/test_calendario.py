"""Vencimientos ficticios: las copias del cash no se vuelven a sumar."""
import datetime
from dashboard.datos import calendario
from dashboard.app import render


def ejemplo():
    def fila(importe=100, fecha='2026-10-25', **kw):
        return dict(fecha=fecha, importe=importe, unidad='A', contraparte='Ejemplo', **kw)
    cuota = fila(debito_automatico='Sí')
    impuesto = fila(debito_automatico='No')
    return dict(deuda_droguerias=[fila(), fila(50, '2026-10-20'), fila(999, estado='REVISAR')],
        deuda_bancaria={'cuotas': [cuota, fila(70, debito_automatico='No')]},
        deuda_impositiva=[impuesto, fila(999, estado='pagado')],
        egresos_cashflow=[cuota, impuesto, fila(tipo='CHEQUE'), fila(30, '2026-10-01', tipo='CHEQUE')],
        cuentas_a_cobrar_droguerias=[fila(), fila(100, '2026-10-14'), fila(300, '2026-10-04'),
                                    fila(999, estado='cobrado'), fila(40, unidad_otro='AA')],
        cartera_cheques=[fila(100, estado='EN CARTERA'), fila(20, '2026-10-01', estado='EN CARTERA'),
                         fila(999, estado='DEPOSITADO')])


def test_grupos_vencidos_atraso_y_copias():
    c = ejemplo()
    c['cuentas_a_cobrar_droguerias'][-1]['unidad'] = 'AA'
    r = calendario(c, 'A', '2026-10-24')
    e, i = r['egresos'], r['ingresos']
    assert (e['meses'][0]['firme'], e['meses'][0]['decision']) == (200, 270)
    assert (i['meses'][0]['firme'], i['meses'][0]['decision']) == (100, 100)
    assert e['vencido'] == 50 and i['vencido'] == 400
    assert i['atraso_dias'] == 17.5  # (100 × 10 + 300 × 20) / 400
    assert e['cheques_pasados'] == {'cantidad': 1, 'importe': 30}
    assert i['cheques_pasados'] == {'cantidad': 1, 'importe': 20}
    d = e['meses'][0]['dias'][24]
    assert len(d['items']) == 5 and d['total'] == 470
    assert [x['firme'] for x in d['items']] == [True, True, False, False, False]
    assert all(not x['items'] for x in e['meses'][0]['dias'][:23])
    assert calendario(c, 'GRUPO', '2026-10-24')['ingresos']['meses'][0]['decision'] == 140
    assert calendario(c, 'AA', '2026-10-24')['ingresos']['meses'][0]['decision'] == 40


def test_mes_actual_mas_tres_y_anio_bisiesto():
    r = calendario({}, 'GRUPO', '2027-12-24')['egresos']['meses']
    assert [m['nombre'] for m in r] == ['diciembre 2027', 'enero 2028', 'febrero 2028', 'marzo 2028']
    assert len(r[2]['dias']) == 29
    assert r[0]['maximo'] == 0 and r[0]['mayor'] is None


def test_render_calendario():
    html = render({'cliente': 'Ejemplo', 'fecha': '2026-10-24', 'unidades': ['A'],
                   'datos': {'A': {'calendario': calendario(ejemplo(), 'A', '2026-10-24')}}})
    assert 'data-capa="calendario"' in html and 'function verCalendario()' in html
    assert html.index('data-capa="cobrar"') < html.index('data-capa="calendario"')
    assert '<script src=' not in html


def test_lector_conserva_cheques_pasados_y_marcas(tmp_path):
    import openpyxl
    from lector.pruebas.test_tablero_cajas_deuda import _libro
    from lector.cash_limpio import leer
    ruta = tmp_path / 'inventado.xlsx'
    _libro(ruta)
    wb = openpyxl.load_workbook(ruta)
    for nombre in ('Deuda Bancaria', 'Deuda Impositiva'):
        ws = wb[nombre]
        if nombre == 'Deuda Bancaria':
            ws.cell(6, 9, 'Observaciones')
            ws.cell(7, 9, 'REVISAR: ejemplo inventado')
        else:
            ws.cell(1, 8, 'Observaciones')
            ws.cell(2, 8, 'REVISAR: ejemplo inventado')
    wb.save(ruta)
    c = leer(ruta, 'navar', '2026-10-07')
    r = calendario(c, 'GRUPO', '2026-10-07')
    assert r['ingresos']['cheques_pasados'] == {'cantidad': 1, 'importe': 15}
    assert r['egresos']['cheques_pasados'] == {'cantidad': 1, 'importe': 40}
    assert r['egresos']['meses'][0]['firme'] == 0
    assert r['egresos']['meses'][0]['decision'] == 30


def test_no_confunde_documentos_distintos_del_mismo_importe():
    c = {'cuentas_a_cobrar_droguerias': [
        {'fecha': '2026-10-25', 'importe': 100, 'unidad': 'A', 'factura': n,
         'contraparte': 'Cliente inventado'} for n in ('F-1', 'F-2')]}
    r = calendario(c, 'A', '2026-10-24')['ingresos']['meses'][0]
    assert r['decision'] == 200 and len(r['dias'][24]['items']) == 2
