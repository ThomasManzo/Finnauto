"""Ejemplos inventados: un archivo roto no deja a los otros bancos sin publicar."""
import datetime as dt
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from lector import extractos as e
from clientes.navar.herramientas import vigilante as v


class ExtractosTolerantes(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)

    def excel(self, nombre, filas):
        ruta = self.base / nombre
        ruta.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        for fila in filas:
            wb.active.append(fila)
        wb.save(ruta)
        wb.close()
        return str(ruta)

    def bbva(self, nombre, nueva=False, falta=False, saldo=110):
        cab = [["Empresa:", "Empresa ejemplo(30000000000)"], ["Cuenta:", "111-000111/1"],
               ["Saldo:", str(saldo)]]
        enc = ["Fecha", "Concepto", "Crédito", "Débito"]
        fila = ["23-09-2026", "TRANSFERENCIA", 10, ""]
        if falta:
            enc.pop(2)
            fila.pop(2)
        enc += ["Saldo Parcial", "Nro de cheque"] if nueva else ["Detalle", "Número Documento"]
        fila += [110 if nueva else "Saldo Disponible: 110,00", "123"]
        return self.excel(nombre, cab + [enc, fila])

    def test_dos_formatos_y_columna_imprescindible_ausente(self):
        self.bbva('bbva/viejo.xlsx')
        self.bbva('bbva/nuevo.xlsx', nueva=True)
        self.bbva('bbva/roto.xlsx', falta=True)
        res = e.procesar(self.base)
        self.assertEqual(res['leidos'], 2)
        self.assertEqual(len(res['errores']), 1)
        self.assertIn('Crédito', res['errores'][0]['motivo'])
        self.assertEqual(res['saldos'][0]['saldo'], 110)
        self.assertEqual(len(res['movimientos']), 1)  # No duplica la misma operación.
        self.assertIn('no pude leer bbva/roto.xlsx:', e.resumen(res, 'ejemplo.xlsx', dt.date.today()))
        self.assertEqual(res['movimientos'][0]['Referencia'], '123')

    def test_saldo_dudoso_no_pisa_el_bueno(self):
        self.bbva('bbva/a_viejo.xlsx')
        self.bbva('bbva/z_nuevo.xlsx', nueva=True, saldo=150)
        res = e.procesar(self.base)
        self.assertEqual(res['saldos'][0]['saldo'], 110)
        lectura = res['bancos'][0]['lecturas'][1]
        self.assertEqual(lectura['saldos_por_dia'], {})
        self.assertEqual(len(lectura['movimientos']), 1)
        self.assertIn('REVISAR saldo', lectura['nota'])

    def test_opcionales_y_columnas_reordenadas(self):
        fecha = dt.datetime(2026, 9, 23)
        gal = self.excel('galicia/ok.xlsx', [
            [' crédito ', 'concepto', 'FECHA', 'debito'], [0, 'PAGO', fecha, 12]])
        macro = self.excel('macro/ok.xlsx', [
            ['Monto', 'Descripción', ' fecha '], [-7, 'PAGO', fecha], ['', '', 'Fecha de descarga: hoy']])
        self.assertEqual(e.leer_planilla_galicia(gal)['movimientos'][0]['importe'], -12)
        self.assertIsNone(e.leer_planilla_galicia(gal)['movimientos'][0]['saldo'])
        self.assertEqual(e.leer_planilla_macro(macro)['movimientos'][0]['importe'], -7)
        for banco, lector in [('galicia', e.leer_planilla_galicia), ('macro', e.leer_planilla_macro)]:
            with self.subTest(banco=banco):
                ruta = self.excel(banco + '/sin_fecha.xlsx', [['Concepto'], ['PAGO']])
                with self.assertRaisesRegex(ValueError, 'Fecha'):
                    lector(ruta)

    def test_pdf_y_excel_corruptos_no_frenan_otro_banco(self):
        self.bbva('bbva/ok.xlsx')
        (self.base / 'macro').mkdir()
        (self.base / 'macro/roto.xls').write_bytes(b'no es Excel')
        (self.base / 'macro/roto.pdf').write_bytes(b'no es PDF')
        res = e.procesar(self.base)
        self.assertEqual(res['leidos'], 1)
        self.assertEqual(len(res['errores']), 2)
        self.assertEqual(len(res['movimientos']), 1)
        self.assertEqual(res['bancos'][1]['n'], 0)

    def test_fallo_a_mitad_del_archivo_no_deja_movimientos_sueltos(self):
        ruta = self.excel('macro/roto.xlsx', [['Fecha', 'Concepto', 'Importe'],
                        [dt.datetime(2026, 9, 23), 'PAGO', -10],
                        [dt.datetime(2026, 9, 23), 'PAGO', 'ilegible']])
        self.bbva('bbva/ok.xlsx')
        res = e.procesar(self.base)
        self.assertEqual(len(res['errores']), 1)
        self.assertTrue(all(m['Origen'] == 'Extracto BBVA' for m in res['movimientos']))

    def publicacion(self, carpeta, fecha, macro=True):
        carpeta.mkdir(exist_ok=True)
        movs = []
        if macro:
            for i in range(30):
                movs.append(dict(zip(e.ENC_MOV, [i + 1, dt.date(2026, 9, 20), 'A', 'Egreso',
                    'Otros', 'PAGO', -10, '', 'MACRO ejemplo', 'Extracto MACRO', 'Real', None, None, ''])))
        movs.append(dict(zip(e.ENC_MOV, [99, dt.date(2026, 9, 23), 'A', 'Ingreso',
                    'Otros', 'COBRO', 5, '', 'BBVA ejemplo', 'Extracto BBVA', 'Real', None, None, ''])))
        return e.escribir_para_pegar({'saldos': [], 'movimientos': movs}, carpeta, fecha)

    def test_publica_lo_sano_y_conserva_historia_del_banco_fallido(self):
        salida, entrada = self.base / 'publicados', self.base / 'entrada'
        self.publicacion(salida, dt.date(2026, 9, 23))
        nuevo = self.publicacion(entrada, dt.date(2026, 9, 24), macro=False)
        resumen = entrada / 'resumen_bancos_2026-09-24.md'
        resumen.write_text('- no pude leer macro/roto.xls: archivo ilegible\n')
        with patch.object(v, 'SALIDA', str(salida)), patch.object(v, 'log'):
            self.assertIsNotNone(v.control_contra_anterior(nuevo))
            self.assertIsNone(v.mover_salidas(str(entrada)))
        wb = openpyxl.load_workbook(salida / Path(nuevo).name)
        self.addCleanup(wb.close)
        self.assertEqual(wb['Movimientos'].max_row, 32)
        filas = list(wb['Movimientos'].values)[1:]
        self.assertEqual(sum(r[8] == 'MACRO ejemplo' for r in filas), 30)
        self.assertTrue(all(r[1].date() == dt.date(2026, 9, 20) for r in filas if r[8] == 'MACRO ejemplo'))
        self.assertIn('30 filas recuperadas', (salida / resumen.name).read_text())

    def test_saldos_recuperados_conservan_fecha_y_no_pisan_nuevos(self):
        salida, entrada = self.base / 'publicados', self.base / 'entrada'
        previo = self.publicacion(salida, dt.date(2026, 9, 24))
        nuevo = self.publicacion(entrada, dt.date(2026, 9, 24), macro=False)
        for ruta, saldo in [(previo, 100), (nuevo, 200)]:
            wb = openpyxl.load_workbook(ruta)
            wb['Saldos Bancarios'].append([dt.date(2026, 9, 20), 'MACRO', 'A', 'ejemplo', saldo, 'Extracto MACRO', ''])
            if ruta == previo:
                wb['Saldos Bancarios'].append([dt.date(2026, 9, 19), 'MACRO', 'A', 'ejemplo', 50, 'Extracto MACRO', ''])
            wb.save(ruta)
            wb.close()
        (entrada / 'resumen_bancos_2026-09-24.md').write_text('- no pude leer macro\\roto.xls: error\n')
        with patch.object(v, 'SALIDA', str(salida)), patch.object(v, 'log'):
            self.assertIsNone(v.mover_salidas(str(entrada)))
        wb = openpyxl.load_workbook(salida / Path(nuevo).name)
        self.addCleanup(wb.close)
        saldos = {r[0].date(): r[4] for r in list(wb['Saldos Bancarios'].values)[1:]}
        self.assertEqual(saldos, {dt.date(2026, 9, 19): 50, dt.date(2026, 9, 20): 200})

    def test_sin_publicacion_anterior_igual_publica_lo_sano(self):
        salida, entrada = self.base / 'publicados', self.base / 'entrada'
        nuevo = self.publicacion(entrada, dt.date(2026, 9, 24), macro=False)
        (entrada / 'resumen_bancos_2026-09-24.md').write_text('- no pude leer macro\\roto.xls: error\n')
        with patch.object(v, 'SALIDA', str(salida)), patch.object(v, 'log'):
            self.assertIsNone(v.mover_salidas(str(entrada)))
        self.assertTrue((salida / Path(nuevo).name).exists())

    def test_todo_ilegible_no_publica_vacio_pero_deja_resumen(self):
        self.bbva('bbva/roto.xlsx', falta=True)
        args = ['extractos.py', '--carpeta', str(self.base), '--cliente', 'ejemplo', '--hoy', '2026-09-24']
        with patch('sys.argv', args), patch('sys.stdout'), patch('sys.stderr'):
            with self.assertRaisesRegex(RuntimeError, 'ningún archivo'):
                e.main()
        self.assertFalse(list(self.base.glob('para_pegar*')))
        self.assertIn('salteados: 1', (self.base / 'resumen_bancos_2026-09-24.md').read_text())

    def test_recuperar_dos_veces_no_duplica(self):
        salida, entrada = self.base / 'publicados', self.base / 'entrada'
        self.publicacion(salida, dt.date(2026, 9, 23))
        nuevo = self.publicacion(entrada, dt.date(2026, 9, 24), macro=False)
        (entrada / 'resumen_bancos_2026-09-24.md').write_text('- no pude leer macro\\roto.xls: error\n')
        with patch.object(v, 'SALIDA', str(salida)), patch.object(v, 'log'):
            v.conservar_bancos_con_errores(nuevo)
            v.conservar_bancos_con_errores(nuevo)
        wb = openpyxl.load_workbook(nuevo)
        self.addCleanup(wb.close)
        self.assertEqual(wb['Movimientos'].max_row, 32)


if __name__ == '__main__':
    unittest.main()
