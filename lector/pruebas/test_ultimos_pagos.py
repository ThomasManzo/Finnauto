"""Ejemplos inventados: último pago, archivos incompletos y separación de tesorerías."""
import datetime as dt
import io
import json
import os
import subprocess
import sys
import tempfile
import time
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

import openpyxl
from lector import ultimos_pagos as lector
from clientes.navar.herramientas import vigilante


ENC = ["Tipo", "Comprobante", "Fecha", "Clase", "Total (cte)", "Cód. relacionado",
       "Desc. relacionado", "Fecha de emisión"]


def fila(tipo="REC", numero="9", fecha=dt.date(2026, 9, 1), clase="Cobros",
         importe=100, codigo="001", nombre="Cliente  Ejemplo", emision=None):
    return [tipo, numero, fecha, clase, importe, codigo, nombre, emision]


def excel(ruta, filas, encabezado=ENC):
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(encabezado)
    for f in filas:
        ws.append(f)
    wb.save(ruta)
    wb.close()
    return ruta


class UltimosPagos(unittest.TestCase):
    def test_codigo_en_descripcion_con_letra_de_tipo(self):
        casos = [
            ("C", "900001 - CLIENTE INVENTADO S.A.", "900001", "CLIENTE INVENTADO S.A."),
            ("P", " U&G001 - U & G  SRL ", "U&G001", "U & G  SRL"),
            ("X", "NUÑ004 - PROVEEDOR INVENTADO", "NUÑ004", "PROVEEDOR INVENTADO"),
            (None, "C&D000 - OTRO INVENTADO", "C&D000", "OTRO INVENTADO"),
            ("NUÑ004", "NUÑ004 - PROVEEDOR INVENTADO", "NUÑ004", "PROVEEDOR INVENTADO"),
            ("REAL01", "900001 - CLIENTE INVENTADO", "REAL01", "900001 - CLIENTE INVENTADO"),
            ("C", "AB - NOMBRE", "C", "AB - NOMBRE"),
            ("C", "ABCDEFGHI - NOMBRE", "C", "ABCDEFGHI - NOMBRE"),
            ("C", "AB CD - NOMBRE", "C", "AB CD - NOMBRE"),
            ("C", "ABC - NOMBRE", "ABC", "NOMBRE"),
            ("P", "12345678 - NOMBRE", "12345678", "NOMBRE"),
            ("C", "ABC - ", "C", "ABC -"),
            ("C", " - ", "C", ""),
            ("P", "", "P", ""),
            ("C", "NOMBRE SIN PREFIJO", "C", "NOMBRE SIN PREFIJO"),
            ("1", "ABC - NOMBRE", "1", "ABC - NOMBRE"),
        ]
        for codigo, descripcion, esperado_codigo, esperado_nombre in casos:
            with self.subTest(codigo=codigo, descripcion=descripcion):
                self.assertEqual(lector._relacionado(codigo, descripcion),
                                 (esperado_codigo, esperado_nombre))

    def test_agrupa_codigos_extraidos_y_descarta_relacionado_vacio(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = excel(Path(tmp) / "A.xlsx", [
                fila(codigo="C", nombre="900001 - CLIENTE  INVENTADO", numero="9"),
                fila(codigo="C", nombre="900002 - CLIENTE  INVENTADO", numero="10",
                     fecha=dt.date(2026, 9, 2)),
                fila(codigo="P", nombre=" - ", tipo="OPF", clase="Pagos"),
                fila(codigo="P", nombre="", tipo="OPF", clase="Pagos"),
            ])
            filas, control = lector.leer(ruta, "A")
            self.assertEqual(len(filas), 1)
            self.assertEqual(filas[0]["Razon Social"], "CLIENTE  INVENTADO")
            self.assertEqual(filas[0]["Codigo"], "900001 / 900002")
            self.assertEqual(filas[0]["Comprobante"], "REC 10")
            self.assertEqual(filas[0]["Fecha Ultimo Pago"], dt.date(2026, 9, 2))
            self.assertEqual(control["ignorados"], {"sin relacionado": 2})
            salida = lector.escribir(filas, tmp, dt.date(2026, 9, 29))
            wb = openpyxl.load_workbook(salida)
            try:
                self.assertEqual(wb.active["C2"].value, "900001 / 900002")
                self.assertEqual(wb.active["D2"].value, "CLIENTE  INVENTADO")
            finally:
                wb.close()

    def test_dos_empresas_reglas_y_salida(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            a = excel(root / "Tesoreria A" / "A.xlsx", [
                fila(nombre=" 001 - Cliente  Ejemplo  "),
                fila(numero="10", codigo="002", nombre=" Cliente  Ejemplo "),
                fila(tipo="REV", numero="99", fecha=dt.date(2026, 9, 29)),
                fila(numero="100", fecha=dt.date(2026, 9, 28), importe=-200),
                fila(tipo="FAC", numero="003", codigo="003", nombre="Contado Ejemplo"),
                fila(tipo="O/P", clase="Pagos", codigo="P01", nombre="Proveedor Ejemplo"),
                fila(tipo="OPF", clase="Pagos", numero="10", codigo="P01", nombre="Proveedor Ejemplo"),
                fila(tipo="FPR", clase="Pagos", numero="0002", codigo="P02", nombre="Contado Proveedor",
                     fecha=None, emision=dt.date(2026, 9, 20)),
                fila(tipo="EXT", clase="Otros"),
                fila(nombre=""),
                fila(codigo="F01", nombre="Sin fecha", fecha="mal"),
                fila(numero="", nombre="Sin comprobante"),
                fila(importe="mal", nombre="Sin importe"),
            ])
            enc_ext = [c.replace("Total (cte)", "Total (ext)") for c in ENC]
            aa = excel(root / "Tesoreria AA" / "AA.xlsx", [
                fila(codigo=None, nombre="009 - Cliente  Ejemplo", numero="0004"),
                fila(codigo="NAV007", nombre="Pase de caja", tipo="O/P", clase="Pagos"),
                fila(codigo="INT", nombre="NAVAR S.A.", tipo="OPF", clase="Pagos"),
                fila(tipo="FPR", clase="Pagos", codigo="P02", nombre="Proveedor Ejemplo"),
            ], enc_ext)
            with redirect_stdout(io.StringIO()):
                self.assertEqual(lector.main(["--a", str(a), "--aa", str(aa), "--hoy", "2026-09-29"]), 0)
            salida = a.parent / "para_pegar_ultimos_pagos_2026-09-29.xlsx"
            wb = openpyxl.load_workbook(salida)
            try:
                self.assertEqual(wb.sheetnames, ["Ultimos Pagos"])
                datos = list(wb.active.values)
                self.assertEqual(list(datos[0]), lector.ENCABEZADOS)
                por_clave = {(f[0], f[1], f[3]): f for f in datos[1:]}
                self.assertEqual(len(por_clave), 6)
                cliente = por_clave["A", "Cliente", "Cliente  Ejemplo"]
                self.assertEqual(cliente[2], "001 / 002")
                self.assertEqual(cliente[4], dt.datetime(2026, 9, 1))
                self.assertEqual(cliente[5], "REC 10")
                self.assertEqual(cliente[6], "Tango tesorería · A.xlsx")
                self.assertEqual(por_clave["A", "Proveedor", "Proveedor Ejemplo"][5], "OPF 10")
                self.assertEqual(por_clave["A", "Proveedor", "Contado Proveedor"][4], dt.datetime(2026, 9, 20))
                self.assertEqual(por_clave["AA", "Cliente", "Cliente  Ejemplo"][2], "009")
                for row in wb.active.iter_rows(min_row=2):
                    self.assertEqual(row[2].number_format, "@")
                    self.assertEqual(row[5].data_type, "s")
            finally:
                wb.close()
            md = (a.parent / "resumen_ultimos_pagos_2026-09-29.md").read_text()
            for motivo in ("anulación REV", "importe negativo", "tipo o clase fuera de alcance",
                           "sin relacionado", "fecha ausente o inválida", "sin comprobante",
                           "importe ausente o inválido", "pase interno de AA"):
                self.assertIn(motivo, md)
            self.assertIn("Clientes: 2 · proveedores: 2", md)
            self.assertIn("2026-09-01 a 2026-09-29", md)

    def test_encabezados_reordenados_y_signo_en_ambas_monedas(self):
        with tempfile.TemporaryDirectory() as tmp:
            enc = list(reversed(ENC)) + ["Total (ext)"]
            ruta = excel(Path(tmp) / "A.xlsx", [list(reversed(fila())) + [-1],
                         list(reversed(fila(nombre="Aceptado", numero="012"))) + [0]], enc)
            filas, control = lector.leer(ruta, "A")
            self.assertEqual([f["Razon Social"] for f in filas], ["Aceptado"])
            self.assertEqual(control["ignorados"]["importe negativo"], 1)

    def test_ultimo_por_fecha_antes_que_por_numero_y_sin_cruzar_tipo(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = excel(Path(tmp) / "A.xlsx", [fila(numero="999"),
                         fila(numero="1", fecha=dt.date(2026, 9, 2)),
                         fila(tipo="OPF", clase="Pagos", numero="555")])
            filas, _ = lector.leer(ruta, "A")
            self.assertEqual([f["Comprobante"] for f in filas], ["REC 1", "OPF 555"])
            self.assertEqual(lector._relacionado("001", "Nombre - Sucursal"), ("001", "Nombre - Sucursal"))
            self.assertEqual(lector._relacionado(None, "CLI - Cliente  Ejemplo"), ("CLI", "Cliente  Ejemplo"))

    def test_falta_o_es_invalido_un_archivo_no_escribe(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            a = excel(root / "A.xlsx", [fila()])
            aa = root / "AA.xlsx"
            for estado in ("ausente", "sin encabezados"):
                with self.subTest(estado=estado):
                    if estado != "ausente":
                        excel(aa, [["cualquier dato"]], ["Otra columna"])
                    with self.assertRaises(ValueError):
                        lector.main(["--a", str(a), "--aa", str(aa)])
                    self.assertEqual([], list(root.glob("para_pegar*")))
                    self.assertEqual([], list(root.glob("resumen*")))
            # La entrada de consola también comunica el error con código distinto de cero.
            r = subprocess.run([sys.executable, lector.__file__, "--a", str(a), "--aa", str(root / "falta.xlsx")],
                               capture_output=True, text=True)
            self.assertNotEqual(r.returncode, 0)
            self.assertIn("falta el archivo", r.stderr)
            self.assertEqual([], list(root.glob("para_pegar*")))

    def test_nombres_y_referencias_no_se_convierten_en_formulas(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = excel(Path(tmp) / "A.xlsx", [fila(nombre="Normal")])
            filas, _ = lector.leer(ruta, "A")
            filas[0]["Razon Social"] = "=Nombre ficticio"
            salida = lector.escribir(filas, tmp, dt.date(2026, 9, 29))
            wb = openpyxl.load_workbook(salida)
            try:
                self.assertEqual(wb.active["D2"].data_type, "s")
                self.assertEqual(wb.active["C2"].value, "001")
            finally:
                wb.close()


class VigilanteUltimosPagos(unittest.TestCase):
    def test_archivo_nuevo_a_dispara_lista_pero_no_caja_aa(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            a = excel(root / "Tesoreria A" / "A anterior.xlsx", [fila()])
            aa = excel(root / "Tesoreria AA" / "AA.xlsx", [fila()])
            salida = root / "_para la Sheet"
            estado = root / "estado.json"
            with mock.patch.object(vigilante, "DRIVE", str(root)), \
                    mock.patch.object(vigilante, "SALIDA", str(salida)), \
                    mock.patch.object(vigilante, "ESTADO", str(estado)), \
                    mock.patch.object(vigilante, "LOG", str(root / "prueba.log")), \
                    mock.patch.object(vigilante, "ESPERA_SEG", 0), \
                    mock.patch.object(sys, "argv", ["vigilante", "--hoy", "2026-09-29"]):
                fuentes = vigilante.fuentes(dt.date(2026, 9, 29))
                firmas = {f["nombre"]: vigilante.firma(f["archivos"]) for f in fuentes}
                fuente = next(f for f in fuentes if f["nombre"] == "ultimos_pagos")
                firmas["ultimos_pagos"] += "|" + "|".join(str(os.stat(r).st_mtime_ns) for r, _ in fuente["archivos"])
                estado.write_text(json.dumps(firmas))
                nuevo = excel(a.parent / "A nuevo.xlsx", [fila(numero="10")])
                os.utime(a, (time.time() - 300, time.time() - 300))
                # Una subcarpeta más nueva no debe seleccionarse.
                excel(a.parent / "subcarpeta" / "ignorar.xlsx", [fila(numero="99")])
                correr = subprocess.run
                with mock.patch.object(vigilante.subprocess, "run", wraps=correr) as comando, redirect_stdout(io.StringIO()):
                    self.assertEqual(vigilante.main(), 0)
                    self.assertEqual(comando.call_count, 1)
                    args = comando.call_args.args[0]
                    self.assertEqual(Path(args[1]).name, "ultimos_pagos.py")
                    self.assertEqual(args[args.index("--a") + 1], str(nuevo))
                    self.assertEqual(args[args.index("--aa") + 1], str(aa))
                    self.assertTrue((salida / "para_pegar_ultimos_pagos_2026-09-29.xlsx").is_file())
                    self.assertTrue((salida / "resumen_ultimos_pagos_2026-09-29.md").is_file())
                    self.assertEqual(vigilante.main(), 0)
                    self.assertEqual(comando.call_count, 1)
                    # Mismo nombre y tamaño, nueva fecha de modificación: se reprocesa.
                    ns = os.stat(nuevo).st_mtime_ns + 1000000
                    os.utime(nuevo, ns=(ns, ns))
                    self.assertEqual(vigilante.main(), 0)
                    self.assertEqual(comando.call_count, 2)
                    nuevo_aa = excel(aa.parent / "AA nuevo.xlsx", [fila(numero="22")])
                    os.utime(aa, (time.time() - 300, time.time() - 300))
                    self.assertEqual(vigilante.main(), 0)
                    self.assertEqual(comando.call_count, 4)
                    ultimas = [c.args[0] for c in comando.call_args_list[-2:]]
                    self.assertEqual([Path(c[1]).name for c in ultimas],
                                     ["tesoreria_aa.py", "ultimos_pagos.py"])
                    self.assertIn(str(nuevo_aa), ultimas[0])
                    self.assertNotIn(str(nuevo), ultimas[0])
                    nuevo_aa.unlink()
                aa.unlink()
                fuente = next(f for f in vigilante.fuentes(dt.date(2026, 9, 29)) if f["nombre"] == "ultimos_pagos")
                self.assertEqual(fuente["archivos"], [])

    def test_aa_con_tilde_y_sin_a_no_publica_lista(self):
        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(vigilante, "DRIVE", tmp):
            aa = excel(Path(tmp) / "Tesorería AA" / "AA.xlsx", [fila()])
            fuentes = {f["nombre"]: f for f in vigilante.fuentes(dt.date(2026, 9, 29))}
            self.assertEqual(fuentes["ultimos_pagos"]["archivos"], [])
            self.assertIn(str(aa), fuentes["tesoreria_aa"]["cmd"]())
            excel(Path(tmp) / "Tesoreria A" / "A.xlsx", [fila()])
            fuentes = {f["nombre"]: f for f in vigilante.fuentes(dt.date(2026, 9, 29))}
            self.assertIn(str(aa), fuentes["ultimos_pagos"]["cmd"]())
            self.assertEqual(len(fuentes["tesoreria_aa"]["archivos"]), 1)


if __name__ == "__main__":
    unittest.main()
