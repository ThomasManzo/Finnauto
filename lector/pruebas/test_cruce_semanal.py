"""Datos inventados: el cruce semanal arma el Excel para la administración y los datos del mail."""
import datetime as dt
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import openpyxl
from lector import cruce_semanal as cs
from lector.pruebas.test_cruce import CFG, ENC_MOV, ENC_TANGO

D = dt.datetime


class CruceSemanal(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.drive = Path(self.tmp.name) / "NAVAR - Datos"
        self.mov, self.tg = [], []

    def banco(self, fecha, cuenta, importe, concepto="MOV", categoria="Otros"):
        self.mov.append([len(self.mov) + 1, fecha, "A", "", categoria, concepto, importe, "", cuenta,
                         "Extracto X", "Real", "REF%d" % len(self.mov), None, ""])

    def tango(self, fecha, tipo, cod, importe, leyenda=""):
        debe, haber = (importe, 0) if importe >= 0 else (0, -importe)
        self.tg.append(["X", tipo, " %010d/0" % (len(self.tg) + 1), len(self.tg) + 1, cod, "BANCO", debe, haber,
                        abs(importe), None, None, None, None, leyenda, fecha, "BANCO"])

    def correr(self, hoy="2026-10-07"):
        for carpeta in ("_para la Sheet", "Tesoreria A detalle"):
            (self.drive / carpeta).mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Movimientos"
        ws.append(ENC_MOV)
        for f in self.mov:
            ws.append(f)
        wb.save(self.drive / "_para la Sheet" / "para_pegar_bancos_2026-10-07.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Detalle de comprobantes"
        wb.active.append(ENC_TANGO)
        for f in self.tg:
            wb.active.append(f)
        wb.save(self.drive / "Tesoreria A detalle" / "A tesoreria detalle 2026-10-07.xlsx")
        destino = Path(self.tmp.name) / "salida"
        with patch.object(cs.cz, "cargar_config", return_value=dict(CFG, cuit_propio="30999999995")), \
                redirect_stdout(io.StringIO()):
            rc = cs.main(["--cliente", "prueba", "--hoy", hoy, "--drive", str(self.drive), "--destino", str(destino)])
        datos = json.loads((destino / ("cruce_semanal_%s.json" % hoy)).read_text(encoding="utf-8"))
        excel = openpyxl.load_workbook(destino / datos["archivo"])
        return rc, datos, excel

    def test_lo_que_falta_cargar_y_el_detalle(self):
        self.banco(D(2026, 9, 30), "GALICIA 111", -24000.0, "CUOTA DE PRESTAMO 6", "Prestamo")    # falta en Tango
        self.banco(D(2026, 10, 5), "GALICIA 111", -800.0, "TRANSF RECIENTE", "Proveedores")      # menos de 7 días
        self.banco(D(2026, 9, 10), "GALICIA 111", 5000.0, "TRANSF 20111111112")                   # concilia
        self.tango(D(2026, 9, 10), "REC", 5, 5000.0)
        self.banco(D(2026, 9, 15), "GALICIA 111", -12.5, "IMP DEBITOS Y CREDITOS", "Impuestos")   # gasto
        self.tango(D(2026, 9, 20), "O/P", 5, -3000.0, "PAGO A PROVEEDOR")                         # solo en Tango
        rc, d, excel = self.correr()
        self.assertEqual(rc, 0)
        self.assertEqual(d["meses"], ["septiembre", "octubre"])
        self.assertEqual(d["falta_total"]["cantidad"], 1)                 # la de octubre todavía no cuenta
        self.assertEqual(d["falta"][0]["concepto"], "CUOTA DE PRESTAMO 6")
        self.assertEqual(d["gastos"], [{"banco": "Galicia 111", "cantidad": 1, "importe": -12.5}])
        self.assertEqual(d["bancos"][0]["extracto_hasta"], "2026-10-05")
        self.assertTrue(d["control_ok"])
        self.assertEqual(excel.sheetnames, ["Por banco", "Falta cargar en Tango", "Posibles errores",
                                            "Detalle de lo que no cruza"])
        lados = [r[0] for r in excel["Detalle de lo que no cruza"].iter_rows(min_row=2, values_only=True)]
        self.assertIn("En Tango, no aparece en el banco", lados)
        self.assertEqual(sum(1 for x in lados if x.startswith("En el banco, falta en Tango")), 3)
        refs = [r[3] for r in excel["Falta cargar en Tango"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(refs, ["REF0"])                                   # la referencia del banco, para buscarlo

    def test_error_de_tipeo_y_fecha_imposible(self):
        self.banco(D(2026, 9, 14), "MACRO 222", -12345678.0, "TARJETA", "Otros")
        self.tango(D(2026, 8, 31), "EXT", 25, -12354678.0)
        self.tango(D(2036, 9, 2), "BDM", 25, 10.0)
        rc, d, _ = self.correr()
        textos = " | ".join(d["errores"])
        self.assertIn("error de tipeo", textos)
        self.assertIn("fecha imposible", textos)
        self.assertEqual(len([e for e in d["errores"] if "fecha imposible" in e]), 1)   # no se repite por mes


if __name__ == "__main__":
    unittest.main()
