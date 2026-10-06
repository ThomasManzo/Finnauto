# -*- coding: utf-8 -*-
"""Pruebas de lector/cajas.py con detalles de tesorería INVENTADOS (nada real: el repo es público)."""
import os
import sys
import datetime
import tempfile
import unittest

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, BASE)
from lector import cajas  # noqa: E402

ENC = ["Fecha", "Fecha de emisión", "Cód. comprobante", "Comprobante", "Nro. interno", "Renglón",
       "Cód. cuenta", "Desc. contable", "Desc. cuenta", "Debe (cte) (renglón)", "Haber (cte) (renglón)",
       "Razón social (encab.)", "Leyenda", "COD_TIPO_CUENTA"]

CFG = {
    "A": {"nombre_en_el_cash": "Caja A", "cuentas": ["CAJA CHICA"]},
    "AA": {"nombre_en_el_cash": "Caja AA", "cuentas": ["CAJA FUERTE", "CAJA CHICA"]},
    "cuentas_de_pase": ["BILLETERA DIGITAL", "EMPRESA HERMANA"],
}


def renglon(interno, tipo, cuenta, debe=0, haber=0, tipo_cuenta="O", razon="", leyenda="", fecha="2026-10-01"):
    return {"Fecha": datetime.datetime.fromisoformat(fecha), "Fecha de emisión": datetime.datetime.fromisoformat(fecha),
            "Cód. comprobante": tipo, "Comprobante": "0000-%05d" % interno, "Nro. interno": interno, "Renglón": 1,
            "Cód. cuenta": 1, "Desc. contable": cuenta, "Desc. cuenta": cuenta, "Debe (cte) (renglón)": debe,
            "Haber (cte) (renglón)": haber, "Razón social (encab.)": razon, "Leyenda": leyenda,
            "COD_TIPO_CUENTA": tipo_cuenta}


def escribir(ruta, filas, encabezado=ENC):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(encabezado)
    for f in filas:
        ws.append([f.get(c) for c in encabezado])
    wb.save(ruta)


class CajasTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()

    def mov(self, filas, empresa="A"):
        ruta = os.path.join(self.dir, "%s.xlsx" % empresa)
        escribir(ruta, filas)
        return cajas.movimientos(cajas.leer_detalle(ruta), empresa, CFG, os.path.basename(ruta))

    def test_pago_por_caja_un_renglon_negativo_con_cuenta_y_leyenda(self):
        m = self.mov([renglon(1, "OPF", "GASTOS VARIOS", debe=120000),
                      renglon(1, "OPF", "CAJA CHICA 01/01", haber=120000, leyenda="VIATICOS VIAJE INVENTADO")])
        self.assertEqual(1, len(m))
        x = m[0]
        self.assertEqual((-120000, "Egreso", "Otros", "Caja A"), (x["Importe"], x["Tipo"], x["Categoria"], x["Banco / Cuenta"]))
        self.assertEqual("CAJA CHICA 01/01", x["Cuenta Tango"])
        self.assertEqual("VIATICOS VIAJE INVENTADO", x["Leyenda"])
        self.assertTrue(x["Origen"].startswith("Tango caja A"))

    def test_recibo_mitad_caja_mitad_banco_solo_la_parte_de_caja_y_es_cobranza(self):
        m = self.mov([renglon(2, "REC", "CAJA CHICA", debe=300, razon="CLIENTE INVENTADO"),
                      renglon(2, "REC", "BANCO INVENTADO", debe=700, tipo_cuenta="B"),
                      renglon(2, "REC", "DEUDORES VARIOS", haber=1000)])
        self.assertEqual([(300, "Ingreso", "Cobranza Facturas")], [(x["Importe"], x["Tipo"], x["Categoria"]) for x in m])

    def test_deposito_al_banco_es_cobranza_en_negativo_para_no_contar_dos_veces(self):
        # El extracto trae el depósito como cobro; en la caja resta: el cobro queda una sola vez.
        m = self.mov([renglon(3, "BDC", "CAJA CHICA", haber=500), renglon(3, "BDC", "BANCO X", debe=500, tipo_cuenta="B")])
        self.assertEqual([(-500, "Ingreso", "Cobranza Facturas")], [(x["Importe"], x["Tipo"], x["Categoria"]) for x in m])
        self.assertIn("Depósito", m[0]["Concepto / Detalle"])
        # y su anulación lo devuelve en la misma fila
        m = self.mov([renglon(4, "REV", "CAJA CHICA", debe=500), renglon(4, "REV", "BANCO X", haber=500, tipo_cuenta="B")])
        self.assertEqual([(500, "Ingreso", "Cobranza Facturas")], [(x["Importe"], x["Tipo"], x["Categoria"]) for x in m])

    def test_del_banco_a_la_caja_y_pases_son_internos_y_la_cartera_es_cheques(self):
        m = self.mov([renglon(3, "EXT", "CAJA CHICA", debe=500), renglon(3, "EXT", "BANCO X", haber=500, tipo_cuenta="B"),
                      renglon(4, "EXT", "CAJA CHICA", debe=80), renglon(4, "EXT", "VALORES A DEPOSITAR", haber=80, tipo_cuenta="C"),
                      renglon(5, "EXT", "CAJA CHICA", haber=40), renglon(5, "EXT", "BILLETERA DIGITAL", debe=40)])
        self.assertEqual([(500, "Transferencia Interna"), (80, "Cheques"), (-40, "Transferencia Interna")],
                         [(x["Importe"], x["Categoria"]) for x in m])
        self.assertEqual("Ingreso", m[1]["Tipo"])

    def test_lo_que_esta_del_mismo_lado_que_la_caja_no_decide(self):
        # Pago a un productor: la caja y el convenio gremial que se le descuenta están en el Haber;
        # el destino (acreedores) en el Debe. Es pago a proveedor, no sueldos.
        m = self.mov([renglon(20, "O/P", "ACREEDORES VARIOS", debe=1000), renglon(20, "O/P", "CAJA CHICA", haber=700),
                      renglon(20, "O/P", "CONVENIO CORRESP. GREMIAL", haber=300)])
        self.assertEqual("Proveedores MP y Logist.", m[0]["Categoria"])
        # Recibo con vuelto: la caja da plata que va al banco junto con el pago del cliente.
        m = self.mov([renglon(21, "REC", "CAJA CHICA", haber=50), renglon(21, "REC", "DEUDORES VARIOS", haber=950),
                      renglon(21, "REC", "BANCO X", debe=1000, tipo_cuenta="B")])
        self.assertEqual([(-50, "Ingreso", "Cobranza Facturas")], [(x["Importe"], x["Tipo"], x["Categoria"]) for x in m])

    def test_pase_entre_dos_cajas_de_aa_genera_dos_renglones_internos(self):
        m = self.mov([renglon(6, "EXT", "CAJA FUERTE", haber=90), renglon(6, "EXT", "CAJA CHICA", debe=90)], "AA")
        self.assertEqual([(-90, "Transferencia Interna"), (90, "Transferencia Interna")],
                         [(x["Importe"], x["Categoria"]) for x in m])
        self.assertTrue(all(x["Origen"].startswith("Tango AA") for x in m))

    def test_categorias_por_la_otra_cuenta(self):
        m = self.mov([
            renglon(7, "OPF", "SUELDOS Y JORNALES", debe=50), renglon(7, "OPF", "CAJA CHICA", haber=50),
            renglon(8, "O/P", "ACREEDORES VARIOS", debe=60), renglon(8, "O/P", "RETENCIONES GANANCIAS", haber=5),
            renglon(8, "O/P", "CAJA CHICA", haber=55),
            renglon(9, "OPF", "GASTOS VARIOS", debe=7), renglon(9, "OPF", "CAJA CHICA", haber=7, leyenda="APORTES SINDICATO MAYO"),
        ])
        self.assertEqual(["Sueldos y Jornales", "Proveedores MP y Logist.", "Sueldos y Jornales"], [x["Categoria"] for x in m])
        m = self.mov([renglon(8, "O/P", "ACREEDORES VARIOS", debe=60), renglon(8, "O/P", "CAJA FUERTE", haber=60)], "AA")
        self.assertEqual("Proveedores AA", m[0]["Categoria"])

    def test_pase_de_aa_a_la_caja_de_a_se_refleja_en_la_caja_de_a(self):
        # En el Tango de AA la caja de A es una cuenta más; A no anota la entrada en el suyo.
        cfg = dict(CFG, AA=dict(CFG["AA"], espejo={"empresa": "A", "cuentas": ["CAJA DE LA OTRA"]}),
                   cuentas_de_pase=CFG["cuentas_de_pase"] + ["CAJA DE LA OTRA"])
        ruta = os.path.join(self.dir, "aa.xlsx")
        escribir(ruta, [renglon(30, "EXT", "CAJA FUERTE", haber=20, leyenda="DEPOSITO BANCO INVENTADO"),
                        renglon(30, "EXT", "CAJA DE LA OTRA EMPRESA", debe=20),
                        renglon(31, "EXT", "CAJA FUERTE", debe=5), renglon(31, "EXT", "CAJA DE LA OTRA EMPRESA", haber=5)])
        m = cajas.movimientos(cajas.leer_detalle(ruta), "AA", cfg, "aa.xlsx")
        self.assertEqual([("AA", "Caja AA", -20), ("A", "Caja A", 20), ("AA", "Caja AA", 5), ("A", "Caja A", -5)],
                         [(x["Empresa"], x["Banco / Cuenta"], x["Importe"]) for x in m])
        self.assertEqual({"Transferencia Interna"}, {x["Categoria"] for x in m})
        self.assertTrue(m[1]["Origen"].startswith("Tango caja A"))   # suma en el saldo de la Caja A
        self.assertEqual("DEPOSITO BANCO INVENTADO", m[1]["Leyenda"])
        # Sin espejo en el perfil, A no recibe nada.
        self.assertEqual(["AA", "AA"], [x["Empresa"] for x in cajas.movimientos(cajas.leer_detalle(ruta), "AA", CFG, "aa.xlsx")])

    def test_pago_a_la_empresa_hermana_es_interno(self):
        m = self.mov([renglon(10, "O/P", "ACREEDORES VARIOS", debe=10), renglon(10, "O/P", "CAJA FUERTE", haber=10, razon="NAVAR S.A.")], "AA")
        self.assertEqual("Transferencia Interna", m[0]["Categoria"])

    def test_cobro_con_cliente_es_cobranza_aunque_toque_otra_cuenta(self):
        m = self.mov([renglon(11, "REC", "CAJA CHICA", debe=900), renglon(11, "REC", "DEUDORES VARIOS", haber=950),
                      renglon(11, "REC", "SUELDOS Y JORNALES", debe=50)])
        self.assertEqual("Cobranza Facturas", m[0]["Categoria"])

    def test_anulacion_resta_en_la_misma_fila(self):
        m = self.mov([renglon(12, "REC", "CAJA CHICA", debe=100), renglon(12, "REC", "DEUDORES VARIOS", haber=100),
                      renglon(13, "REV", "CAJA CHICA", haber=100), renglon(13, "REV", "DEUDORES VARIOS", debe=100)])
        self.assertEqual([("Ingreso", "Cobranza Facturas", 100), ("Ingreso", "Cobranza Facturas", -100)],
                         [(x["Tipo"], x["Categoria"], x["Importe"]) for x in m])

    def test_aporte_de_un_socio_y_su_anulacion_se_cancelan(self):
        m = self.mov([renglon(14, "REC", "CAJA CHICA", debe=1000), renglon(14, "REC", "CTA. PARTICULAR SOCIO", haber=1000),
                      renglon(15, "REV", "CAJA CHICA", haber=1000), renglon(15, "REV", "CTA. PARTICULAR SOCIO", debe=1000)])
        self.assertEqual({("Ingreso", "Otros")}, {(x["Tipo"], x["Categoria"]) for x in m})
        self.assertEqual(0, sum(x["Importe"] for x in m))

    def test_cuenta_que_no_es_caja_no_entra(self):
        self.assertEqual([], self.mov([renglon(16, "OPF", "GASTOS VARIOS", debe=1), renglon(16, "OPF", "OTRA CUENTA", haber=1)]))

    def test_columnas_en_otro_orden_y_falta_columna(self):
        ruta = os.path.join(self.dir, "x.xlsx")
        escribir(ruta, [renglon(17, "OPF", "CAJA CHICA", haber=3)], list(reversed(ENC)))
        self.assertEqual(-3, cajas.movimientos(cajas.leer_detalle(ruta), "A", CFG, "x")[0]["Importe"])
        escribir(ruta, [renglon(17, "OPF", "CAJA CHICA", haber=3)], [c for c in ENC if c != "Desc. cuenta"])
        with self.assertRaisesRegex(ValueError, "faltan columnas"):
            cajas.leer_detalle(ruta)

    def test_main_exige_los_dos_archivos_y_escribe_con_las_columnas_nuevas(self):
        a = os.path.join(self.dir, "A tesoreria detalle 2026-10-07.xlsx")
        escribir(a, [renglon(1, "OPF", "GASTOS VARIOS", debe=5), renglon(1, "OPF", "CAJA CONTADO", haber=5, leyenda="EJEMPLO")])
        with self.assertRaisesRegex(ValueError, "no se publica ninguna caja"):
            cajas.main(["--a", a, "--aa", os.path.join(self.dir, "no-existe.xlsx"), "--salidas", self.dir])
        self.assertFalse(any(n.startswith("para_pegar_cajas") for n in os.listdir(self.dir)))
        aa = os.path.join(self.dir, "AA tesoreria detalle 2026-10-07.xlsx")
        escribir(aa, [])
        import io
        from contextlib import redirect_stdout
        with redirect_stdout(io.StringIO()):
            cajas.main(["--a", a, "--aa", aa, "--salidas", self.dir, "--hoy", "2026-10-07", "--cliente", "navar"])
        ws = openpyxl.load_workbook(os.path.join(self.dir, "para_pegar_cajas_2026-10-07.xlsx")).active
        enc = [c.value for c in ws[1]]
        self.assertEqual(["Cuenta Tango", "Leyenda"], enc[-2:])
        self.assertEqual("Movimientos", ws.title)

    def test_origen_de_a_no_cae_en_el_patron_de_aa(self):
        import fnmatch
        self.assertFalse(fnmatch.fnmatchcase(cajas.MARCAS["A"] + " · cajas · x", "Tango AA*"))
        self.assertTrue(fnmatch.fnmatchcase(cajas.MARCAS["AA"] + " · cajas · x", "Tango AA*"))

    def test_vigilante_corre_cajas_con_el_mas_nuevo_de_cada_empresa_y_solo_si_estan_las_dos(self):
        import time
        from unittest import mock
        sys.path.insert(0, os.path.join(BASE, "clientes", "navar", "herramientas"))
        import vigilante
        viejo = os.path.join(self.dir, "Tesoreria A detalle", "A tesoreria detalle 2026-10-05.xlsx")
        nuevo = os.path.join(self.dir, "Tesoreria A detalle", "A tesoreria detalle 2026-10-06.xlsx")
        for ruta in (viejo, nuevo):
            os.makedirs(os.path.dirname(ruta), exist_ok=True)
            escribir(ruta, [])
        os.utime(viejo, (time.time() - 300, time.time() - 300))
        with mock.patch.object(vigilante, "DRIVE", self.dir):
            f = {x["nombre"]: x for x in vigilante.fuentes(datetime.date(2026, 10, 6))}
            self.assertEqual([], f["cajas"]["archivos"])            # falta la de AA: no corre
            aa = os.path.join(self.dir, "Tesoreria AA detalle", "AA tesoreria detalle 2026-10-06.xlsx")
            os.makedirs(os.path.dirname(aa))
            escribir(aa, [])
            f = {x["nombre"]: x for x in vigilante.fuentes(datetime.date(2026, 10, 6))}
            cmd = f["cajas"]["cmd"]()
        self.assertEqual("cajas.py", os.path.basename(cmd[1]))
        self.assertEqual(nuevo, cmd[cmd.index("--a") + 1])
        self.assertEqual(aa, cmd[cmd.index("--aa") + 1])
        self.assertEqual(vigilante.STAGING_CAJAS, cmd[cmd.index("--salidas") + 1])
        self.assertNotIn("tesoreria_aa", f)


if __name__ == "__main__":
    unittest.main()
