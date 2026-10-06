"""Ejemplos inventados: cada regla del cruce banco ↔ Tango y la cuenta de control."""
import datetime as dt
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from lector import cruce as cz

D = dt.datetime
CFG = {
    "empresa": "A", "cuit_propio": "30-99999999-5",
    "cuentas": {"GALICIA 111": 5, "MACRO 222": 25, "BBVA 333/1": 89, "BBVA 333/2": 89},
    "dias_antes": 3, "dias_despues": 5, "tolerancia": 1,
    "categorias_gastos": ["Impuestos", "Gastos Bancarios"],
    "categorias_deposito": ["Cheques", "Descuento de Cheques"],
    "comprobantes_deposito": ["BDM", "BDG"],
    "intereses_descuento": {"todas": ["INTER"], "alguna": ["VTA", "VALORES"]},
    "cuentas_efectivo": ["CAJA"], "cuenta_cheques_terceros": "VALORES A DEPOSITAR",
    "comprobantes_cobro": ["REC"], "comprobantes_pago": ["O/P", "OPF"],
}
ENC_MOV = ["ID", "Fecha", "Empresa", "Tipo", "Categoria", "Concepto / Detalle", "Importe", "Medio de Pago",
           "Banco / Cuenta", "Origen", "Estado", "Referencia", "Semana (lunes)", "Observaciones"]
ENC_TANGO = ["Banco", "Cód. comprobante", "Comprobante", "Nro. interno", "Cód. cuenta", "Desc. contable",
             "Debe (cte) (renglón)", "Haber (cte) (renglón)", "Total comp. (cte)", "CUIT cliente (encab.)",
             "CUIT proveedor (encab.)", "Razón social (encab.)", "Proveedor (encab.)", "Leyenda", "Fecha de emisión",
             "Desc. cuenta"]
CUIT_A = "20-11111111-2"      # CUITs inventados con dígito verificador válido
CUIT_B = "27-22222222-8"


class _Base(unittest.TestCase):
    """Arma los dos Excel inventados y corre el cruce. No tiene pruebas propias."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.mov, self.tg, self.n = [], [], 0

    # ---- armado de los dos Excel inventados
    def banco(self, fecha, cuenta, importe, concepto="MOV", categoria="Otros", origen="Extracto X", obs=""):
        self.mov.append([len(self.mov) + 1, fecha, "A", "", categoria, concepto, importe, "", cuenta, origen,
                         "Real", None, None, obs])

    def tango(self, fecha, tipo, cod, importe, cuenta="BANCO", cuit="", nombre="", leyenda="", interno=None,
              banco="X"):
        self.n += 1
        debe, haber = (importe, 0) if importe >= 0 else (0, -importe)
        self.tg.append([banco if cod in (5, 25, 89, 106) else None, tipo, " %010d/0" % self.n, interno or self.n,
                        cod, cuenta, debe, haber, abs(importe), cuit if tipo == "REC" else None,
                        cuit if tipo != "REC" else None, nombre if tipo == "REC" else None,
                        nombre if tipo != "REC" else None, leyenda, fecha])

    def correr(self, mes="2026-08", parchar=True):
        sheet, tango = self.base / "sheet.xlsx", self.base / "tango.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Cash"
        ws = wb.create_sheet("Movimientos")
        ws.append(ENC_MOV)
        for f in self.mov:
            ws.append(f)
        wb.save(sheet)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Detalle de comprobantes"
        # la columna "Fecha" (la del movimiento) solo va si alguna prueba la cargó
        ancho = max([len(f) for f in self.tg] + [len(ENC_TANGO)])
        ws.append(ENC_TANGO + ["Fecha"][:ancho - len(ENC_TANGO)])
        for f in self.tg:
            ws.append(f)
        ws.append([None] * len(ENC_TANGO))           # la fila casi vacía del final del export
        wb.save(tango)
        def correr():
            return cz.correr("prueba", mes, str(sheet), str(tango), str(self.base / "salidas"),
                             hoy=dt.date(2026, 9, 28))
        if parchar:
            with patch.object(cz, "cargar_config", return_value=cz_config()):
                inf, md = correr()
        else:
            inf, md = correr()
        self.assertTrue(inf["control_ok"], md)
        return inf

    @staticmethod
    def regla_de(inf, importe):
        for p in inf["pares"]:
            if any(abs(b["c"] - round(importe * 100)) <= 1 for b in p["banco"]):
                return p["regla"]
        return None


class Cruce(_Base):
    def test_exacto_banco_despues_y_antes(self):
        self.banco(D(2026, 8, 5), "GALICIA 111", 1000.0)          # 2 días después de Tango
        self.tango(D(2026, 8, 3), "REC", 5, 1000.0)
        self.banco(D(2026, 8, 9), "GALICIA 111", -2500.5)         # 1 día antes que Tango
        self.tango(D(2026, 8, 10), "O/P", 5, -2500.5)
        inf = self.correr()
        self.assertEqual(self.regla_de(inf, 1000), "exacto")
        self.assertEqual(self.regla_de(inf, -2500.5), "exacto")
        self.assertEqual(inf["solo_banco"], [])
        self.assertEqual(inf["solo_tango"], [])

    def test_fecha_corrida(self):
        self.banco(D(2026, 8, 24), "MACRO 222", -4321.0, "DEBITO PRESTAMOS", "Prestamo")
        self.tango(D(2026, 8, 31), "FPR", 25, -4321.0)               # 7 días después: fuera de la ventana
        inf = self.correr()
        self.assertEqual(self.regla_de(inf, -4321), "fecha corrida")

    def test_fecha_corrida_con_competencia_no_empareja(self):
        self.banco(D(2026, 8, 24), "MACRO 222", -500.0, "DEBITO", "Prestamo")
        self.banco(D(2026, 8, 20), "MACRO 222", -500.0, "DEBITO", "Prestamo")
        self.tango(D(2026, 8, 31), "FPR", 25, -500.0)
        inf = self.correr()
        self.assertEqual(inf["pares"], [])

    def test_muy_lejos_no_empareja_y_deja_pista(self):
        self.banco(D(2026, 8, 30), "GALICIA 111", -700.0, categoria="Proveedores")
        self.tango(D(2026, 8, 10), "O/P", 5, -700.0)                # 20 días antes: ni corrida
        inf = self.correr()
        self.assertEqual(len(inf["solo_banco"]), 1)
        self.assertIn("mismo importe", inf["solo_banco"][0]["pista"])
        self.assertEqual(inf["solo_tango"][0]["estado"], "no aparece en el banco: revisar")

    def test_desempate_por_cuit(self):
        self.banco(D(2026, 8, 12), "MACRO 222", -500.0, "TRANSF 27222222228 VAR")
        self.banco(D(2026, 8, 12), "MACRO 222", -500.0, "TRANSF 20111111112 VAR")
        self.tango(D(2026, 8, 12), "O/P", 25, -500.0, cuit=CUIT_A, nombre="PROVEEDOR A")
        self.tango(D(2026, 8, 12), "O/P", 25, -500.0, cuit=CUIT_B, nombre="PROVEEDOR B")
        inf = self.correr()
        for p in inf["pares"]:
            self.assertEqual(p["banco"][0]["cuit"], p["tango"][0]["cuit"])

    def test_mismo_cuit_una_op_contra_dos_transferencias(self):
        self.banco(D(2026, 8, 14), "MACRO 222", -300.0, "TRANSF 20111111112 VAR")
        self.banco(D(2026, 8, 15), "MACRO 222", -200.0, "TRANSF 20111111112 VAR")
        self.tango(D(2026, 8, 14), "O/P", 25, -500.0, cuit=CUIT_A)
        inf = self.correr()
        self.assertEqual(self.regla_de(inf, -300), "mismo CUIT")
        self.assertEqual(len(inf["pares"][0]["banco"]), 2)

    def test_descuento_neto_de_intereses(self):
        self.banco(D(2026, 8, 6), "MACRO 222", 9600.0, "033000111", "Descuento de Cheques")
        self.banco(D(2026, 8, 6), "MACRO 222", 1900.0, "033000222", "Descuento de Cheques")
        self.tango(D(2026, 8, 6), "BDM", 25, 10000.0)
        self.tango(D(2026, 8, 6), "BDM", 25, 2000.0)
        self.tango(D(2026, 8, 6), "FPR", 25, -400.0, leyenda="INTERESES VTA.CH.")
        self.tango(D(2026, 8, 6), "FPR", 25, -100.0, leyenda="INTERESES VTA VALORES MACRO")
        inf = self.correr()
        self.assertEqual(self.regla_de(inf, 9600), "descuento neto")
        self.assertEqual(self.regla_de(inf, 1900), "descuento neto")
        self.assertEqual(inf["solo_tango"], [])

    def test_bloque_del_dia(self):
        for v in (100.0, 250.0, 650.0):
            self.banco(D(2026, 8, 26), "GALICIA 111", v, "CREDITO DESCUENTO DOCUMENTO", "Descuento de Cheques")
        self.tango(D(2026, 8, 26), "BDG", 5, 1000.0)
        inf = self.correr()
        self.assertEqual(self.regla_de(inf, 100), "bloque del día")
        self.assertEqual(len(inf["pares"][0]["banco"]), 3)

    def test_boleta_acreditada_al_dia_siguiente_en_dos_renglones(self):
        self.banco(D(2026, 8, 12), "MACRO 222", 300.0, "ACREDITACION CHEQUE", "Cheques")
        self.banco(D(2026, 8, 12), "MACRO 222", 700.0, "ACREDITACION CHEQUE", "Cheques")
        self.banco(D(2026, 8, 12), "MACRO 222", 55.0, "ACREDITACION CHEQUE", "Cheques")   # de otra boleta
        self.banco(D(2026, 8, 12), "MACRO 222", 80.0, "ACREDITACION CHEQUE", "Cheques")   # sin boleta: el bloque no da
        self.tango(D(2026, 8, 11), "BDM", 25, 1000.0)
        self.tango(D(2026, 8, 13), "BDM", 25, 55.0)
        inf = self.correr()
        self.assertEqual(self.regla_de(inf, 300), "agrupado")
        self.assertEqual(self.regla_de(inf, 55), "exacto")
        self.assertEqual([b["c"] for b in inf["solo_banco"]], [8000])

    def test_ambiguo_va_a_revisar(self):
        # 100+200 y 150+150 dan 300: dos combinaciones posibles → no se empareja
        for v in (100.0, 200.0, 150.0, 150.0):
            self.banco(D(2026, 8, 18), "MACRO 222", -v, "PAGO", "Proveedores")
        self.tango(D(2026, 8, 18), "OPF", 25, -300.0)
        inf = self.correr()
        self.assertEqual(inf["pares"], [])
        self.assertTrue(any(r["lado"] == "Tango" and "más de una combinación" in r["motivo"] for r in inf["revisar"]))
        self.assertEqual(len(inf["solo_tango"]), 0)

    def test_transferencia_entre_bancos_propios(self):
        self.banco(D(2026, 8, 10), "MACRO 222", -25000.0, "TRANSF 30999999995 VAR", "Transferencia Interna")
        self.banco(D(2026, 8, 10), "GALICIA 111", 25000.0, "TRANSFERENCIA DE CUENTA PROPIA", "Transferencia Interna")
        self.tango(D(2026, 8, 10), "EXT", 25, -25000.0, interno=900)
        self.tango(D(2026, 8, 10), "EXT", 5, 25000.0, interno=900)
        inf = self.correr()
        self.assertEqual(len(inf["pares"]), 2)
        self.assertEqual({p["regla"] for p in inf["pares"]}, {"exacto"})

    def test_internas_entre_las_dos_bbva(self):
        self.banco(D(2026, 8, 6), "BBVA 333/2", 4000.0, "OPERACION DE RECAUDO", "Cobranza Facturas")
        self.banco(D(2026, 8, 6), "BBVA 333/2", -4000.0, "TRANSFERENCIA CUENTAS PAIS", "Transferencia Interna")
        self.banco(D(2026, 8, 6), "BBVA 333/1", 4000.0, "TRANSFERENCIA CUENTAS PAIS", "Transferencia Interna")
        self.tango(D(2026, 8, 6), "REC", 89, 4000.0, cuit=CUIT_A)
        inf = self.correr()
        self.assertEqual(len(inf["internas"]), 1)
        self.assertEqual(len(inf["pares"]), 1)
        self.assertEqual(inf["pares"][0]["banco"][0]["cuenta"], "BBVA 333/2")
        self.assertEqual(inf["solo_banco"], [])

    def test_solo_tango_de_cuenta_compartida_se_cuenta_una_vez(self):
        self.banco(D(2026, 8, 6), "BBVA 333/1", -1.0, "COMISION", "Gastos Bancarios")
        self.banco(D(2026, 8, 6), "BBVA 333/2", -1.0, "COMISION", "Gastos Bancarios")
        self.tango(D(2026, 8, 6), "O/P", 89, -777.0)
        inf = self.correr()
        self.assertEqual(sorted(r["n_solo_tango"] for r in inf["resumen"]), [0, 1])

    def test_gastos_van_a_su_bloque(self):
        self.banco(D(2026, 8, 3), "GALICIA 111", -12.34, "IMP DEBITOS Y CREDITOS", "Impuestos")
        self.banco(D(2026, 8, 3), "GALICIA 111", -50.0, "COMISION", "Gastos Bancarios")
        inf = self.correr()
        self.assertEqual(len(inf["gastos"]), 2)
        self.assertEqual(inf["solo_banco"], [])

    def test_no_pasa_por_banco(self):
        self.tango(D(2026, 8, 7), "REC", 1, 800.0, cuenta="CAJA", cuit=CUIT_A, nombre="CLIENTE", interno=50)
        self.tango(D(2026, 8, 7), "REC", 3, 1200.0, cuenta="VALORES A DEPOSITAR", cuit=CUIT_A, interno=50)
        self.tango(D(2026, 8, 7), "REC", 10, -2000.0, cuenta="DEUDORES POR VENTAS", cuit=CUIT_A, interno=50)
        self.tango(D(2026, 8, 9), "O/P", 3, -900.0, cuenta="VALORES A DEPOSITAR", cuit=CUIT_B, interno=60)
        self.tango(D(2026, 8, 9), "O/P", 17, 900.0, cuenta="PROVEEDORES", cuit=CUIT_B, interno=60)
        self.banco(D(2026, 8, 9), "GALICIA 111", -1.0, "COMISION", "Gastos Bancarios")
        inf = self.correr()
        clases = sorted(g["clase"] for g in inf["no_banco"])
        self.assertEqual(clases, ["Cheque de tercero endosado", "Efectivo (cobro)"])
        self.assertEqual(inf["cheques_recibidos"], [1, 120000])

    def test_fecha_imposible(self):
        self.banco(D(2026, 8, 13), "MACRO 222", 8980.0, "ACREDITACION CHEQUE", "Cheques")
        self.tango(D(2036, 8, 12), "BDM", 25, 8980.0)
        inf = self.correr()
        self.assertEqual(inf["pares"], [])
        self.assertTrue(any("fecha imposible" in r["motivo"] and "12/08/2026" in r["motivo"] for r in inf["revisar"]))
        self.assertIn("fecha imposible", inf["solo_banco"][0]["pista"])

    def test_par_cruzado_de_mes(self):
        self.banco(D(2026, 9, 2), "GALICIA 111", 5000.0, "TRANSF 20111111112", "Cobranza Facturas")
        self.tango(D(2026, 8, 31), "REC", 5, 5000.0, cuit=CUIT_A)
        inf = self.correr()
        self.assertEqual(len(inf["pares"]), 1)
        self.assertEqual(inf["solo_tango"], [])

    def test_ignora_manual_y_caja_aa(self):
        self.banco(D(2026, 8, 5), "GALICIA 111", 999.0, origen="Manual")
        self.banco(D(2026, 8, 5), "Caja AA", 999.0, origen="Tango AA · movimientos")
        self.banco(D(2026, 8, 5), "GALICIA 111", 10.0, "COMISION", "Gastos Bancarios")
        inf = self.correr()
        self.assertEqual(inf["resumen"][0]["n"], 1)
        self.assertEqual(inf["sin_par_banco"], {})

    def test_anulado_con_reversion(self):
        self.tango(D(2026, 8, 18), "OPF", 25, -609.0, leyenda="PAGO GREMIO")
        self.tango(D(2026, 8, 18), "REV", 25, 609.0, leyenda="PAGO GREMIO")
        self.banco(D(2026, 8, 18), "MACRO 222", -1.0, "COMISION", "Gastos Bancarios")
        inf = self.correr()
        self.assertTrue(all(t.get("anulado") for t in inf["solo_tango"]))
        self.assertEqual(inf["resumen"][0]["n_solo_tango"], 0)

    def test_excel_y_resumen_se_escriben(self):
        self.banco(D(2026, 8, 5), "GALICIA 111", 1000.0)
        self.tango(D(2026, 8, 5), "REC", 5, 1000.0)
        self.correr()
        wb = openpyxl.load_workbook(self.base / "salidas" / "cruce_2026-08.xlsx")
        self.assertEqual(wb.sheetnames, ["Resumen", "Conciliado", "Solo en banco", "Gastos bancarios",
                                         "Solo en Tango", "No pasa por banco", "Revisar"])
        self.assertTrue((self.base / "salidas" / "resumen_cruce_2026-08.md").exists())


class ReglasDeSeptiembre(_Base):
    """Casos que aparecieron con los datos de septiembre (tarea 45). Datos inventados."""

    def par_de(self, inf, importe):
        for p in inf["pares"]:
            if any(abs(b["c"] - round(importe * 100)) <= 1 for b in p["banco"]):
                return p
        return None

    def test_impuestos_agrupados_en_dias_distintos(self):
        self.banco(D(2026, 9, 16), "MACRO 222", -2210.0 * 1000, "IMP. AFIP", "Impuestos")
        self.tango(D(2026, 9, 15), "OPF", 25, -2000.0 * 1000, leyenda="PAGO APORTES")
        self.tango(D(2026, 9, 15), "OPF", 25, -200.0 * 1000, leyenda="INTERES APORTES")
        self.tango(D(2026, 9, 16), "OPF", 25, -10.0 * 1000, leyenda="PAGO INTERES APORTE")
        par = self.par_de(self.correr("2026-09"), -2210000)
        self.assertEqual(("impuestos", "Sugerido", 3), (par["regla"], par["nivel"], len(par["tango"])))
        self.assertIn("pagos de impuestos", par["criterio"])

    def test_plan_de_pagos_una_orden_varios_vep(self):
        for v in (490.0, 1010.0, 2672.5, 15845.0):
            self.banco(D(2026, 8, 18), "MACRO 222", -v * 1000, "ARCA", "Impuestos")
        self.banco(D(2026, 8, 18), "MACRO 222", -60.0, "IMPDBCR 25413", "Impuestos")    # chico: no entra
        self.tango(D(2026, 8, 18), "OPF", 25, -20017.5 * 1000, leyenda="Planes de pago ARCA")
        inf = self.correr()
        par = self.par_de(inf, -490000)
        self.assertEqual(("impuestos", "Sugerido", 4), (par["regla"], par["nivel"], len(par["banco"])))
        self.assertEqual(len(inf["gastos"]), 1)

    def test_impuestos_no_toma_pagos_a_proveedores_ni_montos_chicos(self):
        self.banco(D(2026, 9, 16), "MACRO 222", -300.0 * 1000, "IMP. AFIP", "Impuestos")
        self.tango(D(2026, 9, 16), "OPF", 25, -100.0 * 1000, cuit=CUIT_A)   # tiene CUIT: es un tercero
        self.tango(D(2026, 9, 16), "OPF", 25, -200.0 * 1000)
        self.banco(D(2026, 9, 17), "MACRO 222", -500.0, "IMP. AFIP", "Impuestos")   # menos del mínimo
        self.tango(D(2026, 9, 17), "OPF", 25, -300.0)
        self.tango(D(2026, 9, 17), "OPF", 25, -200.0)
        inf = self.correr("2026-09")
        self.assertEqual(inf["pares"], [])
        self.assertEqual(len(inf["gastos"]), 2)

    def test_impuestos_ambiguo_va_a_revisar(self):
        self.banco(D(2026, 9, 16), "MACRO 222", -300.0 * 1000, "IMP. AFIP", "Impuestos")
        for v in (100.0, 200.0, 150.0, 150.0):
            self.tango(D(2026, 9, 16), "OPF", 25, -v * 1000)
        inf = self.correr("2026-09")
        self.assertEqual(inf["pares"], [])
        self.assertTrue(any("impuestos agrupados" in r["motivo"] for r in inf["revisar"]))
        self.assertEqual(inf["gastos"], [])

    def test_descuento_neto_corrido(self):
        # el interés se carga el día del banco y la boleta dos días después
        self.banco(D(2026, 9, 14), "MACRO 222", 9600.0, "033000111", "Descuento de Cheques")
        self.tango(D(2026, 9, 16), "BDM", 25, 10000.0)
        self.tango(D(2026, 9, 14), "FPR", 25, -400.0, leyenda="INTERESES VTA VALORES MACRO")
        par = self.par_de(self.correr("2026-09"), 9600)
        self.assertEqual(("descuento neto", "Sugerido"), (par["regla"], par["nivel"]))
        self.assertIn("cargados en días distintos", par["criterio"])

    def test_descuento_del_mismo_dia_sigue_siendo_seguro(self):
        self.banco(D(2026, 9, 14), "MACRO 222", 9600.0, "033000111", "Descuento de Cheques")
        self.tango(D(2026, 9, 14), "BDM", 25, 10000.0)
        self.tango(D(2026, 9, 14), "FPR", 25, -400.0, leyenda="INTERESES VTA VALORES MACRO")
        self.assertEqual("Seguro", self.par_de(self.correr("2026-09"), 9600)["nivel"])

    def test_deposito_acreditado_en_dos_dias(self):
        self.banco(D(2026, 8, 5), "MACRO 222", 1600.0, "DEPOSITO CANJE", "Cheques")
        self.banco(D(2026, 8, 6), "MACRO 222", 1200.0, "ACREDITACION CHEQUE", "Cheques")
        self.banco(D(2026, 8, 6), "MACRO 222", 77.0, "ACREDITACION CHEQUE", "Cheques")    # de otra cosa
        self.tango(D(2026, 8, 5), "BDM", 25, 2800.0)
        inf = self.correr()
        par = self.par_de(inf, 1600)
        self.assertEqual(("agrupado", "Posible", 2), (par["regla"], par["nivel"], len(par["banco"])))
        self.assertIn("acreditados entre el 05/08 y el 06/08", par["criterio"])
        self.assertEqual([b["c"] for b in inf["solo_banco"]], [7700])

    def test_pista_de_error_de_tipeo(self):
        self.banco(D(2026, 9, 14), "MACRO 222", -12345678.0, "TARJETA DE CREDITO", "Otros")
        self.tango(D(2026, 8, 31), "EXT", 25, -12354678.0)
        inf = self.correr("2026-09")
        self.assertIn("error de tipeo", inf["solo_banco"][0]["pista"])


class FechaDelMovimiento(_Base):
    """Tango tiene dos fechas; el cruce usa 'Fecha' (la del movimiento) y no 'Fecha de emisión'."""

    def test_usa_fecha_y_no_la_de_emision(self):
        self.banco(D(2026, 8, 13), "MACRO 222", 8980.0, "ACREDITACION CHEQUE", "Cheques")
        self.tango(D(2036, 8, 12), "BDM", 25, 8980.0)            # emisión mal cargada
        self.tg[-1] += [None, D(2026, 8, 12)]                      # Desc. cuenta, Fecha (la buena)
        inf = self.correr()
        self.assertEqual(len(inf["pares"]), 1)
        self.assertEqual(inf["revisar"], [])
        self.assertFalse(inf["fecha_de_emision"])

    def test_sin_columna_fecha_avisa(self):
        self.banco(D(2026, 8, 13), "MACRO 222", 10.0, "X")
        self.tango(D(2026, 8, 13), "REC", 25, 10.0)
        self.assertTrue(self.correr()["fecha_de_emision"])


class MejorasTarea46(_Base):
    """Cheques propios, neto con cargos, varios días, anulaciones, interés mal escrito y mes vencido."""

    def cheques(self, filas):
        ruta = self.base / "A cheques propios 2026-10-05.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["Nro. de cheque", "Banco", "Razón social", "Fecha de emisión", "Fecha del cheque",
                          "Importe mon. cta.", "Estado", "Cuenta emisión"])
        for f in filas:
            wb.active.append(f)
        wb.save(ruta)
        return str(ruta)

    def correr_con_cheques(self, filas, mes="2026-08"):
        ruta = self.cheques(filas)
        original = cz.correr
        with patch.object(cz, "correr", lambda *a, **k: original(*a, cheques=ruta, **k)):
            return self.correr(mes)

    def par_de(self, inf, importe):
        for p in inf["pares"]:
            if any(abs(b["c"] - round(importe * 100)) <= 1 for b in p["banco"]):
                return p
        return None

    def tango_banco(self, fecha, tipo, cod, importe, nombre_cuenta, **kw):
        """Como tango(), pero con 'Desc. cuenta' (el nombre del banco que usan los cheques)."""
        self.tango(fecha, tipo, cod, importe, **kw)
        self.tg[-1].append(nombre_cuenta)

    def test_cheque_propio_contra_debito_por_canje(self):
        self.banco(D(2026, 8, 11), "MACRO 222", -25000.0, "48HS. CANJE ZONAL", "Cheques")
        self.tango_banco(D(2026, 8, 2), "REC", 25, 1.0, "BANCO MACRO INVENTADO")   # solo para conocer la cuenta
        self.banco(D(2026, 8, 2), "MACRO 222", 1.0, "TRANSF")
        inf = self.correr_con_cheques([[1001, "MACRO", "PROVEEDOR X", D(2025, 12, 4), D(2026, 8, 10), 25000.0,
                                        "Al Cobro", "BANCO MACRO INVENTADO"]])
        par = self.par_de(inf, -25000)
        self.assertEqual(("cheque propio", "Seguro"), (par["regla"], par["nivel"]))
        self.assertIn("cheque 1001", par["criterio"])

    def test_cheques_semanales_del_mismo_importe(self):
        self.tango_banco(D(2026, 8, 2), "REC", 25, 1.0, "BANCO MACRO INVENTADO")
        self.banco(D(2026, 8, 2), "MACRO 222", 1.0, "TRANSF")
        for d in (11, 19, 26):
            self.banco(D(2026, 8, d), "MACRO 222", -25000.0, "48HS. CANJE ZONAL", "Cheques")
        inf = self.correr_con_cheques([[n, "MACRO", "PROVEEDOR X", D(2025, 12, 4), D(2026, 8, d), 25000.0,
                                        "Al Cobro", "BANCO MACRO INVENTADO"] for n, d in ((1, 10), (2, 17), (3, 24))])
        pares = sorted((p for p in inf["pares"] if p["regla"] == "cheque propio"), key=lambda p: p["banco"][0]["fecha"])
        self.assertEqual(["cheque 1", "cheque 2", "cheque 3"], [p["tango"][0]["comprobante"] for p in pares])
        self.assertEqual({"Sugerido"}, {p["nivel"] for p in pares})
        self.assertEqual(3, inf["cheques_leidos"])

    def test_orden_de_pago_con_cheques_diferidos_no_es_diferencia(self):
        self.tango_banco(D(2026, 8, 3), "O/P", 25, -30000.0, "BANCO MACRO INVENTADO", nombre="PROVEEDOR X")
        self.banco(D(2026, 8, 3), "MACRO 222", -1.0, "COMISION", "Gastos Bancarios")
        inf = self.correr_con_cheques([
            [1, "MACRO", "PROVEEDOR X", D(2026, 8, 3), D(2026, 9, 15), 10000.0, "Al Cobro", "BANCO MACRO INVENTADO"],
            [2, "MACRO", "PROVEEDOR X", D(2026, 8, 3), D(2026, 10, 15), 20000.0, "Al Cobro", "BANCO MACRO INVENTADO"]])
        self.assertEqual([t for t in inf["solo_tango"] if not t.get("anulado")], [])
        self.assertEqual(["Pago con cheques propios diferidos"], [g["clase"] for g in inf["no_banco"]])
        self.assertIn("2 cheques diferidos", inf["no_banco"][0]["texto"])

    def test_sin_cheques_sigue_andando(self):
        self.banco(D(2026, 8, 11), "MACRO 222", -25000.0, "48HS. CANJE ZONAL", "Cheques")
        inf = self.correr()
        self.assertEqual(len(inf["solo_banco"]), 1)

    def test_prestamo_neto_de_sellos(self):
        self.banco(D(2026, 7, 31), "MACRO 222", 99000.0, "N/C OPERAC PRESTAMOS", "Prestamo")
        self.tango(D(2026, 7, 31), "REC", 25, 100000.0, leyenda="PRESTAMO 24 CUOTAS")
        self.tango(D(2026, 7, 31), "OPF", 25, -1000.0, leyenda="IMP. A SELLOS PRESTAMO")
        par = self.par_de(self.correr("2026-07"), 99000)
        self.assertEqual(("neto con cargos", "Sugerido", 2), (par["regla"], par["nivel"], len(par["tango"])))

    def test_echeq_en_tres_dias(self):
        for f, v in ((13, 2600.0), (14, 500.0), (14, 9800.0), (15, 3000.0)):
            self.banco(D(2026, 7, f), "GALICIA 111", v, "G.DE ECHEQ", "Cheques")
        self.tango(D(2026, 7, 13), "BDG", 5, 15900.0)
        par = self.par_de(self.correr("2026-07"), 2600)
        self.assertEqual(("agrupado", "Posible", 4), (par["regla"], par["nivel"], len(par["banco"])))
        self.assertIn("entre el 13/07 y el 15/07", par["criterio"])

    def test_anulacion_que_no_es_rev(self):
        self.banco(D(2026, 7, 16), "GALICIA 111", -2000.0, "TRANSF CP", "Transferencia Interna")
        self.tango(D(2026, 7, 16), "EXT", 5, -2000.0)            # la que va con el banco
        self.tango(D(2026, 7, 16), "EXT", 5, 2000.0)             # la anulación (EXT .../1)
        self.tango(D(2026, 7, 16), "EXT", 5, -2000.0)            # la nueva carga
        inf = self.correr("2026-07")
        self.assertEqual(len(inf["pares"]), 1)
        self.assertTrue(all(t.get("anulado") for t in inf["solo_tango"]))
        self.assertEqual(inf["resumen"][0]["n_solo_tango"], 0)

    def test_interes_mal_escrito(self):
        self.banco(D(2026, 7, 30), "MACRO 222", 9600.0, "0330009", "Descuento de Cheques")
        self.tango(D(2026, 7, 30), "BDM", 25, 10000.0)
        self.tango(D(2026, 7, 30), "FPR", 25, -400.0, leyenda="INTERSES VTA.CH.VS.")
        self.assertEqual("descuento neto", self.par_de(self.correr("2026-07"), 9600)["regla"])

    def test_cuota_cargada_a_mes_vencido(self):
        self.banco(D(2026, 8, 4), "MACRO 222", -12345.0, "DEBITO PAGO PRESTAMO", "Prestamo")
        self.tango(D(2026, 9, 1), "FPR", 25, -12345.0, leyenda="PRESTAMO CUOTA 1")
        par = self.par_de(self.correr("2026-08"), -12345)
        self.assertEqual(("mes vencido", "Sugerido"), (par["regla"], par["nivel"]))
        self.assertIn("Tango lo cargó el 01/09; el banco lo debitó el 04/08", par["criterio"])

    def test_cuotas_no_cuentan_como_cargos_del_banco(self):
        self.banco(D(2026, 8, 4), "MACRO 222", -5.0, "COMISION", "Gastos Bancarios")
        self.tango(D(2026, 8, 20), "FPR", 25, -7000.0, nombre="BANCO INVENTADO", leyenda="PRESTAMO BANCO")
        self.tango(D(2026, 8, 20), "FPR", 25, -300.0, nombre="BANCO INVENTADO", leyenda="INT. Y COMIS.")
        r = self.correr()["resumen"][0]
        self.assertEqual(r["gasto_tango"], -30000)


class Niveles(_Base):
    """Cada par sale con nivel (Seguro / Sugerido / Posible) y un criterio en castellano."""

    def par_de(self, inf, importe):
        for p in inf["pares"] + inf["pares_internas"]:
            if any(abs(b["c"] - round(importe * 100)) <= 1 for b in p["banco"]):
                return p
        self.fail("no hay par para %s" % importe)

    def test_seguro_exacto_unico(self):
        self.banco(D(2026, 8, 5), "GALICIA 111", 1000.0, "TRANSF 20111111112")
        self.tango(D(2026, 8, 3), "REC", 5, 1000.0, cuit=CUIT_A)
        par = self.par_de(self.correr(), 1000)
        self.assertEqual(par["nivel"], "Seguro")
        self.assertIn("importe exacto", par["criterio"])
        self.assertIn("CUIT coincide", par["criterio"])
        self.assertIn("banco 2 días después", par["criterio"])

    def test_seguro_exacto_con_cuit_entre_varios(self):
        self.banco(D(2026, 8, 12), "MACRO 222", -500.0, "TRANSF 27222222228 VAR")
        self.banco(D(2026, 8, 12), "MACRO 222", -500.0, "TRANSF 20111111112 VAR")
        self.tango(D(2026, 8, 12), "O/P", 25, -500.0, cuit=CUIT_A)
        self.tango(D(2026, 8, 12), "O/P", 25, -500.0, cuit=CUIT_B)
        inf = self.correr()
        self.assertEqual({p["nivel"] for p in inf["pares"]}, {"Seguro"})
        self.assertTrue(any("ganó el del mismo CUIT" in p["criterio"] for p in inf["pares"]))
        self.assertTrue(all("CUIT coincide" in p["criterio"] for p in inf["pares"]))

    def test_sugerido_exacto_desempatado_por_fecha(self):
        self.banco(D(2026, 8, 12), "MACRO 222", -500.0, "DEBITO")
        self.banco(D(2026, 8, 14), "MACRO 222", -500.0, "DEBITO")
        self.tango(D(2026, 8, 12), "OPF", 25, -500.0)
        self.tango(D(2026, 8, 13), "OPF", 25, -500.0)
        inf = self.correr()
        self.assertEqual({p["nivel"] for p in inf["pares"]}, {"Sugerido"})
        criterios = sorted(p["criterio"] for p in inf["pares"])
        self.assertTrue(any("ganó la fecha más cercana" in c for c in criterios))
        self.assertTrue(any("después de otro desempate" in c for c in criterios))

    def test_sugerido_bloque_y_seguro_descuento_neto(self):
        for v in (100.0, 900.0):
            self.banco(D(2026, 8, 26), "GALICIA 111", v, "CREDITO DESCUENTO DOCUMENTO", "Descuento de Cheques")
        self.tango(D(2026, 8, 26), "BDG", 5, 1000.0)
        self.banco(D(2026, 8, 6), "MACRO 222", 9600.0, "033000111", "Descuento de Cheques")
        self.tango(D(2026, 8, 6), "BDM", 25, 10000.0)
        self.tango(D(2026, 8, 6), "FPR", 25, -400.0, leyenda="INTERESES VTA.CH.")
        inf = self.correr()
        self.assertEqual(self.par_de(inf, 100)["nivel"], "Sugerido")
        self.assertIn("2 banco ↔ 1 Tango", self.par_de(inf, 100)["criterio"])
        self.assertEqual(self.par_de(inf, 9600)["nivel"], "Seguro")

    def test_posible_agrupado_y_fecha_corrida(self):
        self.banco(D(2026, 8, 12), "MACRO 222", 300.0, "ACREDITACION CHEQUE", "Cheques")
        self.banco(D(2026, 8, 12), "MACRO 222", 700.0, "ACREDITACION CHEQUE", "Cheques")
        self.banco(D(2026, 8, 12), "MACRO 222", 80.0, "ACREDITACION CHEQUE", "Cheques")
        self.tango(D(2026, 8, 11), "BDM", 25, 1000.0)
        self.banco(D(2026, 8, 24), "MACRO 222", -4321.0, "DEBITO PRESTAMOS", "Prestamo")
        self.tango(D(2026, 8, 31), "FPR", 25, -4321.0)
        inf = self.correr()
        self.assertEqual(self.par_de(inf, 300)["nivel"], "Posible")
        self.assertEqual(self.par_de(inf, -4321)["nivel"], "Posible")
        self.assertIn("banco 7 días antes", self.par_de(inf, -4321)["criterio"])

    def test_posible_exacto_con_cuit_distinto(self):
        self.banco(D(2026, 8, 5), "GALICIA 111", -1000.0, "TRANSF 27222222228")
        self.tango(D(2026, 8, 5), "O/P", 5, -1000.0, cuit=CUIT_A)
        par = self.par_de(self.correr(), -1000)
        self.assertEqual(par["nivel"], "Posible")
        self.assertIn("CUIT distinto", par["criterio"])

    def test_internas_son_seguro(self):
        self.banco(D(2026, 8, 6), "BBVA 333/2", -4000.0, "TRANSFERENCIA CUENTAS PAIS", "Transferencia Interna")
        self.banco(D(2026, 8, 6), "BBVA 333/1", 4000.0, "TRANSFERENCIA CUENTAS PAIS", "Transferencia Interna")
        inf = self.correr()
        self.assertEqual([p["nivel"] for p in inf["pares_internas"]], ["Seguro"])

    def test_resumen_por_nivel_y_cuenta(self):
        self.banco(D(2026, 8, 5), "GALICIA 111", 1000.0)
        self.tango(D(2026, 8, 5), "REC", 5, 1000.0)
        self.banco(D(2026, 8, 24), "GALICIA 111", -250.0, "DEBITO", "Prestamo")
        self.tango(D(2026, 8, 31), "FPR", 5, -250.0)
        r = self.correr()["resumen"][0]
        self.assertEqual(r["por_nivel"]["Seguro"], {"n": 1, "c": 100000})
        self.assertEqual(r["por_nivel"]["Posible"], {"n": 1, "c": 25000})
        self.assertEqual(r["por_nivel"]["Sugerido"], {"n": 0, "c": 0})

    def test_el_perfil_cambia_el_nivel(self):
        self.banco(D(2026, 8, 24), "MACRO 222", -4321.0, "DEBITO PRESTAMOS", "Prestamo")
        self.tango(D(2026, 8, 31), "FPR", 25, -4321.0)
        cfg = cz_config()
        cfg["niveles"] = {"fecha corrida": "Sugerido"}
        with patch.object(cz, "cargar_config", return_value=cfg):
            inf = self.correr(parchar=False)
        self.assertEqual(self.par_de(inf, -4321)["nivel"], "Sugerido")

    def test_nivel_invalido_en_el_perfil_frena(self):
        ruta = self.base / "clientes" / "x" / "perfil.json"
        ruta.parent.mkdir(parents=True)
        ruta.write_text('{"cruce": {"cuentas": {"B 1": 5}, "niveles": {"agrupado": "Casi"}}}', encoding="utf-8")
        with patch.object(cz, "BASE_REPO", str(self.base)):
            with self.assertRaises(SystemExit):
                cz.cargar_config("x")


class Cuits(unittest.TestCase):
    def test_cuit_del_banco(self):
        self.assertEqual(cz.cuit_del_banco("TRANSF:WY7ZEPN6-20111111112"), "20111111112")
        self.assertEqual(cz.cuit_del_banco("", "CUIT 27-22222222-8: proveedor"), "27222222228")
        self.assertEqual(cz.cuit_del_banco("CBU 28500331300000083019"), "")          # parte de un número largo
        self.assertEqual(cz.cuit_del_banco("TRANSF 30999999995", propio="30999999995"), "")
        self.assertEqual(cz.cuit_del_banco("TRANSF 20111111113"), "")               # verificador mal


def cz_config():
    cfg = dict(CFG)
    cfg["cuit_propio"] = "30999999995"
    return cfg


if __name__ == "__main__":
    unittest.main()
