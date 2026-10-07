# -*- coding: utf-8 -*-
"""Tarea 53: las dos cajas y la deuda en dos bloques en el tablero. Datos inventados."""
import os
import sys
import datetime
import tempfile
import unittest

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, BASE)
from lector import cash_limpio as CL   # noqa: E402
from dashboard import datos as DATOS   # noqa: E402

HOY = "2026-10-07"


def contrato_inventado():
    def mov(fecha, banco, importe, origen, unidad, estado="Real"):
        return {"fecha": fecha, "banco": banco.upper(), "importe_caja": importe, "importe": abs(importe),
                "origen": origen, "estado_caja": estado, "estado": "REAL", "unidad": unidad}
    return {
        "_origen": {"lector": "lector/cash_limpio.py"},
        "caja_hoy": 140.0, "caja_por_unidad": {"A": 40.0, "AA": 100.0},
        "saldos": [
            {"fecha": "2026-10-01", "banco": "VARIOS", "unidad": "AA", "cuenta": "", "saldo": 100.0, "origen": "Manual"},
            {"fecha": "2026-10-01", "banco": "CAJA A", "unidad": "A", "cuenta": "", "saldo": 50.0, "origen": "Manual"},
            {"fecha": "2026-10-05", "banco": "BANCO X", "unidad": "A", "cuenta": "111", "saldo": -10.0, "origen": "Extracto BANCO X"},
        ],
        "movimientos": [
            mov("2026-10-03", "Caja A", -20.0, "Tango caja A · cajas · x.xlsx", "A"),
            mov("2026-09-30", "Caja A", -999.0, "Tango caja A · cajas · x.xlsx", "A"),       # antes del arqueo
            mov("2026-10-08", "Caja A", -999.0, "Tango caja A · cajas · x.xlsx", "A"),       # después de hoy
            mov("2026-10-03", "Caja A", -999.0, "Manual", "A"),                              # no es de Tango
            mov("2026-10-03", "BANCO X", -999.0, "Extracto BANCO X", "A"),                   # es del banco
        ],
        "cobros_previstos": [
            mov("2026-10-02", "Caja AA", 30.0, "Tango AA · cajas · y.xlsx", "AA"),
            mov("2026-10-04", "Caja A", 5.0, "Tango AA · cajas · y.xlsx", "A"),              # pase desde AA (espejo)
        ],
    }


class Cajas(unittest.TestCase):
    def test_cada_caja_es_su_arqueo_mas_sus_movimientos(self):
        c = CL.actualizar_cajas(contrato_inventado(), HOY)
        self.assertEqual(35.0, c["cajas"]["A"]["saldo"])        # 50 − 20 + 5
        self.assertEqual(130.0, c["cajas"]["AA"]["saldo"])      # 100 + 30
        self.assertEqual(155.0, c["caja_hoy"])                  # 35 + 130 − 10
        self.assertEqual({"A": 25.0, "AA": 130.0}, c["caja_por_unidad"])
        self.assertEqual(130.0, c["caja_aa"]["saldo"])          # lo que lee el tablero viejo

    def test_se_puede_aplicar_dos_veces(self):
        una = CL.actualizar_cajas(contrato_inventado(), HOY)
        dos = CL.actualizar_cajas(una, HOY)
        self.assertEqual((una["caja_hoy"], una["cajas"]["A"]["saldo"]), (dos["caja_hoy"], dos["cajas"]["A"]["saldo"]))
        self.assertEqual(3, len(dos["saldos"]))

    def test_sin_arqueo_de_a_lo_dice(self):
        c = contrato_inventado()
        c["saldos"] = [s for s in c["saldos"] if s["banco"] != "CAJA A"]
        r = CL.actualizar_cajas(c, HOY)
        self.assertEqual("sin arqueo", r["cajas"]["A"]["estado"])
        self.assertEqual("calculada", r["cajas"]["AA"]["estado"])

    def test_cuentas_de_hoy_suman_lo_mismo_que_la_caja(self):
        c = CL.actualizar_cajas(contrato_inventado(), HOY)
        c["saldos"].append({"fecha": "2026-09-16", "banco": "BANCO X", "unidad": "A", "cuenta": "222 333",
                            "saldo": 0.0, "origen": "Extracto BANCO X"})
        cuentas = DATOS.cuentas_de_hoy(c, DATOS.GRUPO)
        self.assertEqual(["Caja A", "Caja AA"], [x["nombre"] for x in cuentas if x["es_caja"]])
        self.assertAlmostEqual(c["caja_hoy"], sum(x["saldo"] for x in cuentas))
        bancos = [x["nombre"] for x in cuentas if not x["es_caja"]]
        self.assertEqual(2, len(set(bancos)))                    # dos cuentas del mismo banco, distinguidas
        self.assertIn("arqueo del 01/10 + 2 movimientos de Tango", [x["nota"] for x in cuentas if x["nombre"] == "Caja A"][0])
        # En la empresa A: primero la caja, después los bancos de mayor a menor saldo.
        self.assertEqual(["Caja A", "Banco X · cta. …2333", "Banco X · cta. …111"],
                         [x["nombre"] for x in DATOS.cuentas_de_hoy(c, "A")])


def _libro(ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    ws.append(["ID", "Fecha", "Empresa", "Tipo", "Categoria", "Concepto / Detalle", "Importe", "Medio de Pago",
               "Banco / Cuenta", "Origen", "Estado"])
    s = wb.create_sheet("Saldos Bancarios")
    s.append(["Fecha", "Banco", "Empresa", "Cuenta / Nro", "Saldo", "Origen"])
    s.append([datetime.datetime(2026, 10, 6), "Caja A", "A", "", 50, "Manual"])
    s.append([datetime.datetime(2026, 10, 9), "Caja A", "A", "", 999, "Manual"])     # arqueo de mañana: no
    b = wb.create_sheet("Deuda Bancaria")
    b.append(["Banco", "Empresa", "Linea / Producto", "Capital Original", "Capital Vigente", "Debito Automatico"])
    b.append(["BANCO X", "A", "Préstamo inventado", 1000, 800, "Si"])
    b.append(["BANCO Y", "A", "Tarjeta inventada", 300, 200, ""])
    b.append([])
    b.append(["B) Cronograma de Vencimientos (cuotas)"])
    b.append(["Banco", "Empresa", "Linea / Producto", "Nro Cuota", "Fecha Vencimiento", "Importe Total Cuota",
              "Estado", "Debito Automatico"])
    b.append(["BANCO X", "A", "Préstamo inventado", 3, datetime.datetime(2026, 10, 20), 100, "Pendiente", "Si"])
    i = wb.create_sheet("Deuda Impositiva")
    i.append(["Impuesto", "Empresa", "Periodo", "Fecha Vencimiento", "Importe", "Estado", "Debito Automatico"])
    i.append(["PLAN INVENTADO", "A", "cuota 1", datetime.datetime(2026, 10, 16), 70, "Pendiente", "Si"])
    i.append(["IMPUESTO INVENTADO", "A", "2026", datetime.datetime(2026, 10, 20), 30, "Pendiente", "No"])
    ch = wb.create_sheet("Cartera de Cheques")
    ch.append(["Tipo", "Empresa", "Nro Cheque", "Beneficiario / Librador", "Banco", "Fecha Pago / Cobro",
               "Importe", "Estado", "Observaciones"])
    ch.append(["Propio", "A", "0001", "PROVEEDOR INVENTADO", "BANCO X", datetime.datetime(2026, 10, 1), 40,
               "En Cartera", "REVISAR: fecha de pago pasada · inventado"])
    ch.append(["Terceros", "A", "0002", "CLIENTE INVENTADO", "BANCO Y", datetime.datetime(2026, 10, 1), 15,
               "En Cartera", "REVISAR: fecha de cobro pasada · inventado"])
    wb.save(ruta)


class Lector(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.ruta = os.path.join(self.tmp.name, "inventada.xlsx")
        _libro(self.ruta)
        self.c = CL.leer(self.ruta, "navar", HOY)

    def tearDown(self):
        self.tmp.cleanup()

    def test_el_titulo_del_cronograma_no_es_un_banco(self):
        self.assertEqual(["BANCO X", "BANCO Y"], [l["banco"] for l in self.c["deuda_bancaria"]["lineas"]])

    def test_lee_el_debito_automatico(self):
        db = self.c["deuda_bancaria"]
        self.assertEqual(["Si", ""], [l["debito_automatico"] for l in db["lineas"]])
        self.assertEqual(["Si"], [x["debito_automatico"] for x in db["cuotas"]])
        self.assertEqual(["Si", "No"], [x["debito_automatico"] for x in self.c["deuda_impositiva"]])
        grupos = {g["id"]: g["total"] for g in DATOS.obligaciones_por_debito(db["lineas"], db["cuotas"], HOY)}
        self.assertEqual({"automatico": 800.0, "decision": 0.0, "sin_definir": 200.0}, grupos)
        imp = {g["id"]: g["en_30"] for g in DATOS.deuda_impositiva_por_debito(self.c, DATOS.GRUPO, HOY)}
        self.assertEqual({"automatico": 70.0, "decision": 30.0, "sin_definir": 0.0}, imp)

    def test_cheque_propio_vencido_va_al_vencido_y_el_de_terceros_a_hallazgos(self):
        # Tarea 54: como en el Cash ("Cheques propios vencidos sin debitar").
        vencidos = [x for x in self.c["egresos_cashflow"] if x["tipo"] == "CHEQUE"]
        self.assertEqual([(40.0, True)], [(x["importe"], x["vencido_pendiente"]) for x in vencidos])
        self.assertEqual(["Cheques de terceros: fecha de cobro pasada"],
                         [h["titulo"] for h in self.c["hallazgos_del_lector"]])
        stock = DATOS.proyeccion(self.c, DATOS.GRUPO, HOY, modo=DATOS.modo_vencido("navar"))["vencido_stock"]
        self.assertEqual(40.0, stock["por_tipo"]["cheques"])

    def test_la_caja_a_arranca_del_arqueo_de_hoy_y_no_del_de_manana(self):
        self.assertEqual(50.0, self.c["cajas"]["A"]["saldo"])
        self.assertEqual("2026-10-06", self.c["cajas"]["A"]["fecha_arqueo"])
        self.assertFalse(any("Caja A" in a for a in self.c["avisos"] if "anterior a hoy" in a))


if __name__ == "__main__":
    unittest.main()
