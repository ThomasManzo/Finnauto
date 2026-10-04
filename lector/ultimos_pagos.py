# -*- coding: utf-8 -*-
"""Último cobro/pago registrado en Tango, por empresa, tipo y razón social.

Lee las dos tesorerías completas antes de escribir: publicar una sola empresa
borraría la otra en la Sheet. No calcula importes ni arma el ranking de deudas.
Uso: python lector/ultimos_pagos.py --a <tesoreria A.xlsx> --aa <tesoreria AA.xlsx>
Deja el Excel y el resumen al lado del archivo de A; el vigilante los publica.
"""
import argparse
import datetime
import math
import os
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl

ENCABEZADOS = ["Empresa", "Tipo", "Codigo", "Razon Social", "Fecha Ultimo Pago",
               "Comprobante", "Origen"]
MARCA = "Tango tesorería"


def _texto(v):
    return "" if v is None else str(v).strip()


def _norm(v):
    # Se normalizan encabezados y etiquetas, nunca la razón social que usa la Sheet.
    return " ".join(unicodedata.normalize("NFD", _texto(v)).encode(
        "ascii", "ignore").decode().lower().split())


def _fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    for formato in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(_texto(v)[:10], formato).date()
        except ValueError:
            pass
    return None


def _numero(v):
    if isinstance(v, str):
        v = v.strip().replace(" ", "")
        if "," in v:
            v = v.replace(".", "").replace(",", ".")
    n = float(v)
    if not math.isfinite(n):
        raise ValueError("importe inválido")
    return n


def _relacionado(codigo, nombre):
    codigo, nombre = _texto(codigo), _texto(nombre)
    if nombre in ("", "-"):
        return codigo, ""
    prefijo = re.match(r"^(\S{3,8}) +-[ ]+(.+)$", nombre)
    # La API pone C/P (el tipo) en Cód. relacionado. El código verdadero está
    # adelante del nombre y puede tener Ñ o &. Un código distinto se conserva.
    es_tipo = len(codigo) == 1 and codigo.isalpha()
    if prefijo and (not codigo or es_tipo or codigo == prefijo[1]):
        codigo = prefijo[1]
        nombre = prefijo[2].strip()
    return codigo, nombre


def _orden_comprobante(numero, tipo):
    # Comparación numérica: 10 gana a 9, también si vinieron sin ceros adelante.
    return tuple(int(n) for n in re.findall(r"\d+", numero)), numero, tipo


def leer(ruta, empresa):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    grupos, ignorados, fechas = {}, Counter(), []
    try:
        filas = wb.active.iter_rows(values_only=True)
        encabezado = [_norm(v) for v in next(filas, ())]
        if len([c for c in encabezado if c]) != len(set(c for c in encabezado if c)):
            raise ValueError("%s: encabezados repetidos" % Path(ruta).name)
        requeridos = {"tipo", "clase", "comprobante", "desc. relacionado"}
        faltan = requeridos - set(encabezado)
        if faltan or not {"fecha", "fecha de emision"}.intersection(encabezado) or not {
                "total (cte)", "total (ext)"}.intersection(encabezado):
            raise ValueError("%s: faltan encabezados de tesorería (tipo, clase, comprobante, "
                             "relacionado, fecha o total)" % Path(ruta).name)
        for valores in filas:
            if not any(v is not None and _texto(v) for v in valores):
                continue
            r = dict(zip(encabezado, valores))
            fecha = _fecha(r.get("fecha") if _texto(r.get("fecha")) else r.get("fecha de emision"))
            if fecha:
                fechas.append(fecha)
            tipo, clase = _texto(r.get("tipo")).upper(), _norm(r.get("clase"))
            motivo = None
            if tipo == "REV":
                motivo = "anulación REV"
            else:
                try:
                    montos = [_numero(r[c]) for c in ("total (cte)", "total (ext)")
                              if _texto(r.get(c))]
                    if any(m < 0 for m in montos):
                        motivo = "importe negativo"
                    elif not montos:
                        motivo = "importe ausente o inválido"
                except (TypeError, ValueError, OverflowError):
                    motivo = "importe ausente o inválido"
            quien = ("Cliente" if clase == "cobros" and tipo in ("REC", "FAC") else
                     "Proveedor" if clase == "pagos" and tipo in ("O/P", "OPF", "FPR") else None)
            codigo, nombre = _relacionado(r.get("cod. relacionado"), r.get("desc. relacionado"))
            if not motivo and not quien:
                motivo = "tipo o clase fuera de alcance"
            if not motivo and not nombre:
                motivo = "sin relacionado"
            nombre_interno = re.sub(r"[^A-Z0-9]", "", nombre.upper())
            if not motivo and empresa == "AA" and (codigo.upper() == "NAV007" or nombre_interno == "NAVARSA"):
                motivo = "pase interno de AA"
            if not motivo and not fecha:
                motivo = "fecha ausente o inválida"
            comprobante = _texto(r.get("comprobante"))
            if not motivo and not comprobante:
                motivo = "sin comprobante"
            if motivo:
                ignorados[motivo] += 1
                continue
            clave = (empresa, quien, nombre)
            orden = (fecha, _orden_comprobante(comprobante, tipo))
            if clave not in grupos:
                grupos[clave] = {"codigos": set(), "orden": orden, "comprobante": tipo + " " + comprobante}
            g = grupos[clave]
            if codigo:
                g["codigos"].add(codigo)
            if orden > g["orden"]:
                g.update(orden=orden, comprobante=tipo + " " + comprobante)
    finally:
        wb.close()
    salida = []
    for (emp, tipo, nombre), g in sorted(grupos.items()):
        salida.append(dict(zip(ENCABEZADOS, [emp, tipo, " / ".join(sorted(g["codigos"])), nombre,
                      g["orden"][0], g["comprobante"], MARCA + " · " + Path(ruta).name])))
    return salida, {"empresa": empresa, "archivo": Path(ruta).name, "ignorados": ignorados,
                    "desde": min(fechas) if fechas else None, "hasta": max(fechas) if fechas else None}


def escribir(filas, carpeta, hoy):
    ruta = Path(carpeta) / ("para_pegar_ultimos_pagos_%s.xlsx" % hoy.isoformat())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ultimos Pagos"
    ws.append(ENCABEZADOS)
    for fila in filas:
        ws.append([fila[c] for c in ENCABEZADOS])
        for c in ws[ws.max_row]:
            if isinstance(c.value, str):
                c.data_type = "s"  # Los nombres y comprobantes son texto, nunca fórmulas.
        ws.cell(ws.max_row, 3).number_format = "@"
        ws.cell(ws.max_row, 5).number_format = "dd/mm/yyyy"
        ws.cell(ws.max_row, 6).number_format = "@"
    temporal = ruta.with_suffix(".parte.xlsx")
    try:
        wb.save(temporal)
        os.replace(temporal, ruta)
    finally:
        wb.close()
        if temporal.exists():
            temporal.unlink()
    return ruta


def resumen(filas, controles, hoy):
    lineas = ["# Últimos cobros y pagos registrados · %s" % hoy.strftime("%d/%m/%Y"), "",
              "Fechas de Tango dentro de los archivos leídos; no confirma acreditación bancaria.", ""]
    for control in controles:
        empresa = control["empresa"]
        conteo = Counter(f["Tipo"] for f in filas if f["Empresa"] == empresa)
        rango = ("%s a %s" % (control["desde"], control["hasta"]) if control["desde"] else "sin fechas válidas")
        lineas += ["## %s · %s" % (empresa, control["archivo"]),
                   "- Clientes: %d · proveedores: %d" % (conteo["Cliente"], conteo["Proveedor"]),
                   "- Rango de fechas del archivo: " + rango]
        lineas += ["- Ignorados (%s): %d" % (motivo, n) for motivo, n in sorted(control["ignorados"].items())]
        if not control["ignorados"]:
            lineas.append("- Ignorados: 0")
        lineas.append("")
    return "\n".join(lineas)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Último cobro/pago de A y AA para la Sheet")
    ap.add_argument("--a", required=True)
    ap.add_argument("--aa", required=True)
    ap.add_argument("--hoy", default=datetime.date.today().isoformat())
    ap.add_argument("--cliente", default="navar")
    a = ap.parse_args(argv)
    hoy = datetime.date.fromisoformat(a.hoy)
    # Comprobar las dos entradas antes de crear cualquier salida.
    for ruta in (a.a, a.aa):
        if not Path(ruta).is_file():
            raise ValueError("falta el archivo de tesorería: %s; no se publica ninguna empresa" % ruta)
    if Path(a.a).resolve() == Path(a.aa).resolve():
        raise ValueError("A y AA necesitan archivos distintos")
    filas, controles = [], []
    for empresa, ruta in (("A", a.a), ("AA", a.aa)):
        nuevas, control = leer(ruta, empresa)
        filas.extend(nuevas)
        controles.append(control)
    carpeta = Path(a.a).resolve().parent
    md = resumen(filas, controles, hoy)
    ruta = escribir(filas, carpeta, hoy)
    (carpeta / ("resumen_ultimos_pagos_%s.md" % hoy.isoformat())).write_text(md + "\n", encoding="utf-8")
    print(md)
    print("- `%s`: solapa Ultimos Pagos" % ruta)
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (OSError, ValueError) as ex:
        sys.exit(str(ex))
