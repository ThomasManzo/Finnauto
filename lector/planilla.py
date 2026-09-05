# -*- coding: utf-8 -*-
"""
lector.planilla — la radiografía de una planilla que uno ve por primera vez.

PARA QUÉ ES (y para qué NO)
---------------------------
Thomas lo planteó así:

    "Yo sé perfectamente que voy a tener que sentarme con quien escribe los
     números para entender qué es cada cosa. Pero me parece importante tener dos
     herramientas que apenas llegue a una empresa las aplique y me resuma un
     poco lo que ve. Lo que quiero vender no es un programa que hace pim pam:
     es una asesoría integral."

Entonces esto NO es un lector automático que entiende cualquier Excel. Es un
instrumento: mira la planilla y devuelve dos cosas.

    1. LO QUE VE      -> hojas, tablas, columnas, tipos de dato, valores
    2. LO QUE NO ENTIENDE -> preguntas concretas para hacerle al cliente

La segunda lista es la más valiosa. Es la diferencia entre entrar a la primera
reunión a explorar y entrar con una agenda: "esta columna tiene 14 valores
distintos, ¿qué son?", "acá hay 200 fechas que no pude leer, ¿cómo las cargan?".

QUÉ BUSCA
---------
No asume una estructura concreta. Busca, en cada hoja, dónde arranca una tabla
(la primera fila que parece encabezado y tiene datos consistentes debajo) y qué
contiene cada columna, mirando los VALORES y no el título. Un título puede decir
"Importe" y tener texto adentro; los valores no mienten.

Uso:
    python lector/planilla.py --archivo "Cash de la empresa.xlsx"
    python lector/planilla.py --archivo x.xlsx --hoja MOVIMIENTOS
"""

import io
import os
import re
import sys
import json
import argparse
import datetime
from collections import Counter, defaultdict

BASE_REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_REPO not in sys.path:
    sys.path.insert(0, BASE_REPO)

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl. Instalalo con:  pip install openpyxl")
    sys.exit(1)


MAX_FILAS = 5000          # con esto alcanza para reconocer la forma
MAX_BUSCAR_ENCABEZADO = 30


# ------------------------------------------------------------------ tipos
def _es_fecha(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return True
    if not isinstance(v, str):
        return False
    return bool(re.match(r"^\s*\d{1,4}[-/]\d{1,2}[-/]\d{1,4}\s*$", v))


def _es_numero(v):
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    if not isinstance(v, str):
        return False
    s = v.strip().replace("$", "").replace(" ", "")
    if not s:
        return False
    # Formato argentino ($1.234.567,89) y tambien el ingles.
    return bool(re.match(r"^-?\(?\d{1,3}([.,]\d{3})*([.,]\d+)?\)?$", s))


def _vacio(v):
    return v is None or (isinstance(v, str) and not v.strip())


def perfil_columna(valores):
    """Qué hay realmente en esta columna, mirando los valores y no el título."""
    llenos = [v for v in valores if not _vacio(v)]
    if not llenos:
        return {"clase": "vacia", "llenos": 0, "distintos": 0, "ejemplos": []}

    n = len(llenos)
    fechas = sum(1 for v in llenos if _es_fecha(v))
    numeros = sum(1 for v in llenos if _es_numero(v))
    distintos = len(set(str(v) for v in llenos))

    if fechas / float(n) > 0.8:
        clase = "fecha"
    elif numeros / float(n) > 0.8:
        clase = "numero"
    elif distintos <= max(30, n * 0.05) and distintos <= n * 0.5:
        # Pocos valores QUE SE REPITEN = es una CATEGORIA (un desplegable).
        # Estas columnas son las mas importantes: de ahi sale el catalogo de
        # tipos del cliente, que es lo que despues maneja todo el motor.
        #
        # El segundo chequeo no sobra: sin el, una columna con 23 valores todos
        # DISTINTOS pasaba como categoria solo por tener menos de 30. Una
        # categoria se define porque los valores se REPITEN, no porque sean pocos.
        clase = "categoria"
    else:
        clase = "texto"

    comunes = Counter(str(v).strip() for v in llenos).most_common(20)
    return {"clase": clase, "llenos": n, "distintos": distintos,
            "cobertura": n / float(len(valores)) if valores else 0,
            "ejemplos": [c[0][:40] for c in comunes[:5]],
            "valores": comunes if clase == "categoria" else [],
            "mixta_fecha": fechas, "mixta_numero": numeros}


# ------------------------------------------------------------------ tablas
def _fila_encabezado(filas):
    """La fila que mejor funciona como encabezado.

    Criterio: la que tiene más celdas de texto no vacías, y debajo tiene datos.
    No se busca ningún nombre de columna en particular: eso ataría el lector a
    una planilla concreta, que es justo lo que no queremos.
    """
    mejor, mejor_p = None, -1
    for i, fila in enumerate(filas[:MAX_BUSCAR_ENCABEZADO]):
        textos = sum(1 for v in fila if isinstance(v, str) and v.strip())
        if textos < 2:
            continue
        abajo = filas[i + 1:i + 12]
        if not abajo:
            continue
        llenos_abajo = sum(1 for f in abajo for v in f if not _vacio(v))
        p = textos * 2 + llenos_abajo * 0.1
        if p > mejor_p:
            mejor, mejor_p = i, p
    return mejor


def leer_hoja(ws):
    filas = []
    for fila in ws.iter_rows(max_row=MAX_FILAS, values_only=True):
        filas.append(list(fila))
    return filas


def analizar_hoja(nombre, filas):
    usadas = [f for f in filas if any(not _vacio(v) for v in f)]
    if not usadas:
        return {"hoja": nombre, "vacia": True}

    h = _fila_encabezado(filas)
    if h is None:
        return {"hoja": nombre, "vacia": False, "sin_tabla": True,
                "filas_con_datos": len(usadas)}

    encabezados = [str(v).strip() if not _vacio(v) else "" for v in filas[h]]
    cuerpo = [f for f in filas[h + 1:] if any(not _vacio(v) for v in f)]

    cols = []
    for c, titulo in enumerate(encabezados):
        valores = [f[c] if c < len(f) else None for f in cuerpo]
        p = perfil_columna(valores)
        p["titulo"] = titulo or "(sin titulo)"
        p["col"] = c
        cols.append(p)

    return {"hoja": nombre, "vacia": False, "fila_encabezado": h + 1,
            "filas_datos": len(cuerpo), "columnas": cols}


# ------------------------------------------------------------------ preguntas
def preguntas(analisis):
    """Lo que el lector NO pudo resolver, en forma de preguntas para el cliente.

    Es la salida más útil del módulo: convierte una primera reunión exploratoria
    en una con agenda.
    """
    qs = []
    for a in analisis:
        if a.get("vacia") or a.get("sin_tabla"):
            continue
        hoja = a["hoja"]
        cols = a["columnas"]

        fechas = [c for c in cols if c["clase"] == "fecha"]
        numeros = [c for c in cols if c["clase"] == "numero"]
        cats = [c for c in cols if c["clase"] == "categoria"]

        if not fechas:
            qs.append((hoja, "No encontre ninguna columna de FECHA. "
                             "¿Como se registra cuando pasa cada movimiento?"))
        if not numeros:
            qs.append((hoja, "No encontre ninguna columna de IMPORTE. "
                             "¿Los montos estan en otra hoja?"))
        if len(numeros) > 1:
            qs.append((hoja, "Hay %d columnas de numeros (%s). ¿Cual es el importe "
                             "del movimiento y que son las otras?" % (
                                 len(numeros), ", ".join(c["titulo"] for c in numeros[:5]))))
        for c in cats:
            vals = ", ".join(v for v, _ in c["valores"][:8])
            if c["distintos"] == 1:
                # Una columna con un solo valor no es un desplegable para
                # discutir: o sobra, o alguien se olvido de completarla.
                qs.append((hoja, "La columna \"%s\" tiene siempre el mismo valor (%s). "
                                 "¿Se usa para algo o quedo de otra epoca?"
                           % (c["titulo"], vals)))
                continue
            qs.append((hoja, "La columna \"%s\" tiene %d valores distintos (%s%s). "
                             "¿Que significa cada uno? ¿Cual se puede postergar y cual no?"
                       % (c["titulo"], c["distintos"], vals,
                          ", ..." if c["distintos"] > 8 else "")))

        # Una columna con nombre y sin un solo dato es una pregunta, no un
        # detalle: o dejo de usarse, o alguien deberia estar completandola.
        for c in cols:
            if c["clase"] == "vacia" and c["titulo"] != "(sin titulo)":
                qs.append((hoja, "La columna \"%s\" existe pero no tiene ni un dato. "
                                 "¿Se dejo de usar o falta cargarla?" % c["titulo"]))
        for c in cols:
            if c["clase"] in ("vacia",):
                continue
            if 0 < c["cobertura"] < 0.5:
                qs.append((hoja, "La columna \"%s\" esta cargada solo en el %.0f%% de "
                                 "las filas. ¿Es opcional o quedo a medio llenar?"
                           % (c["titulo"], 100 * c["cobertura"])))
            # Una columna que mezcla fechas y numeros casi siempre es un error de
            # carga, y es el tipo de cosa que rompe cualquier calculo despues.
            if c["clase"] == "texto" and c["mixta_fecha"] > 3 and c["mixta_numero"] > 3:
                qs.append((hoja, "La columna \"%s\" mezcla fechas, numeros y texto. "
                                 "¿Se carga a mano?" % c["titulo"]))
        if a["filas_datos"] == 0:
            qs.append((hoja, "Encontre encabezados pero ninguna fila de datos. "
                             "¿Esta hoja se usa?"))
    return qs


# ------------------------------------------------------------------ salida
def imprimir(analisis, archivo):
    L = 78
    print("=" * L)
    print("  RADIOGRAFIA DE LA PLANILLA")
    print("  %s" % os.path.basename(archivo))
    print("=" * L)

    utiles = [a for a in analisis if not a.get("vacia")]
    print("  %d hoja(s), %d con datos\n" % (len(analisis), len(utiles)))

    for a in analisis:
        if a.get("vacia"):
            print("  - %-28s (vacia)" % a["hoja"])
            continue
        if a.get("sin_tabla"):
            print("  - %-28s %d filas con datos, pero no reconoci una tabla"
                  % (a["hoja"], a["filas_con_datos"]))
            continue
        print("  " + "-" * (L - 4))
        print("  HOJA: %s" % a["hoja"])
        print("  encabezados en la fila %d  .  %d filas de datos"
              % (a["fila_encabezado"], a["filas_datos"]))
        print("")
        print("    %-26s %-10s %8s %9s  %s" % ("COLUMNA", "CONTIENE", "LLENAS", "DISTINTOS", "EJEMPLOS"))
        for c in a["columnas"]:
            if c["clase"] == "vacia":
                continue
            print("    %-26s %-10s %8d %9d  %s" % (
                c["titulo"][:26], c["clase"], c["llenos"], c["distintos"],
                " | ".join(c["ejemplos"][:2])[:34]))
        vacias = [c for c in a["columnas"] if c["clase"] == "vacia"]
        if vacias:
            print("    (%d columna(s) sin ningun dato)" % len(vacias))
        print("")

    qs = preguntas(analisis)
    print("=" * L)
    print("  PARA PREGUNTARLE AL CLIENTE  (%d)" % len(qs))
    print("=" * L)
    if not qs:
        print("  Nada pendiente: la planilla se entiende sola.")
    else:
        print("  Esto es lo que el lector NO puede deducir mirando. Son las")
        print("  preguntas que conviene llevar a la primera reunion.\n")
        por_hoja = defaultdict(list)
        for hoja, q in qs:
            por_hoja[hoja].append(q)
        for hoja, lista in por_hoja.items():
            print("  --- %s ---" % hoja)
            for q in lista:
                # Envolver a mano: son textos largos y el ancho importa.
                linea = "   . "
                for palabra in q.split():
                    if len(linea) + len(palabra) + 1 > L - 2:
                        print(linea)
                        linea = "     "
                    linea += palabra + " "
                print(linea.rstrip())
            print("")


def analizar(archivo, solo_hoja=None):
    wb = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    out = []
    for nombre in wb.sheetnames:
        if solo_hoja and nombre != solo_hoja:
            continue
        out.append(analizar_hoja(nombre, leer_hoja(wb[nombre])))
    wb.close()
    return out


def main():
    # El cmd de Windows no siempre habla UTF-8. Sin esto, un acento o un signo
    # de pregunta invertido pueden cortar la corrida entera con UnicodeEncodeError
    # justo cuando estas adelante del cliente.
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    ap = argparse.ArgumentParser(description="Radiografia de una planilla desconocida")
    ap.add_argument("--archivo", required=True)
    ap.add_argument("--hoja", help="analizar solo una hoja")
    ap.add_argument("--json", help="guardar el analisis crudo en un archivo")
    args = ap.parse_args()

    if not os.path.exists(args.archivo):
        print("No existe: %s" % args.archivo)
        sys.exit(1)

    analisis = analizar(args.archivo, args.hoja)
    imprimir(analisis, args.archivo)

    if args.json:
        with io.open(args.json, "w", encoding="utf-8") as f:
            f.write(json.dumps({"archivo": args.archivo, "hojas": analisis,
                                "preguntas": preguntas(analisis)},
                               ensure_ascii=False, indent=2, default=str))
        print("  Analisis crudo: %s" % args.json)


if __name__ == "__main__":
    main()
