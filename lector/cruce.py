# -*- coding: utf-8 -*-
"""
lector/cruce — cruce banco ↔ Tango de un mes: qué está en los dos, qué solo en el banco y qué solo
en Tango.

POR QUÉ
    Tango "no se concilia": hay cobranzas que entran al banco y nadie imputa, órdenes de pago
    cargadas que el banco nunca debitó y gastos bancarios que no están en Tango. Este informe
    empareja cada movimiento del extracto con lo que Tango registró en tesorería y deja a la vista
    las diferencias, para revisarlas a mano. NO toca la Sheet ni sube nada a Drive.

LAS DOS PUNTAS
    Banco  = solapa Movimientos de un export de la Sheet (--sheet). Solo las filas cuyo Origen
             empieza con "Extracto" (las de Tango AA, Manual y Proyección no son banco).
    Tango  = export "Detalle de comprobantes" de Tesorería de la empresa A (--tango). Trae un
             renglón por cada cuenta que toca el comprobante (el banco, la caja, los cheques,
             retenciones...). Solo los renglones de una CUENTA DE BANCO con extracto tienen que
             aparecer en el banco: esos son los que se emparejan. En esas cuentas, Debe − Haber es el
             movimiento del banco con el mismo signo que el extracto (Debe = entra, Haber = sale).

LAS REGLAS (en este orden; cada par guarda el nombre de la regla)
    Siempre: misma cuenta (según perfil → cruce → cuentas), mismo signo, diferencia ≤ tolerancia, y
    la fecha del banco entre (fecha Tango − dias_antes) y (fecha Tango + dias_despues): el banco
    acredita unos días después, y a veces Tango se carga tarde y queda después del banco.
    0. internas de la misma cuenta: transferencias entre dos cuentas del extracto que van a la misma
       cuenta de Tango (las dos de BBVA). Se cancelan entre sí y Tango no las registra.
    0b. descuento neto: un crédito del banco = una boleta de depósito de Tango (bruto) + su interés
       (FPR del mismo día, negativo). Así carga Tango los descuentos de cheques de algunos bancos.
       Si no da así, se prueba "corrido": el interés cargado el día del banco y la boleta otro día
       (hasta dias_fecha_corrida).
    1. exacto: uno a uno por importe. Primero los que tienen un solo candidato de cada lado (así el
       resultado no depende del orden); después desempata el mismo CUIT y la fecha más cercana.
    2. mismo CUIT: uno contra varios (en cualquier dirección) del mismo CUIT que suman lo mismo.
    3. bloque del día: depósitos y descuentos de cheques. El banco acredita un renglón por cheque y
       Tango carga una boleta por el total; en Macro la boleta va en bruto y los intereses aparte
       (FPR), y el banco acredita el neto. Se compara el total del día de cada lado.
    4. agrupado: último recurso, uno contra una combinación del mismo día sin CUIT. Solo si la
       combinación es única.
    4b. impuestos agrupados: un débito de impuestos del banco (VEP de ARCA) = varias órdenes de pago
       de Tango sin CUIT de un tercero (el impuesto y sus intereses), aunque estén en días distintos.
    4c. agrupado en dos días: una boleta de depósito = cheques que el banco acreditó en dos días
       seguidos (cada cheque se acredita cuando lo compensa la cámara).
    5. fecha corrida: lo que sobró y tiene un único par por importe exacto a menos de
       dias_fecha_corrida días. Es lo mismo cargado con otra fecha: conciliado, pero marcado.
    Cuando una regla encuentra más de una combinación posible, no empareja: va a "Revisar".

QUÉ TAN SEGURO ES CADA PAR (perfil → cruce → niveles)
    Cada par sale con un nivel y un "Criterio" en castellano (qué coincidió). La idea: el que revisa
    mira solo lo Sugerido y lo Posible, no cada renglón.
    Seguro   = no hay otra lectura posible: exacto con un solo candidato de cada lado o con el mismo
               CUIT, descuento neto (boleta − interés da el neto al centavo), internas.
    Sugerido = lo más probable, pero hubo que elegir: exacto desempatado por fecha, mismo CUIT uno
               contra varios, bloque del día, descuento neto corrido, impuestos agrupados.
    Posible  = cierra por importe pero conviene mirarlo: agrupado (combinación sin CUIT, en uno o dos
               días), fecha corrida, exacto con CUIT distinto de cada lado.

USO
    python lector/cruce.py --cliente navar --mes 2026-08 \\
        --sheet "<export de la Sheet>.xlsx" --tango "<detalle de comprobantes de A>.xlsx" \\
        --salidas clientes/navar/privado/cruce
    Deja cruce_<AAAA-MM>.xlsx y resumen_cruce_<AAAA-MM>.md. Termina con error si la cuenta de
    control no da.
"""

import io
import os
import re
import sys
import glob
import json
import argparse
import datetime
import itertools
import unicodedata
from collections import defaultdict, Counter

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE_REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

MARGEN_CARGA = 10          # días antes y después del mes que se leen, para encontrar pares cruzados
MAX_COMBINAR = {"mismo CUIT": 10, "agrupado": 12}   # candidatos como máximo al buscar combinaciones

# Qué tan seguro es cada forma de emparejar. Se puede cambiar en perfil → cruce → niveles; esto es lo
# que vale si el perfil no dice nada. La clave es la "variante": la regla más el detalle de cómo se
# decidió (p. ej. un "exacto" con un solo candidato no es lo mismo que uno desempatado por fecha).
NIVELES = ("Seguro", "Sugerido", "Posible")
NIVELES_POR_DEFECTO = {
    "exacto único": "Seguro",          # un solo candidato de cada lado
    "exacto con CUIT": "Seguro",       # había varios, ganó el del mismo CUIT
    "descuento neto": "Seguro",        # boleta − interés = neto del banco, al centavo
    "internas": "Seguro",              # entre dos cuentas del extracto que son la misma cuenta de Tango
    "exacto por fecha": "Sugerido",    # había varios del mismo importe, ganó la fecha más cercana
    "mismo CUIT": "Sugerido",          # uno contra varios del mismo CUIT
    "bloque del día": "Sugerido",      # total del día de cada lado
    "descuento neto corrido": "Sugerido",  # boleta − interés = neto, pero cargados en días distintos
    "impuestos agrupados": "Sugerido", # varios pagos de impuestos de Tango = un débito del banco
    "agrupado en dos días": "Posible", # cheques acreditados en dos días seguidos = una boleta
    "agrupado": "Posible",             # combinación sin CUIT
    "fecha corrida": "Posible",        # mismo importe pero fuera de la ventana de días
    "exacto CUIT distinto": "Posible", # mismo importe, pero el CUIT del banco y el de Tango no coinciden
}


# ------------------------------------------------------------------ helpers
def _norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(s.upper().split())


def _m(centavos):
    v = centavos / 100.0
    return ("-" if v < 0 else "") + "$" + format(int(round(abs(v))), ",d").replace(",", ".")


def _centavos(v):
    """Importe (número o texto argentino) → centavos enteros, para comparar sin redondeos."""
    if v in (None, ""):
        return 0
    if isinstance(v, (int, float)):
        return int(round(float(v) * 100))
    s = str(v).strip().replace("$", "").replace(" ", "")
    if "," in s:                                  # 1.234,56 → formato argentino
        s = s.replace(".", "").replace(",", ".")
    try:
        return int(round(float(s) * 100))
    except ValueError:
        return 0


def _fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = str(v or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    return None


def cuit_valido(c):
    """CUIT de 11 dígitos con el dígito verificador correcto."""
    if not c or len(c) != 11 or not c.isdigit():
        return False
    pesos = (5, 4, 3, 2, 7, 6, 5, 4, 3, 2)
    resto = 11 - sum(int(d) * p for d, p in zip(c[:10], pesos)) % 11
    dv = 0 if resto == 11 else resto
    return dv != 10 and dv == int(c[10])


def _cuit_de_tango(v):
    c = re.sub(r"\D", "", str(v or ""))
    return c if cuit_valido(c) else ""


def cuit_del_banco(*textos, propio=""):
    """El primer CUIT válido (que no sea el propio) que aparezca en el concepto o las observaciones.
    Puede venir pegado a otro texto ("TRANSF:ABC123-20111111112") o con guiones ("20-11111111-2");
    no se toman 11 dígitos que sean parte de un número más largo (un CBU, un nro de operación)."""
    for t in textos:
        for m in re.finditer(r"(?<!\d)(\d{2})-?(\d{8})-?(\d)(?!\d)", str(t or "")):
            c = "".join(m.groups())
            if c != propio and cuit_valido(c):
                return c
    return ""


def _cuit_lindo(c):
    return "%s-%s-%s" % (c[:2], c[2:10], c[10:]) if c else ""


def _col(enc, *alternativas):
    """Índice de la primera columna cuyo encabezado coincide (sin tildes ni mayúsculas)."""
    normal = [_norm(e) for e in enc]
    for a in alternativas:
        if _norm(a) in normal:
            return normal.index(_norm(a))
    return None


def cargar_config(cliente):
    ruta = os.path.join(BASE_REPO, "clientes", cliente, "perfil.json")
    with io.open(ruta, encoding="utf-8") as f:
        perfil = json.load(f)
    cfg = perfil.get("cruce")
    if not cfg:
        sys.exit("el perfil de %s no tiene el bloque 'cruce'" % cliente)
    cfg = dict(cfg)
    cfg["cuentas"] = {k: int(v) for k, v in cfg["cuentas"].items() if not k.startswith("_")}
    cfg.setdefault("dias_antes", 3)
    cfg.setdefault("dias_despues", 5)
    cfg.setdefault("tolerancia", 1)
    cfg.setdefault("categorias_gastos", ["Impuestos", "Gastos Bancarios"])
    cfg.setdefault("categorias_deposito", ["Cheques", "Descuento de Cheques"])
    cfg.setdefault("comprobantes_deposito", ["BDM", "BDG", "BNA", "BDC", "BDF"])
    cfg.setdefault("intereses_descuento", {"todas": ["INTERES"], "alguna": ["VTA", "VALORES"]})
    cfg.setdefault("cuentas_efectivo", ["CAJA"])
    cfg.setdefault("cuenta_cheques_terceros", "VALORES A DEPOSITAR")
    cfg.setdefault("comprobantes_cobro", ["REC"])
    cfg.setdefault("comprobantes_pago", ["O/P", "OPF"])
    cfg.setdefault("categorias_impuestos", ["Impuestos"])
    cfg.setdefault("comprobantes_impuesto", ["OPF"])
    cfg.setdefault("impuestos_minimo", 100000)
    cfg["cuit_propio"] = re.sub(r"\D", "", str(cfg.get("cuit_propio") or ""))
    niveles = dict(NIVELES_POR_DEFECTO)
    niveles.update({k: v for k, v in (cfg.get("niveles") or {}).items() if not k.startswith("_")})
    malos = {k: v for k, v in niveles.items() if v not in NIVELES}
    if malos:
        sys.exit("perfil → cruce → niveles: %s (los niveles posibles son %s)" % (malos, ", ".join(NIVELES)))
    cfg["niveles"] = niveles
    return cfg


# ------------------------------------------------------------------ lectura
def leer_banco(ruta_sheet, cfg, desde, hasta):
    """Movimientos de extracto entre desde y hasta. Devuelve (movimientos, ultima_fecha_por_cuenta,
    cuentas_sin_par) — la última fecha se toma de TODA la solapa, para saber hasta dónde llega el
    extracto de cada cuenta."""
    wb = openpyxl.load_workbook(ruta_sheet, read_only=True, data_only=True)
    if "Movimientos" not in wb.sheetnames:
        sys.exit("el export de la Sheet no tiene la solapa Movimientos: %s" % ruta_sheet)
    filas = wb["Movimientos"].iter_rows(values_only=True)
    enc = list(next(filas))
    ix = {k: _col(enc, k) for k in ("ID", "Fecha", "Empresa", "Categoria", "Concepto / Detalle", "Importe",
                                    "Banco / Cuenta", "Origen", "Observaciones")}
    for k in ("Fecha", "Importe", "Banco / Cuenta", "Origen"):
        if ix[k] is None:
            sys.exit("a la solapa Movimientos le falta la columna %r" % k)

    def val(r, k):
        return r[ix[k]] if ix[k] is not None and ix[k] < len(r) else None

    cuentas = {_norm(k): (k, v) for k, v in cfg["cuentas"].items()}
    movs, ultima, sin_par = [], {}, defaultdict(lambda: [0, 0])
    for n, r in enumerate(filas, start=2):
        if not str(val(r, "Origen") or "").startswith("Extracto"):
            continue
        if cfg.get("empresa") and val(r, "Empresa") not in (None, "", cfg["empresa"]):
            continue
        f = _fecha(val(r, "Fecha"))
        c = _centavos(val(r, "Importe"))
        if not f or not c:
            continue
        cta_txt = str(val(r, "Banco / Cuenta") or "").strip()
        if _norm(cta_txt) not in cuentas:
            if desde <= f <= hasta:
                sin_par[cta_txt][0] += 1
                sin_par[cta_txt][1] += c
            continue
        cta, cod = cuentas[_norm(cta_txt)]
        ultima[cod] = max(ultima.get(cod, f), f)
        if not (desde <= f <= hasta):
            continue
        concepto = str(val(r, "Concepto / Detalle") or "")
        obs = str(val(r, "Observaciones") or "")
        movs.append({
            "lado": "Banco", "id": val(r, "ID") or "fila %d" % n, "fecha": f, "cuenta": cta, "cod": cod,
            "c": c, "categoria": str(val(r, "Categoria") or ""), "texto": concepto, "obs": obs,
            "cuit": cuit_del_banco(concepto, obs, propio=cfg["cuit_propio"]),
        })
    wb.close()
    return movs, ultima, dict(sin_par)


def _es_detalle(wb):
    """El export a mano trae la hoja 'Detalle de comprobantes'; el que baja la API (carpeta
    'Tesoreria A detalle') trae una hoja sin nombre útil, pero con 'Cód. cuenta' y el Debe del renglón."""
    if "Detalle de comprobantes" in wb.sheetnames:
        return True
    enc = next(wb.worksheets[0].iter_rows(max_row=1, values_only=True), ())
    return _col(enc, "Cód. cuenta") is not None and _col(enc, "Debe (cte) (renglón)") is not None


def _archivo_tango(ruta):
    """Si --tango es una carpeta, el Excel más nuevo que sea un detalle de comprobantes."""
    if os.path.isfile(ruta):
        return ruta
    candidatos = []
    for p in glob.glob(os.path.join(ruta, "*.xlsx")):
        nombre = os.path.basename(p)
        if nombre.startswith("~$") or nombre.endswith(".parte.xlsx"):
            continue
        try:
            wb = openpyxl.load_workbook(p, read_only=True)
            if _es_detalle(wb):
                candidatos.append(p)
            wb.close()
        except Exception:
            continue
    if not candidatos:
        sys.exit("no encontré un 'Detalle de comprobantes' de Tango en %s" % ruta)
    return max(candidatos, key=os.path.getmtime)


def leer_tango(ruta, cfg):
    """Todos los renglones del detalle de comprobantes (con fecha), con el importe firmado Debe − Haber.
    No filtra por fechas: el filtro lo hace quien llama, porque las fechas imposibles hay que verlas."""
    ruta = _archivo_tango(ruta)
    # read_only=False: el export de Live declara mal el tamaño de la hoja y en modo read_only
    # openpyxl lee una sola celda.
    wb = openpyxl.load_workbook(ruta)
    ws = wb["Detalle de comprobantes"] if "Detalle de comprobantes" in wb.sheetnames else wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    enc = list(next(filas))
    alt = {
        "fecha": ("Fecha de emisión", "Fecha emisión", "Fecha"),
        "tipo": ("Cód. comprobante", "Tipo", "Tipo de comprobante"),
        "desc_tipo": ("Desc. comprobante",),
        "comprobante": ("Comprobante", "Nro. comprobante", "Número"),
        "interno": ("Nro. interno", "Número interno"),
        "cod": ("Cód. cuenta", "Código cuenta"),
        "desc_cuenta": ("Desc. contable", "Desc. cuenta", "Cuenta"),
        "banco": ("Banco",),
        "debe": ("Debe (cte) (renglón)", "Debe (cte)", "Debe"),
        "haber": ("Haber (cte) (renglón)", "Haber (cte)", "Haber"),
        "total": ("Total comp. (cte)", "Total (cte)"),
        "cuit_cli": ("CUIT cliente (encab.)", "CUIT cliente"),
        "cuit_prov": ("CUIT proveedor (encab.)", "CUIT proveedor"),
        "razon": ("Razón social (encab.)", "Razón social"),
        "proveedor": ("Proveedor (encab.)", "Proveedor"),
        "leyenda": ("Leyenda",),
    }
    ix = {k: _col(enc, *v) for k, v in alt.items()}
    for k in ("fecha", "tipo", "cod", "debe", "haber"):
        if ix[k] is None:
            sys.exit("al detalle de Tango le falta la columna de %r (%s)" % (k, alt[k][0]))

    def val(r, k):
        return r[ix[k]] if ix[k] is not None and ix[k] < len(r) else None

    renglones = []
    for n, r in enumerate(filas, start=2):
        f = _fecha(val(r, "fecha"))
        if not f:
            continue                            # la fila casi vacía del final
        cod = val(r, "cod")
        try:
            cod = int(cod) if cod not in (None, "") else None
        except (TypeError, ValueError):
            cod = None
        cuit = _cuit_de_tango(val(r, "cuit_cli")) or _cuit_de_tango(val(r, "cuit_prov"))
        renglones.append({
            "lado": "Tango", "fila": n, "fecha": f, "cod": cod,
            "tipo": str(val(r, "tipo") or "").strip(),
            "comprobante": " ".join(str(val(r, "comprobante") or "").split()),
            "interno": val(r, "interno"),
            "desc_cuenta": str(val(r, "desc_cuenta") or "").strip(),
            "banco": str(val(r, "banco") or "").strip(),
            "c": _centavos(val(r, "debe")) - _centavos(val(r, "haber")),
            "total": _centavos(val(r, "total")),
            "cuit": cuit,
            "contraparte": str(val(r, "razon") or val(r, "proveedor") or "").strip(),
            "texto": str(val(r, "leyenda") or "").strip(),
        })
    wb.close()
    return renglones, ruta


# ------------------------------------------------------------------ emparejado
class Cruce:
    """Guarda qué renglón ya se usó y los pares/grupos armados. Trabaja con listas de diccionarios:
    no sabe nada de archivos."""

    def __init__(self, banco, tango, cfg):
        self.banco, self.tango, self.cfg = banco, tango, cfg
        self.tol = int(round(float(cfg["tolerancia"]) * 100))
        self.antes, self.despues = int(cfg["dias_antes"]), int(cfg["dias_despues"])
        self.usados = set()             # id() de los renglones ya emparejados
        self.pares = []                 # {"regla", "banco": [...], "tango": [...]}
        self.internas = []              # pares de transferencias entre cuentas de la misma cuenta Tango
        self.ambiguos = {}              # id() → texto de por qué quedó en Revisar
        for i, b in enumerate(banco):
            b["n"] = i
        for i, t in enumerate(tango):
            t["n"] = i

    # --- utilidades
    def libre(self, x):
        return id(x) not in self.usados

    def en_ventana(self, fecha_banco, fecha_tango):
        return -self.antes <= (fecha_banco - fecha_tango).days <= self.despues

    def iguales(self, a, b):
        return abs(a - b) <= self.tol

    def emparejar(self, regla, bs, ts, variante=None, nota=""):
        """variante = la clave de perfil → cruce → niveles (si no se dice, es la regla); nota = lo que
        haya que contar en el Criterio además de lo que se ve en los datos (p. ej. un desempate)."""
        for x in list(bs) + list(ts):
            self.usados.add(id(x))
            self.ambiguos.pop(id(x), None)
        self.pares.append({"regla": regla, "variante": variante or regla, "nota": nota,
                           "banco": list(bs), "tango": list(ts)})

    def es_gasto(self, b):
        return b["categoria"] in self.cfg["categorias_gastos"]

    def es_interes_descuento(self, t):
        ley = _norm(t["texto"])
        reg = self.cfg["intereses_descuento"]
        return (t["tipo"] == "FPR" and t["c"] < 0 and all(p in ley for p in reg.get("todas", []))
                and any(p in ley for p in reg.get("alguna", [""])))

    def combinaciones(self, candidatos, objetivo, maximo):
        """Combinaciones de 2 o más candidatos que suman el objetivo. Corta en 2: alcanza para saber
        si es única. None si hay demasiados candidatos para probar."""
        if len(candidatos) < 2:
            return []
        if len(candidatos) > maximo:
            return None
        encontradas = []
        for k in range(2, len(candidatos) + 1):
            for combo in itertools.combinations(candidatos, k):
                if self.iguales(sum(x["c"] for x in combo), objetivo):
                    encontradas.append(combo)
                    if len(encontradas) > 1:
                        return encontradas
        return encontradas

    # --- regla 0: internas de la misma cuenta de Tango
    def internas_misma_cuenta(self):
        por_cod = defaultdict(set)
        for b in self.banco:
            por_cod[b["cod"]].add(b["cuenta"])
        compartidas = {cod for cod, ctas in por_cod.items() if len(ctas) > 1}
        salidas = [b for b in self.banco if b["cod"] in compartidas and b["c"] < 0]
        for s in sorted(salidas, key=lambda b: (b["fecha"], b["n"])):
            if not self.libre(s):
                continue
            cands = [e for e in self.banco if self.libre(e) and e["cod"] == s["cod"] and e["cuenta"] != s["cuenta"]
                     and e["c"] > 0 and self.iguales(e["c"], -s["c"]) and abs((e["fecha"] - s["fecha"]).days) <= 1]
            if cands:
                e = min(cands, key=lambda e: (abs((e["fecha"] - s["fecha"]).days), e["n"]))
                self.usados.update({id(s), id(e)})
                self.internas.append([s, e])

    # --- regla 0b: descuento neto (una boleta bruta menos su interés = el crédito del banco)
    def descuento_neto(self):
        """En los descuentos de cheques (Macro, Nación) Tango carga la boleta por el BRUTO y el interés
        aparte (un FPR del mismo día), y el banco acredita el NETO en un solo renglón. Va antes que
        'exacto' porque si no, una boleta suelta que casualmente iguala otro crédito se lo roba.
        Solo se empareja si hay una única combinación boleta + interés que da."""
        deps = set(self.cfg["comprobantes_deposito"])
        lejos = int(self.cfg.get("dias_fecha_corrida", 15))
        # Pasada 1: boleta e interés del mismo día, dentro de la ventana normal.
        # Pasada 2 ("corrido", con lo que sobró): el interés cargado cerca del día del banco y la
        # boleta en otro día (hasta dias_fecha_corrida). Pasa cuando la boleta se carga días después.
        for corrido in (False, True):
            for b in sorted(self.banco, key=lambda b: (b["fecha"], b["n"])):
                if not self.libre(b) or b["c"] <= 0 or self.es_gasto(b):
                    continue
                boletas = [t for t in self.tango if self.libre(t) and t["cod"] == b["cod"] and t["tipo"] in deps
                           and t["c"] > b["c"] and (abs((b["fecha"] - t["fecha"]).days) <= lejos if corrido
                                                    else self.en_ventana(b["fecha"], t["fecha"]))]
                if not boletas:
                    continue
                intereses = [t for t in self.tango if self.libre(t) and t["cod"] == b["cod"] and self.es_interes_descuento(t)]
                if corrido:
                    combos = [(d, i) for d in boletas for i in intereses if i["fecha"] != d["fecha"]
                              and self.en_ventana(b["fecha"], i["fecha"]) and self.iguales(d["c"] + i["c"], b["c"])]
                else:
                    combos = [(d, i) for d in boletas for i in intereses
                              if i["fecha"] == d["fecha"] and self.iguales(d["c"] + i["c"], b["c"])]
                if len(combos) != 1:
                    continue
                d, i = combos[0]
                if corrido:
                    self.emparejar("descuento neto", [b], [d, i], "descuento neto corrido",
                                   "boleta del %s e interés del %s: cargados en días distintos"
                                   % (d["fecha"].strftime("%d/%m"), i["fecha"].strftime("%d/%m")))
                else:
                    self.emparejar("descuento neto", [b], [d, i])

    # --- regla 1: exacto
    def candidatos_exacto(self, b):
        return [t for t in self.tango if self.libre(t) and t["cod"] == b["cod"] and self.iguales(t["c"], b["c"])
                and self.en_ventana(b["fecha"], t["fecha"])]

    def exacto(self):
        banco = [b for b in self.banco if self.libre(b)]
        # 1a. primero los pares sin competencia de ningún lado, hasta que no aparezcan más
        cambio = True
        while cambio:
            cambio = False
            for b in banco:
                if not self.libre(b):
                    continue
                cands = self.candidatos_exacto(b)
                if len(cands) != 1:
                    continue
                t = cands[0]
                rivales = [o for o in banco if o is not b and self.libre(o) and o["cod"] == t["cod"]
                           and self.iguales(o["c"], t["c"]) and self.en_ventana(o["fecha"], t["fecha"])]
                if not rivales:
                    self.emparejar("exacto", [b], [t], self._variante_exacto(b, t, "exacto único"))
                    cambio = True
        # 1b. el resto: gana el mismo CUIT, después la fecha más cercana
        for b in sorted(banco, key=lambda b: (b["fecha"], b["n"])):
            if not self.libre(b):
                continue
            cands = self.candidatos_exacto(b)
            if cands:
                t = min(cands, key=lambda t: (0 if b["cuit"] and t["cuit"] == b["cuit"] else 1,
                                              abs((b["fecha"] - t["fecha"]).days), t["fecha"], t["n"]))
                # cuántos competían: candidatos de Tango para este del banco + otros del banco para este de Tango
                rivales = [o for o in banco if o is not b and self.libre(o) and o["cod"] == t["cod"]
                           and self.iguales(o["c"], t["c"]) and self.en_ventana(o["fecha"], t["fecha"])]
                # Acá nada es "único" de verdad: los únicos ya salieron en 1a. Si quedó un solo candidato
                # es porque otro desempate del mismo importe se llevó al resto: depende de esa elección.
                if b["cuit"] and t["cuit"] == b["cuit"]:
                    self.emparejar("exacto", [b], [t], "exacto con CUIT",
                                   "desempate: %d candidatos del mismo importe, ganó el del mismo CUIT"
                                   % (len(cands) + len(rivales)))
                elif len(cands) == 1 and not rivales:
                    self.emparejar("exacto", [b], [t], self._variante_exacto(b, t, "exacto por fecha"),
                                   "quedó solo después de otro desempate del mismo importe")
                else:
                    self.emparejar("exacto", [b], [t], self._variante_exacto(b, t, "exacto por fecha"),
                                   "desempate: %d candidatos del mismo importe, ganó la fecha más cercana"
                                   % (len(cands) + len(rivales)))

    @staticmethod
    def _variante_exacto(b, t, si_no):
        """Un exacto donde los dos lados traen CUIT y no coinciden baja a 'Posible', sea como sea."""
        if b["cuit"] and t["cuit"] and b["cuit"] != t["cuit"]:
            return "exacto CUIT distinto"
        return si_no

    # --- regla 2: mismo CUIT, uno contra varios
    def mismo_cuit(self):
        maximo = MAX_COMBINAR["mismo CUIT"]
        for b in sorted(self.banco, key=lambda b: (b["fecha"], b["n"])):
            if not self.libre(b) or not b["cuit"]:
                continue
            cands = [t for t in self.tango if self.libre(t) and t["cod"] == b["cod"] and t["cuit"] == b["cuit"]
                     and (t["c"] > 0) == (b["c"] > 0) and self.en_ventana(b["fecha"], t["fecha"])]
            self._resolver("mismo CUIT", [b], cands, b["c"], maximo, lado_uno="banco")
        for t in sorted(self.tango, key=lambda t: (t["fecha"], t["n"])):
            if not self.libre(t) or not t["cuit"] or t["cod"] is None:
                continue
            cands = [b for b in self.banco if self.libre(b) and b["cod"] == t["cod"] and b["cuit"] == t["cuit"]
                     and (b["c"] > 0) == (t["c"] > 0) and self.en_ventana(b["fecha"], t["fecha"])]
            self._resolver("mismo CUIT", [t], cands, t["c"], maximo, lado_uno="tango")

    def _resolver(self, regla, uno, cands, objetivo, maximo, lado_uno):
        combos = self.combinaciones(cands, objetivo, maximo)
        if not combos:
            return False
        if len(combos) > 1:
            x = uno[0]
            self.ambiguos.setdefault(id(x), "%s: más de una combinación posible entre %d candidatos" % (regla, len(cands)))
            return False
        otros = list(combos[0])
        if lado_uno == "banco":
            self.emparejar(regla, uno, otros)
        else:
            self.emparejar(regla, otros, uno)
        return True

    # --- regla 3: bloque del día (depósitos y descuentos de cheques)
    def bloque_del_dia(self):
        cats = set(self.cfg["categorias_deposito"])
        deps = set(self.cfg["comprobantes_deposito"])
        dias = sorted({(b["cod"], b["fecha"]) for b in self.banco
                       if self.libre(b) and b["c"] > 0 and b["categoria"] in cats})
        for cod, d in dias:
            bs = [b for b in self.banco if self.libre(b) and b["cod"] == cod and b["fecha"] == d
                  and b["c"] > 0 and b["categoria"] in cats]
            if not bs:
                continue
            total_banco = sum(b["c"] for b in bs)
            # días posibles de Tango, del más cercano al más lejano (a igual distancia, Tango antes)
            offs = sorted(range(-self.antes, self.despues + 1), key=lambda k: (abs(k), -k))
            for k in offs:
                dia_t = d - datetime.timedelta(days=k)
                ts = [t for t in self.tango if self.libre(t) and t["cod"] == cod and t["fecha"] == dia_t
                      and ((t["tipo"] in deps and t["c"] > 0) or self.es_interes_descuento(t))]
                if not any(t["tipo"] in deps for t in ts):
                    continue
                if self.iguales(sum(t["c"] for t in ts), total_banco):
                    self.emparejar("bloque del día", bs, ts)
                    break
                # sin los intereses (bancos que acreditan en bruto)
                solo_dep = [t for t in ts if t["tipo"] in deps]
                if len(solo_dep) < len(ts) and self.iguales(sum(t["c"] for t in solo_dep), total_banco):
                    self.emparejar("bloque del día", bs, solo_dep)
                    break

    # --- regla 4: agrupado (combinación del mismo día, sin CUIT)
    def agrupado(self):
        maximo = MAX_COMBINAR["agrupado"]
        for t in sorted(self.tango, key=lambda t: (t["fecha"], t["n"])):
            if not self.libre(t) or t["cod"] is None:
                continue
            dias = sorted({b["fecha"] for b in self.banco if self.libre(b) and b["cod"] == t["cod"]
                           and self.en_ventana(b["fecha"], t["fecha"])}, key=lambda d: (abs((d - t["fecha"]).days), d))
            for d in dias:
                cands = [b for b in self.banco if self.libre(b) and b["cod"] == t["cod"] and b["fecha"] == d
                         and (b["c"] > 0) == (t["c"] > 0) and not self.es_gasto(b)]
                if self._resolver("agrupado", [t], cands, t["c"], maximo, lado_uno="tango"):
                    break
        for b in sorted(self.banco, key=lambda b: (b["fecha"], b["n"])):
            if not self.libre(b) or self.es_gasto(b):
                continue
            dias = sorted({t["fecha"] for t in self.tango if self.libre(t) and t["cod"] == b["cod"]
                           and self.en_ventana(b["fecha"], t["fecha"])}, key=lambda d: (abs((b["fecha"] - d).days), d))
            for d in dias:
                cands = [t for t in self.tango if self.libre(t) and t["cod"] == b["cod"] and t["fecha"] == d
                         and (t["c"] > 0) == (b["c"] > 0)]
                if self._resolver("agrupado", [b], cands, b["c"], maximo, lado_uno="banco"):
                    break

    # --- regla 4b: impuestos agrupados (un VEP del banco = varias órdenes de pago de Tango)
    def impuestos_agrupados(self):
        """ARCA debita en un solo renglón lo que Tango carga en varias órdenes de pago: el impuesto y
        sus intereses, a veces en días distintos. Se busca, entre las órdenes de pago de impuestos sin
        CUIT de un tercero y dentro de la ventana de días, la única combinación que da el débito."""
        cats = set(self.cfg.get("categorias_impuestos", ["Impuestos"]))
        tipos = set(self.cfg.get("comprobantes_impuesto", ["OPF"]))
        minimo = int(round(float(self.cfg.get("impuestos_minimo", 100000)) * 100))
        for b in sorted(self.banco, key=lambda b: (b["fecha"], b["n"])):
            if not self.libre(b) or b["c"] >= 0 or -b["c"] < minimo or b["categoria"] not in cats:
                continue
            cands = [t for t in self.tango if self.libre(t) and t["cod"] == b["cod"] and t["c"] < 0
                     and t["tipo"] in tipos and not t["cuit"] and self.en_ventana(b["fecha"], t["fecha"])]
            combos = self.combinaciones(cands, b["c"], MAX_COMBINAR["agrupado"])
            if combos is None or not combos:
                continue
            if len(combos) > 1:
                self.ambiguos.setdefault(id(b), "impuestos agrupados: más de una combinación posible entre %d"
                                                " órdenes de pago" % len(cands))
                continue
            self.emparejar("impuestos", [b], list(combos[0]), "impuestos agrupados")

    # --- regla 4c: agrupado en dos días (una boleta = cheques acreditados en dos días seguidos)
    def agrupado_dos_dias(self):
        """Los cheques de una misma boleta se acreditan cuando los compensa la cámara: a veces unos un
        día y otros al siguiente. Se prueba, por cada boleta sin par, cada par de días seguidos con
        créditos de cheques en la ventana; la combinación tiene que usar los dos días y ser única."""
        cats = set(self.cfg["categorias_deposito"])
        deps = set(self.cfg["comprobantes_deposito"])
        maximo = MAX_COMBINAR["agrupado"]
        for t in sorted(self.tango, key=lambda t: (t["fecha"], t["n"])):
            if not self.libre(t) or t["tipo"] not in deps or t["c"] <= 0:
                continue
            dias = sorted({b["fecha"] for b in self.banco if self.libre(b) and b["cod"] == t["cod"] and b["c"] > 0
                           and b["categoria"] in cats and self.en_ventana(b["fecha"], t["fecha"])})
            encontradas = []
            for d1, d2 in zip(dias, dias[1:]):
                cands = [b for b in self.banco if self.libre(b) and b["cod"] == t["cod"] and b["c"] > 0
                         and b["categoria"] in cats and b["fecha"] in (d1, d2)]
                if len(cands) > maximo:
                    continue
                for k in range(2, len(cands) + 1):
                    for combo in itertools.combinations(cands, k):
                        if ({b["fecha"] for b in combo} == {d1, d2}
                                and self.iguales(sum(b["c"] for b in combo), t["c"])):
                            encontradas.append((combo, d1, d2))
                            if len(encontradas) > 1:
                                break
                    if len(encontradas) > 1:
                        break
                if len(encontradas) > 1:
                    break
            if len(encontradas) == 1:
                combo, d1, d2 = encontradas[0]
                self.emparejar("agrupado", list(combo), [t], "agrupado en dos días",
                               "cheques acreditados el %s y el %s" % (d1.strftime("%d/%m"), d2.strftime("%d/%m")))
            elif len(encontradas) > 1:
                self.ambiguos.setdefault(id(t), "agrupado en dos días: más de una combinación posible")

    # --- regla 5: fecha corrida (mismo importe, cargado en Tango bastante antes o después)
    def fecha_corrida(self):
        """Lo que sobró y tiene en la misma cuenta un único par por importe exacto dentro de
        dias_fecha_corrida (y ese par no tiene otro candidato). Es lo mismo, pero Tango lo cargó con
        otra fecha (p. ej. débitos de fin de mes que Tango carga el 1ro del mes siguiente). Cuenta como conciliado
        y queda marcado, porque corre la foto del día en el cash."""
        dias = int(self.cfg.get("dias_fecha_corrida", 15))
        cerca = lambda b, t: abs((b["fecha"] - t["fecha"]).days) <= dias
        for b in sorted(self.banco, key=lambda b: (b["fecha"], b["n"])):
            if not self.libre(b):
                continue
            cands = [t for t in self.tango if self.libre(t) and t["cod"] == b["cod"]
                     and self.iguales(t["c"], b["c"]) and cerca(b, t)]
            if len(cands) != 1:
                continue
            t = cands[0]
            rivales = [o for o in self.banco if o is not b and self.libre(o) and o["cod"] == t["cod"]
                       and self.iguales(o["c"], t["c"]) and cerca(o, t)]
            if not rivales:
                self.emparejar("fecha corrida", [b], [t])

    def correr(self):
        self.internas_misma_cuenta()
        self.descuento_neto()
        self.exacto()
        self.mismo_cuit()
        self.bloque_del_dia()
        self.agrupado()
        self.impuestos_agrupados()
        self.agrupado_dos_dias()
        self.fecha_corrida()
        return self


# ------------------------------------------------------------------ nivel y criterio de cada par
def _dias_txt(d):
    if d == 0:
        return "mismo día"
    return "banco %d día%s %s que Tango" % (abs(d), "" if abs(d) == 1 else "s", "después" if d > 0 else "antes")


def criterio(par):
    """Qué coincidió, en castellano simple: importe, CUIT, días entre banco y Tango, cuántos renglones
    de cada lado y, si hubo, el desempate. Es lo que lee el que revisa para decidir si confía."""
    bs, ts = par["banco"], par["tango"]
    partes = ["%d banco ↔ %d Tango" % (len(bs), len(ts))] if ts else ["%d ↔ %d entre cuentas del banco" % (1, 1)]
    if ts:
        dif = sum(b["c"] for b in bs) - sum(t["c"] for t in ts)
        partes.append("importe exacto" if dif == 0 else "diferencia de %s" % _m(dif))
        cb = {b["cuit"] for b in bs if b["cuit"]}
        ct = {t["cuit"] for t in ts if t["cuit"]}
        if cb and ct:
            partes.append("CUIT coincide" if cb & ct else "CUIT distinto")
        elif cb or ct:
            partes.append("CUIT solo en el " + ("banco" if cb else "Tango") + ": no se pudo comparar")
        else:
            partes.append("sin CUIT")
        partes.append(_dias_txt((min(b["fecha"] for b in bs) - min(t["fecha"] for t in ts)).days))
    else:
        partes += ["importe exacto con signo contrario", _dias_txt((bs[1]["fecha"] - bs[0]["fecha"]).days)
                   .replace("banco", "una").replace("que Tango", "que la otra")]
    extra = {
        "descuento neto corrido": "boleta de Tango menos su interés = lo que acreditó el banco",
        "impuestos agrupados": "varias órdenes de pago de impuestos de Tango (impuesto e intereses) suman el débito",
        "agrupado en dos días": "combinación única de cheques de dos días seguidos, sin CUIT que lo confirme",
    }.get(par.get("variante")) or {
        "descuento neto": "boleta de Tango menos su interés = lo que acreditó el banco",
        "bloque del día": "el total del día del banco = el total de boletas de Tango",
        "agrupado": "combinación única del mismo día, sin CUIT que lo confirme",
        "fecha corrida": "único par con ese importe, pero fuera de la ventana normal de días",
        "mismo CUIT": "varios renglones del mismo CUIT que suman lo mismo",
    }.get(par["regla"])
    if extra:
        partes.append(extra)
    if par.get("nota"):
        partes.append(par["nota"])
    return " · ".join(partes)


# ------------------------------------------------------------------ clasificación
def fechas_imposibles(renglones, hoy):
    """Renglones con una fecha que no puede ser (más de 60 días en el futuro, o antes del 2000)."""
    tope = hoy + datetime.timedelta(days=60)
    return [t for t in renglones if t["fecha"] > tope or t["fecha"].year < 2000]


def candidato_probable(b, tango_banco, cfg, contraparte_por_cuit, imposibles=()):
    """Una pista para encontrar en Tango un movimiento del banco que quedó sin par."""
    tol = int(round(float(cfg["tolerancia"]) * 100))
    for t in imposibles:
        if t["cod"] == b["cod"] and abs(t["c"] - b["c"]) <= tol:
            return "mismo importe en Tango con fecha imposible %s (%s %s)" % (
                t["fecha"].strftime("%d/%m/%Y"), t["tipo"], t["comprobante"])
    mismo = [t for t in tango_banco if abs(t["c"] - b["c"]) <= tol and abs((t["fecha"] - b["fecha"]).days) <= 45]
    misma_cuenta = [t for t in mismo if t["cod"] == b["cod"]]
    if misma_cuenta:
        t = min(misma_cuenta, key=lambda t: abs((t["fecha"] - b["fecha"]).days))
        return "mismo importe el %s (%s %s), fuera de la ventana de fechas" % (t["fecha"].strftime("%d/%m"), t["tipo"], t["comprobante"])
    otra = [t for t in mismo if t["cod"] != b["cod"] and abs((t["fecha"] - b["fecha"]).days) <= 10]
    if otra:
        t = min(otra, key=lambda t: abs((t["fecha"] - b["fecha"]).days))
        return "mismo importe el %s en otra cuenta de Tango (%s): ¿imputado al banco equivocado?" % (
            t["fecha"].strftime("%d/%m"), t["desc_cuenta"] or t["cod"])
    # Mismos dígitos en otro orden (12.345.678 contra 12.354.678): típico error de tipeo al cargar.
    pesos = lambda c: str(abs(c) // 100)
    tipeo = [t for t in tango_banco if t["cod"] == b["cod"] and (t["c"] > 0) == (b["c"] > 0)
             and abs(t["c"]) >= 100000 * 100 and pesos(t["c"]) != pesos(b["c"])
             and sorted(pesos(t["c"])) == sorted(pesos(b["c"])) and abs((t["fecha"] - b["fecha"]).days) <= 45]
    if tipeo:
        t = min(tipeo, key=lambda t: abs((t["fecha"] - b["fecha"]).days))
        return "¿error de tipeo? Tango tiene %s el %s (%s %s): mismos dígitos en otro orden" % (
            _m(t["c"]), t["fecha"].strftime("%d/%m"), t["tipo"], t["comprobante"])
    if b["cuit"]:
        cuit = [t for t in tango_banco if t["cuit"] == b["cuit"] and abs((t["fecha"] - b["fecha"]).days) <= 15]
        if cuit:
            t = min(cuit, key=lambda t: abs((t["fecha"] - b["fecha"]).days))
            return "mismo CUIT el %s por %s (%s %s)" % (t["fecha"].strftime("%d/%m"), _m(t["c"]), t["tipo"], t["comprobante"])
        if b["cuit"] not in contraparte_por_cuit:
            return "CUIT que Tango no conoce"
    return ""


def armar_informe(banco, tango, cfg, mes_desde, mes_hasta, ultima, hoy, sin_par_banco):
    """Corre las reglas y arma todas las listas del informe. Devuelve un diccionario."""
    cods = set(cfg["cuentas"].values())
    carga_desde = mes_desde - datetime.timedelta(days=MARGEN_CARGA)
    carga_hasta = mes_hasta + datetime.timedelta(days=MARGEN_CARGA)
    en_mes = lambda f: mes_desde <= f <= mes_hasta

    imposibles = [t for t in fechas_imposibles(tango, hoy) if t["cod"] in cods]
    ids_imposibles = {id(t) for t in imposibles}
    tango_banco = [t for t in tango if t["cod"] in cods and id(t) not in ids_imposibles
                   and carga_desde <= t["fecha"] <= carga_hasta]
    revisar = []
    for t in imposibles:
        corregida = ""
        try:
            f2 = t["fecha"].replace(year=mes_desde.year)
            if carga_desde <= f2 <= carga_hasta:
                corregida = "¿quiso decir %s?" % f2.strftime("%d/%m/%Y")
        except ValueError:
            pass
        revisar.append(dict(t, motivo="fecha imposible %s %s" % (t["fecha"].strftime("%d/%m/%Y"), corregida)))
    ceros = [t for t in tango_banco if t["c"] == 0]
    tango_banco = [t for t in tango_banco if t["c"] != 0]
    for t in ceros:
        if en_mes(t["fecha"]):
            revisar.append(dict(t, motivo="renglón de banco con importe cero"))
    for t in tango:
        if t["cod"] is None and t["c"] != 0 and en_mes(t["fecha"]):
            revisar.append(dict(t, motivo="renglón sin cuenta"))

    cruce = Cruce(banco, tango_banco, cfg).correr()

    # pares que tocan el mes (por cualquiera de las dos fechas)
    pares = [p for p in cruce.pares if any(en_mes(x["fecha"]) for x in p["banco"] + p["tango"])]
    pares.sort(key=lambda p: min(x["fecha"] for x in p["banco"] + p["tango"]))
    internas = [p for p in cruce.internas if any(en_mes(x["fecha"]) for x in p)]
    niveles = dict(NIVELES_POR_DEFECTO, **(cfg.get("niveles") or {}))
    for par in pares:
        par["nivel"] = niveles.get(par["variante"], "Posible")
        par["criterio"] = criterio(par)
    # las internas también se muestran como pares (sin lado Tango), así entran en la vista por nivel
    pares_internas = [{"regla": "internas", "variante": "internas", "nota": "", "banco": list(p), "tango": [],
                       "nivel": niveles.get("internas", "Seguro")} for p in internas]
    for par in pares_internas:
        par["criterio"] = criterio(par)

    contraparte_por_cuit = {}
    for t in tango:
        if t["cuit"] and t["contraparte"]:
            contraparte_por_cuit.setdefault(t["cuit"], t["contraparte"])

    libres_banco = [b for b in banco if cruce.libre(b) and en_mes(b["fecha"])]
    # para las pistas: lo de Tango que quedó sin par (en todas las cuentas de banco y sin límite de
    # ventana), así "mismo importe en otra fecha" apunta a algo que de verdad está libre
    libres_tango = [t for t in tango if t["cod"] in cods and t["c"] and id(t) not in ids_imposibles
                    and cruce.libre(t)]
    solo_banco, gastos = [], []
    for b in libres_banco:
        if id(b) in cruce.ambiguos:
            revisar.append(dict(b, motivo=cruce.ambiguos[id(b)]))
        elif cruce.es_gasto(b):
            gastos.append(b)
        else:
            solo_banco.append(dict(b, contraparte=contraparte_por_cuit.get(b["cuit"], ""),
                                   pista=candidato_probable(b, libres_tango, cfg, contraparte_por_cuit, imposibles)))

    # Anulados dentro de Tango: un comprobante y su reversión (REV) por el mismo importe con signo
    # contrario se cancelan. No son diferencia con el banco: se muestran aparte y no suman.
    anulados = set()
    for r in sorted((t for t in tango_banco if t["tipo"] == "REV" and cruce.libre(t)), key=lambda t: t["fecha"]):
        pares_rev = [t for t in tango_banco if cruce.libre(t) and id(t) not in anulados and t is not r
                     and t["cod"] == r["cod"] and t["c"] == -r["c"] and abs((t["fecha"] - r["fecha"]).days) <= 31]
        if pares_rev:
            o = min(pares_rev, key=lambda t: (abs((t["fecha"] - r["fecha"]).days), t["fila"]))
            anulados.update({id(r), id(o)})

    solo_tango = []
    for t in tango_banco:
        if not cruce.libre(t) or not en_mes(t["fecha"]):
            continue
        if id(t) in anulados:
            solo_tango.append(dict(t, estado="anulado con su reversión: no es diferencia", anulado=True))
            continue
        if id(t) in cruce.ambiguos:
            revisar.append(dict(t, motivo=cruce.ambiguos[id(t)]))
            continue
        ult = ultima.get(t["cod"])
        if ult is None or t["fecha"] > ult - datetime.timedelta(days=cfg["dias_despues"]):
            estado = "todavía no acreditado (probable)"
        else:
            estado = "no aparece en el banco: revisar"
        solo_tango.append(dict(t, estado=estado))

    # Tango de gastos sin par (para comparar con los gastos del banco, no para emparejar)
    def parece_gasto(t):
        return t["tipo"] in ("FPR", "OPF") and t["c"] < 0 and (not t["cuit"] or "BANCO" in _norm(t["contraparte"]))
    gasto_tango = defaultdict(int)
    for t in solo_tango:
        if parece_gasto(t) and not t.get("anulado"):
            gasto_tango[t["cod"]] += t["c"]

    # lo que no pasa por el banco: efectivo y endosos
    efectivo = {_norm(x) for x in cfg["cuentas_efectivo"]}
    cheques = _norm(cfg["cuenta_cheques_terceros"])
    cobros, pagos = set(cfg["comprobantes_cobro"]), set(cfg["comprobantes_pago"])
    no_banco = defaultdict(lambda: {"c": 0})
    cheques_recibidos = [0, 0]
    for t in tango:
        if not en_mes(t["fecha"]) or not t["c"] or t["tipo"] not in cobros | pagos:
            continue
        cta = _norm(t["desc_cuenta"])
        if cta in efectivo:
            clase = "Efectivo (cobro)" if t["tipo"] in cobros else "Efectivo (pago)"
        elif cta == cheques and t["tipo"] in pagos:
            clase = "Cheque de tercero endosado"
        elif cta == cheques and t["tipo"] in cobros:
            cheques_recibidos[0] += 1
            cheques_recibidos[1] += t["c"]
            continue
        else:
            continue
        k = (clase, t["tipo"], t["comprobante"], t["interno"])
        g = no_banco[k]
        g.update({"clase": clase, "fecha": t["fecha"], "tipo": t["tipo"], "comprobante": t["comprobante"],
                  "contraparte": t["contraparte"], "cuit": t["cuit"], "cuenta": t["desc_cuenta"], "texto": t["texto"]})
        g["c"] += t["c"]
    no_banco = sorted(no_banco.values(), key=lambda g: (g["clase"], g["fecha"]))

    # cuentas de Tango que parecen de banco (tienen "Banco") pero no tienen extracto
    sin_extracto = defaultdict(lambda: [0, 0])
    for t in tango:
        if en_mes(t["fecha"]) and t["cod"] not in cods and t["banco"] and t["c"]:
            k = "%s %s" % (t["cod"], t["desc_cuenta"])
            sin_extracto[k][0] += 1
            sin_extracto[k][1] += t["c"]

    # ---- resumen por cuenta y cuenta de control (sumas tomadas de las listas que se escriben)
    cuentas = sorted({b["cuenta"] for b in banco if en_mes(b["fecha"])} | set(), key=str)
    filas_resumen, control_ok = [], True
    revisar_banco = [r for r in revisar if r["lado"] == "Banco"]
    for cta in cuentas:
        del_mes = [b for b in banco if b["cuenta"] == cta and en_mes(b["fecha"])]
        conc = [b for p in pares for b in p["banco"] if b["cuenta"] == cta and en_mes(b["fecha"])]
        inter = [b for p in internas for b in p if b["cuenta"] == cta and en_mes(b["fecha"])]
        sb = [b for b in solo_banco if b["cuenta"] == cta]
        ga = [b for b in gastos if b["cuenta"] == cta]
        rv = [b for b in revisar_banco if b["cuenta"] == cta]
        cod = cfg["cuentas"][cta]
        # lo de Tango se muestra en UNA sola cuenta del extracto (la primera del perfil con ese código),
        # para no contarlo dos veces cuando dos cuentas van a la misma cuenta de Tango (BBVA)
        principal = next(k for k, v in cfg["cuentas"].items() if v == cod)
        st = [t for t in solo_tango if t["cod"] == cod and not t.get("anulado")] if cta == principal else []
        total = sum(b["c"] for b in del_mes)
        partes = sum(b["c"] for b in conc + inter + sb + ga + rv)
        cantidad_ok = len(del_mes) == len(conc) + len(inter) + len(sb) + len(ga) + len(rv)
        ok = total == partes and cantidad_ok
        control_ok &= ok
        movido = sum(abs(b["c"]) for b in del_mes)
        # conciliado por nivel: cantidad de pares y lo movido en el banco (entradas + salidas sin signo)
        por_nivel = {nv: {"n": 0, "c": 0} for nv in NIVELES}
        for par in pares + pares_internas:
            propios = [b for b in par["banco"] if b["cuenta"] == cta and en_mes(b["fecha"])]
            if propios:
                por_nivel[par["nivel"]]["n"] += 1
                por_nivel[par["nivel"]]["c"] += sum(abs(b["c"]) for b in propios)
        filas_resumen.append({
            "cuenta": cta, "cod": cod, "n": len(del_mes), "por_nivel": por_nivel, "movido": movido,
            "entradas": sum(b["c"] for b in del_mes if b["c"] > 0),
            "salidas": sum(b["c"] for b in del_mes if b["c"] < 0),
            "conc_ent": sum(b["c"] for b in conc if b["c"] > 0),
            "conc_sal": sum(b["c"] for b in conc if b["c"] < 0),
            "pct": (sum(abs(b["c"]) for b in conc + inter) / movido) if movido else 0,
            "solo_banco": sum(b["c"] for b in sb), "n_solo_banco": len(sb),
            "sb_ent": sum(b["c"] for b in sb if b["c"] > 0), "sb_sal": sum(b["c"] for b in sb if b["c"] < 0),
            "st_ent": sum(t["c"] for t in st if t["c"] > 0), "st_sal": sum(t["c"] for t in st if t["c"] < 0),
            "gastos": sum(b["c"] for b in ga), "n_gastos": len(ga),
            "internas": sum(b["c"] for b in inter), "revisar": sum(b["c"] for b in rv),
            "solo_tango": sum(t["c"] for t in st), "n_solo_tango": len(st),
            "gasto_tango": gasto_tango.get(cod, 0) if cta == principal else 0,
            "total": total, "partes": partes, "ok": ok,
        })

    return {
        "pares": pares, "internas": internas, "pares_internas": pares_internas,
        "solo_banco": solo_banco, "gastos": gastos,
        "solo_tango": solo_tango, "no_banco": no_banco, "cheques_recibidos": cheques_recibidos,
        "revisar": revisar, "resumen": filas_resumen, "control_ok": control_ok,
        "sin_extracto": dict(sin_extracto), "sin_par_banco": sin_par_banco,
        "reglas": Counter(p["regla"] for p in pares),
        "niveles": Counter(p["nivel"] for p in pares + pares_internas),
    }


# ------------------------------------------------------------------ salidas
def _hoja(wb, titulo, encabezado, filas, formatos=None, anchos=None):
    ws = wb.create_sheet(titulo)
    ws.append(encabezado)
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in filas:
        ws.append(f)
    formatos = formatos or {}
    for j, nombre in enumerate(encabezado, start=1):
        fmt = formatos.get(nombre)
        if fmt:
            for (celda,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                celda.number_format = fmt
        ancho = (anchos or {}).get(nombre, max(10, min(60, len(str(nombre)) + 2)))
        ws.column_dimensions[get_column_letter(j)].width = ancho
    if filas:
        ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    return ws


PESOS, FECHA, PCT = "#,##0.00", "dd/mm/yyyy", "0.0%"


def escribir_excel(inf, ruta, mes):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    p = lambda c: c / 100.0

    # 1. Resumen
    enc = ["Cuenta", "Movimientos banco", "Entradas", "Salidas", "Conciliado entradas", "Conciliado salidas",
           "% conciliado"] + [x for nv in NIVELES for x in ("Pares %s" % nv, "%s ($ movido)" % nv, "%% %s" % nv)] + [
           "Solo en banco (sin gastos)", "Cant. solo en banco", "Gastos e impuestos (banco)",
           "Cant. gastos", "Gastos que Tango tiene sin par", "Internas", "Revisar (banco)",
           "Solo en Tango", "Cant. solo en Tango", "Control"]
    def por_nivel(pn, movido):
        return [x for nv in NIVELES for x in (pn[nv]["n"], p(pn[nv]["c"]), pn[nv]["c"] / movido if movido else 0)]
    filas = [[r["cuenta"], r["n"], p(r["entradas"]), p(r["salidas"]), p(r["conc_ent"]), p(r["conc_sal"]), r["pct"]]
             + por_nivel(r["por_nivel"], r["movido"]) + [
              p(r["solo_banco"]), r["n_solo_banco"], p(r["gastos"]), r["n_gastos"], p(r["gasto_tango"]), p(r["internas"]),
              p(r["revisar"]), p(r["solo_tango"]), r["n_solo_tango"], "OK" if r["ok"] else "NO DA"] for r in inf["resumen"]]
    R = inf["resumen"]
    tot = lambda k: sum(r[k] for r in R)
    movido = sum(abs(r["entradas"]) + abs(r["salidas"]) for r in R)
    conc = sum(r["pct"] * (abs(r["entradas"]) + abs(r["salidas"])) for r in R)
    pn_tot = {nv: {"n": sum(r["por_nivel"][nv]["n"] for r in R), "c": sum(r["por_nivel"][nv]["c"] for r in R)}
              for nv in NIVELES}
    filas.append(["TOTAL", tot("n"), p(tot("entradas")), p(tot("salidas")), p(tot("conc_ent")), p(tot("conc_sal")),
                  conc / movido if movido else 0] + por_nivel(pn_tot, movido) + [p(tot("solo_banco")), tot("n_solo_banco"), p(tot("gastos")),
                  tot("n_gastos"), p(tot("gasto_tango")), p(tot("internas")), p(tot("revisar")), p(tot("solo_tango")),
                  tot("n_solo_tango"), "OK" if inf["control_ok"] else "NO DA"])
    ws = _hoja(wb, "Resumen", enc, filas,
               {k: (PCT if k.startswith("%") else PESOS) for k in enc
                if k not in ("Cuenta", "Movimientos banco", "Control") and not k.startswith(("Cant.", "Pares"))},
               {"Cuenta": 28})
    ws.auto_filter.ref = None
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append(["Cuenta de control: por cada cuenta, lo que se movió en el banco en el mes = conciliado + internas + "
               "solo en banco + gastos + revisar (lado banco). " + ("Da en todas." if inf["control_ok"] else "NO DA: el informe está mal, no usarlo.")])
    ws.append(["Pares por regla: " + ", ".join("%s %d" % (k, v) for k, v in sorted(inf["reglas"].items()))])
    ws.append(["Niveles: Seguro = no hay otra lectura posible · Sugerido = lo más probable, hubo que elegir · "
               "Posible = cierra por importe, conviene mirarlo. '$ movido' = entradas + salidas del banco sin signo; "
               "'%' = sobre todo lo movido en la cuenta. Las internas cuentan como Seguro."])
    ws.append(["Cheques de terceros recibidos en el mes (entran al banco después, con la boleta de depósito): %d renglones, %s"
               % (inf["cheques_recibidos"][0], _m(inf["cheques_recibidos"][1]))])
    if inf["sin_extracto"]:
        ws.append(["Cuentas de Tango sin extracto (quedan afuera del cruce):"])
        for k, (n, c) in sorted(inf["sin_extracto"].items()):
            ws.append(["   " + k, n, p(c)])
            ws.cell(ws.max_row, 3).number_format = PESOS
    if inf["sin_par_banco"]:
        ws.append(["Cuentas del extracto sin cuenta de Tango en el perfil (quedan afuera):"])
        for k, (n, c) in sorted(inf["sin_par_banco"].items()):
            ws.append(["   " + k, n, p(c)])
            ws.cell(ws.max_row, 3).number_format = PESOS

    # 2. Conciliado
    enc = ["Par", "Nivel", "Regla", "Criterio", "Lado", "Fecha", "Cuenta", "Importe", "Concepto / leyenda",
           "Comprobante Tango", "CUIT", "Contraparte", "Días (banco − Tango)"]
    filas = []
    for i, par in enumerate(inf["pares"] + inf["pares_internas"], start=1):
        f_b = min(b["fecha"] for b in par["banco"])
        dias = (f_b - min(t["fecha"] for t in par["tango"])).days if par["tango"] else ""
        cab = [i, par["nivel"], par["regla"], par["criterio"]]
        for b in par["banco"]:
            filas.append(cab + ["Banco", b["fecha"], b["cuenta"], p(b["c"]), b["texto"], "", _cuit_lindo(b["cuit"]), "", dias])
        for t in par["tango"]:
            filas.append(cab + ["Tango", t["fecha"], t["desc_cuenta"], p(t["c"]), t["texto"],
                                ("%s %s" % (t["tipo"], t["comprobante"])).strip(), _cuit_lindo(t["cuit"]), t["contraparte"], dias])
    _hoja(wb, "Conciliado", enc, filas, {"Fecha": FECHA, "Importe": PESOS},
          {"Concepto / leyenda": 50, "Cuenta": 26, "Contraparte": 30, "Comprobante Tango": 22, "Criterio": 70})

    # 3. Solo en banco
    enc = ["Fecha", "Cuenta", "Categoría", "Concepto", "Importe", "CUIT", "Contraparte (según Tango)",
           "Candidato probable en Tango", "ID en Movimientos"]
    filas = [[b["fecha"], b["cuenta"], b["categoria"], b["texto"], p(b["c"]), _cuit_lindo(b["cuit"]), b["contraparte"],
              b["pista"], b["id"]] for b in sorted(inf["solo_banco"], key=lambda b: (b["cuenta"], b["fecha"]))]
    _hoja(wb, "Solo en banco", enc, filas, {"Fecha": FECHA, "Importe": PESOS},
          {"Concepto": 55, "Cuenta": 26, "Candidato probable en Tango": 55, "Contraparte (según Tango)": 30})

    # 4. Gastos bancarios
    enc = ["Fecha", "Cuenta", "Categoría", "Concepto", "Importe", "ID en Movimientos"]
    filas = [[b["fecha"], b["cuenta"], b["categoria"], b["texto"], p(b["c"]), b["id"]]
             for b in sorted(inf["gastos"], key=lambda b: (b["cuenta"], b["fecha"]))]
    _hoja(wb, "Gastos bancarios", enc, filas, {"Fecha": FECHA, "Importe": PESOS}, {"Concepto": 55, "Cuenta": 26})

    # 5. Solo en Tango
    enc = ["Fecha", "Cuenta", "Tipo", "Comprobante", "Contraparte", "CUIT", "Importe", "Leyenda", "Estado"]
    filas = [[t["fecha"], t["desc_cuenta"], t["tipo"], t["comprobante"], t["contraparte"], _cuit_lindo(t["cuit"]),
              p(t["c"]), t["texto"], t["estado"]] for t in sorted(inf["solo_tango"], key=lambda t: (t["cod"], t["fecha"]))]
    _hoja(wb, "Solo en Tango", enc, filas, {"Fecha": FECHA, "Importe": PESOS},
          {"Cuenta": 26, "Contraparte": 30, "Leyenda": 40, "Estado": 32})

    # 6. No pasa por banco
    enc = ["Clase", "Fecha", "Tipo", "Comprobante", "Contraparte", "CUIT", "Cuenta Tango", "Importe", "Leyenda"]
    filas = [[g["clase"], g["fecha"], g["tipo"], g["comprobante"], g["contraparte"], _cuit_lindo(g["cuit"]), g["cuenta"],
              p(g["c"]), g["texto"]] for g in inf["no_banco"]]
    ws = _hoja(wb, "No pasa por banco", enc, filas, {"Fecha": FECHA, "Importe": PESOS}, {"Contraparte": 30, "Leyenda": 40})
    ws.append([])
    ws.append(["Cheques de terceros recibidos en el mes (solo total)", None, None, None, None, None, None,
               p(inf["cheques_recibidos"][1])])
    ws.cell(ws.max_row, 8).number_format = PESOS

    # 7. Revisar
    enc = ["Lado", "Fecha", "Cuenta", "Tipo / categoría", "Comprobante / ID", "Importe", "Concepto / leyenda", "Motivo"]
    filas = []
    for r in inf["revisar"]:
        if r["lado"] == "Banco":
            filas.append(["Banco", r["fecha"], r["cuenta"], r["categoria"], r["id"], p(r["c"]), r["texto"], r["motivo"]])
        else:
            filas.append(["Tango", r["fecha"], r["desc_cuenta"], r["tipo"], r["comprobante"], p(r["c"]), r["texto"], r["motivo"]])
    _hoja(wb, "Revisar", enc, filas, {"Fecha": FECHA, "Importe": PESOS}, {"Motivo": 60, "Concepto / leyenda": 45, "Cuenta": 26})

    wb.save(ruta)
    return ruta


def resumen_md(inf, mes, ruta_xlsx, archivo_tango, archivo_sheet):
    L = ["# Cruce banco ↔ Tango · %s" % mes, "",
         "Banco: solapa Movimientos de `%s` · Tango: `%s`." % (os.path.basename(archivo_sheet), os.path.basename(archivo_tango)), ""]
    if not inf["control_ok"]:
        L += ["## ⚠️ LA CUENTA DE CONTROL NO DA — el informe está mal, no usarlo", ""]
    L += ["Entradas / salidas por separado (no se netean). Entre paréntesis, cantidad de renglones.", "",
          "| Cuenta | Movs | Movido (entra / sale) | Conciliado | Solo en banco (entra / sale) | Gastos banco | Solo en Tango (entra / sale) |",
          "|---|---|---|---|---|---|---|"]
    for r in inf["resumen"]:
        L.append("| %s | %d | %s / %s | %.0f %% | %s / %s (%d) | %s (%d) | %s / %s (%d) |" % (
            r["cuenta"], r["n"], _m(r["entradas"]), _m(r["salidas"]), r["pct"] * 100, _m(r["sb_ent"]), _m(r["sb_sal"]),
            r["n_solo_banco"], _m(r["gastos"]), r["n_gastos"], _m(r["st_ent"]), _m(r["st_sal"]), r["n_solo_tango"]))
    L += ["", "Pares por regla: " + (", ".join("%s %d" % (k, v) for k, v in sorted(inf["reglas"].items())) or "ninguno")
          + " · internas de la misma cuenta: %d" % len(inf["internas"]), ""]
    L += ["## Qué tan firme es lo conciliado", "",
          "Pares y lo movido en el banco (entradas + salidas, sin signo). Mirar solo Sugerido y Posible.", "",
          "| Cuenta | " + " | ".join(NIVELES) + " |", "|---|" + "---|" * len(NIVELES)]
    for r in inf["resumen"]:
        L.append("| %s | %s |" % (r["cuenta"], " | ".join(
            "%d · %s (%.0f %%)" % (r["por_nivel"][nv]["n"], _m(r["por_nivel"][nv]["c"]),
                                   100.0 * r["por_nivel"][nv]["c"] / r["movido"] if r["movido"] else 0) for nv in NIVELES)))
    L.append("")

    L += ["## Lo más grande que falta imputar en Tango", ""]
    for b in sorted(inf["solo_banco"], key=lambda b: -abs(b["c"]))[:10]:
        L.append("- %s · %s · %s · %s%s" % (b["fecha"].strftime("%d/%m"), b["cuenta"], _m(b["c"]), b["texto"][:70],
                                            (" · " + b["pista"]) if b["pista"] else ""))
    L += ["", "## Lo más grande que Tango tiene y el banco no", ""]
    for t in sorted((t for t in inf["solo_tango"] if not t.get("anulado")), key=lambda t: -abs(t["c"]))[:10]:
        L.append("- %s · %s · %s %s · %s · %s · %s" % (t["fecha"].strftime("%d/%m"), t["desc_cuenta"], t["tipo"], t["comprobante"],
                                                     t["contraparte"] or t["texto"][:40], _m(t["c"]), t["estado"]))
    tg = sum(r["gastos"] for r in inf["resumen"])
    L += ["", "## Gastos e impuestos bancarios", "",
          "En el banco sin par: %s en %d renglones. Lo que Tango tiene de gastos sin par: %s." % (
              _m(tg), sum(r["n_gastos"] for r in inf["resumen"]), _m(sum(r["gasto_tango"] for r in inf["resumen"])))]
    efe = [g for g in inf["no_banco"] if g["clase"].startswith("Efectivo")]
    endo = [g for g in inf["no_banco"] if g["clase"].startswith("Cheque")]
    L += ["", "## No pasa por el banco (informativo)", "",
          "- Efectivo: %d comprobantes, %s" % (len(efe), _m(sum(g["c"] for g in efe))),
          "- Cheques de terceros endosados: %d órdenes de pago, %s" % (len(endo), _m(sum(g["c"] for g in endo))),
          "- Cheques de terceros recibidos: %s (llegan al banco con la boleta de depósito)" % _m(inf["cheques_recibidos"][1])]
    if inf["revisar"]:
        L += ["", "## Revisar (%d)" % len(inf["revisar"]), ""]
        for r in inf["revisar"][:15]:
            L.append("- %s · %s · %s · %s" % (r["lado"], r["fecha"].strftime("%d/%m/%Y"), _m(r["c"]), r["motivo"]))
    L += ["", "Detalle completo: `%s`" % ruta_xlsx]
    return "\n".join(L)


# ------------------------------------------------------------------ main
def correr(cliente, mes, sheet, tango, salidas, hoy=None):
    cfg = cargar_config(cliente)
    anio, m = (int(x) for x in mes.split("-"))
    mes_desde = datetime.date(anio, m, 1)
    mes_hasta = (datetime.date(anio + (m == 12), m % 12 + 1, 1) - datetime.timedelta(days=1))
    hoy = hoy or datetime.date.today()
    banco, ultima, sin_par = leer_banco(sheet, cfg, mes_desde - datetime.timedelta(days=MARGEN_CARGA),
                                        mes_hasta + datetime.timedelta(days=MARGEN_CARGA))
    renglones, archivo_tango = leer_tango(tango, cfg)
    inf = armar_informe(banco, renglones, cfg, mes_desde, mes_hasta, ultima, hoy, sin_par)
    os.makedirs(salidas, exist_ok=True)
    ruta = escribir_excel(inf, os.path.join(salidas, "cruce_%s.xlsx" % mes), mes)
    md = resumen_md(inf, mes, ruta, archivo_tango, sheet)
    with io.open(os.path.join(salidas, "resumen_cruce_%s.md" % mes), "w", encoding="utf-8") as f:
        f.write(md + "\n")
    return inf, md


def main():
    ap = argparse.ArgumentParser(description="Cruce banco ↔ Tango de un mes (informe, no toca la Sheet)")
    ap.add_argument("--cliente", default="navar")
    ap.add_argument("--mes", required=True, help="AAAA-MM")
    ap.add_argument("--sheet", required=True, help="export de la Sheet en Excel (solapa Movimientos)")
    ap.add_argument("--tango", required=True, help="detalle de comprobantes de Tesorería (archivo o carpeta)")
    ap.add_argument("--salidas", required=True)
    ap.add_argument("--hoy", default=None)
    a = ap.parse_args()
    hoy = datetime.date.fromisoformat(a.hoy) if a.hoy else None
    inf, md = correr(a.cliente, a.mes, a.sheet, a.tango, a.salidas, hoy)
    print(md)
    if not inf["control_ok"]:
        sys.exit(1)


if __name__ == "__main__":
    main()
