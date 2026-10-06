# -*- coding: utf-8 -*-
"""
lector/cuotas_prestamos — las cuotas REALES de los préstamos, sacadas de las tablas de los bancos.

PARA QUÉ
    El mapa de deuda (Bancos_Navar.xlsx) trae solo el valor de la cuota y un texto: las cuotas que
    siguen se ESTIMAN. Y la administración no carga los débitos de préstamos en Tango porque el banco
    debita un solo importe y les falta el desglose (capital, interés, IVA, percepción). Las tablas de
    amortización de los bancos traen las dos cosas. Este módulo las guarda en UN Excel
    ("Cuotas de prestamos.xlsx", en NAVAR - Datos/Deuda bancaria) y lo usan:
      · lector/deuda_bancaria.py: cuotas reales en lugar de estimadas, en los préstamos con tabla;
      · lector/cruce_semanal.py: el desglose de cada cuota que el banco debitó y falta cargar en Tango.

EL EXCEL (una fila por cuota; hoja "Cuotas")
    Banco · Línea del mapa · Préstamo · Cuota · Vencimiento · Estado · Fecha de pago · Capital ·
    Interés · IVA · Percepción IVA · Otros · Punitorios · Total · Fuente
    "Línea del mapa" tiene que ser IGUAL a la columna Producto del mapa de deuda: así se sabe qué
    préstamo reemplaza. "Otros" lleva, con signo, lo que no es capital/interés/IVA (p. ej. un subsidio
    de tasa, que resta). "Estado": Pagada · A vencer · Impaga.

DE DÓNDE SALE CADA TABLA
    Galicia: el PDF "Detalle de cuotas" de Office Banking trae texto → se lee solo con este comando.
    Otros bancos (p. ej. el informe de deuda del Nación, que llega escaneado): se cargan a mano o con
    una ayuda aparte, verificando cada cuota con el total que imprime el banco.

USO
    python lector/cuotas_prestamos.py --tabla "<...>/Cuotas de prestamos.xlsx" \\
        --galicia "<PDF de Office Banking>" --linea "Préstamo Nro. 123456789"
      (agrega o reemplaza las cuotas de ese préstamo en la tabla)
    python lector/cuotas_prestamos.py --tabla "<...>/Cuotas de prestamos.xlsx" --ver
"""

import os
import re
import sys
import argparse
import datetime
import unicodedata

import openpyxl
from openpyxl.styles import Font

ARCHIVO = "Cuotas de prestamos.xlsx"
ENCABEZADO = ["Banco", "Línea del mapa", "Préstamo", "Cuota", "Vencimiento", "Estado", "Fecha de pago", "Capital",
              "Interés", "IVA", "Percepción IVA", "Otros", "Punitorios", "Total", "Fuente"]
CLAVES = ["banco", "linea", "prestamo", "cuota", "vto", "estado", "pago", "capital", "interes", "iva", "percepcion",
          "otros", "punitorios", "total", "fuente"]
NUMEROS = ("capital", "interes", "iva", "percepcion", "otros", "punitorios", "total")


def _norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(s.upper().split())


def _fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    return None


def _plata(v):
    s = format(abs(v), ",.2f").replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-$" if v < 0 else "$") + s


# ------------------------------------------------------------------ el Excel
def leer(ruta):
    """Las cuotas del Excel, como diccionarios. Si no existe, lista vacía."""
    if not ruta or not os.path.exists(ruta):
        return []
    ws = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    hoja = ws["Cuotas"] if "Cuotas" in ws.sheetnames else ws.worksheets[0]
    filas = hoja.iter_rows(values_only=True)
    enc = [_norm(x) for x in next(filas)]
    ix = {k: enc.index(_norm(n)) for k, n in zip(CLAVES, ENCABEZADO) if _norm(n) in enc}
    for obligatoria in ("banco", "linea", "cuota", "vto", "capital", "total"):
        if obligatoria not in ix:
            raise SystemExit("a %s le falta la columna %r" % (os.path.basename(ruta), ENCABEZADO[CLAVES.index(obligatoria)]))
    cuotas = []
    for r in filas:
        c = {k: (r[i] if i < len(r) else None) for k, i in ix.items()}
        if not c.get("linea") or c.get("cuota") in (None, ""):
            continue
        c["vto"], c["pago"] = _fecha(c.get("vto")), _fecha(c.get("pago"))
        c["cuota"] = int(c["cuota"])
        for k in NUMEROS:
            c[k] = float(c.get(k) or 0)
        c["banco"] = str(c["banco"] or "").strip().upper()
        c["linea"] = str(c["linea"]).strip()
        c["estado"] = str(c.get("estado") or "").strip() or "A vencer"
        cuotas.append(c)
    return cuotas


def escribir(cuotas, ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cuotas"
    ws.append(ENCABEZADO)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for c in sorted(cuotas, key=lambda c: (c["banco"], c["linea"], c["cuota"])):
        ws.append([c.get(k) for k in CLAVES])
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            if isinstance(celda.value, float):
                celda.number_format = "#,##0.00"
            elif isinstance(celda.value, datetime.date):
                celda.number_format = "dd/mm/yyyy"
    for j, ancho in enumerate((11, 32, 14, 7, 12, 10, 12, 16, 16, 14, 14, 14, 13, 16, 40), start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = ancho
    ws.freeze_panes = "A2"
    tmp = ruta + ".parte.xlsx"
    wb.save(tmp)
    os.replace(tmp, ruta)
    return ruta


def reemplazar_prestamo(cuotas, nuevas):
    """Saca de la tabla las cuotas de los préstamos que vienen en `nuevas` y pone las nuevas."""
    lineas = {(c["banco"], c["linea"]) for c in nuevas}
    return [c for c in cuotas if (c["banco"], c["linea"]) not in lineas] + list(nuevas)


# ------------------------------------------------------------------ Galicia (Office Banking)
def leer_texto_galicia(texto, linea, fuente=""):
    """Texto del PDF 'Detalle de cuotas' de Galicia → cuotas. Cada renglón de cuota es:
    nro · Abonada/A Vencer · aaaa-mm-dd · $ total · $ capital · $ interés · $ IVA interés · $ IVA percepción · $ otros"""
    m = re.search(r"Pr[eé]stamo Nro:\s*(\d+)", texto)
    prestamo = m.group(1) if m else ""
    plata = r"\$\s*([\d.]+)"
    patron = re.compile(r"^\s*(\d+)\s+(Abonada|A Vencer|Impaga|Vencida)\s+(\d{4}-\d{2}-\d{2})\s+" + r"\s+".join([plata] * 6), re.M)
    cuotas = []
    for m in patron.finditer(texto):
        total, capital, interes, iva, percep, otros = (float(x) for x in m.groups()[3:])
        estado = {"Abonada": "Pagada", "A Vencer": "A vencer"}.get(m.group(2), "Impaga")
        if abs(capital + interes + iva + percep + otros - total) > 0.011:
            raise SystemExit("Galicia cuota %s: las partes no suman el total (%s)" % (m.group(1), total))
        cuotas.append({"banco": "GALICIA", "linea": linea, "prestamo": prestamo, "cuota": int(m.group(1)),
                       "vto": _fecha(m.group(3)), "estado": estado, "pago": None, "capital": capital, "interes": interes,
                       "iva": iva, "percepcion": percep, "otros": otros, "punitorios": 0.0, "total": total,
                       "fuente": fuente})
    if not cuotas:
        raise SystemExit("no encontré cuotas en el PDF de Galicia (¿cambió el formato?)")
    return cuotas


def leer_pdf_galicia(ruta, linea):
    from pypdf import PdfReader
    texto = "\n".join((p.extract_text() or "") for p in PdfReader(ruta).pages)
    return leer_texto_galicia(texto, linea, "Galicia Office Banking · " + os.path.basename(ruta))


# ------------------------------------------------------------------ para el cruce y el mail
def desglose(c):
    """'cuota 6 de 123456789: capital $X · interés $Y · IVA $Z · percepción $W'"""
    partes = ["capital " + _plata(c["capital"]), "interés " + _plata(c["interes"])]
    if c["iva"]:
        partes.append("IVA " + _plata(c["iva"]))
    if c["percepcion"]:
        partes.append("percepción IVA " + _plata(c["percepcion"]))
    if c["otros"]:
        partes.append(("cargos (IVA, percepciones) " if c["otros"] > 0 else "subsidio/otros ") + _plata(c["otros"]))
    if c["punitorios"]:
        partes.append("punitorios " + _plata(c["punitorios"]))
    nro = str(c.get("prestamo") or "")
    nombre = c["linea"] if (not nro or nro.lstrip("0") in c["linea"]) else "%s (%s)" % (nro, c["linea"])
    return "cuota %d de %s: %s" % (c["cuota"], nombre, " · ".join(partes))


def para_debito(cuotas, banco, fecha, importe, tolerancia=1.0):
    """La cuota que corresponde a un débito del banco (importe negativo o positivo, en pesos).
    1) una cuota PAGADA ese mismo día por ese importe (la tabla trae la fecha de pago y los punitorios);
    2) si no, una cuota de ese importe que vence entre 10 días después y 45 antes del débito.
    Si hay más de una posible, ninguna (mejor no mostrar un desglose equivocado)."""
    banco, monto = str(banco or "").upper(), abs(importe)
    del_banco = [c for c in cuotas if c["banco"] == banco and abs(c["total"] - monto) <= tolerancia]
    pagadas = [c for c in del_banco if c.get("pago") == fecha]
    if len(pagadas) == 1:
        return pagadas[0]
    cerca = [c for c in del_banco if c["vto"] and -10 <= (fecha - c["vto"]).days <= 45]
    return cerca[0] if len(cerca) == 1 else None


def main(argv=None):
    ap = argparse.ArgumentParser(description="Tabla de cuotas reales de los préstamos")
    ap.add_argument("--tabla", required=True, help="el Excel de cuotas (se crea si no existe)")
    ap.add_argument("--galicia", help="PDF 'Detalle de cuotas' de Galicia Office Banking")
    ap.add_argument("--linea", help="nombre del préstamo en el mapa de deuda (columna Producto)")
    ap.add_argument("--ver", action="store_true")
    a = ap.parse_args(argv)
    cuotas = leer(a.tabla)
    if a.galicia:
        if not a.linea:
            raise SystemExit("falta --linea: el nombre del préstamo en el mapa de deuda")
        nuevas = leer_pdf_galicia(a.galicia, a.linea)
        cuotas = reemplazar_prestamo(cuotas, nuevas)
        escribir(cuotas, a.tabla)
        print("Galicia %s: %d cuotas cargadas en %s" % (a.linea, len(nuevas), a.tabla))
    if a.ver or not a.galicia:
        por = {}
        for c in cuotas:
            por.setdefault((c["banco"], c["linea"]), []).append(c)
        for (banco, linea), cs in sorted(por.items()):
            pend = [c for c in cs if c["estado"] != "Pagada"]
            prox = min(pend, key=lambda c: c["vto"]) if pend else None
            print("%-10s %-32s %2d cuotas · %2d pendientes%s" % (banco, linea, len(cs), len(pend),
                  (" · próxima %d el %s por %s" % (prox["cuota"], prox["vto"].strftime("%d/%m/%Y"), _plata(prox["total"]))) if prox else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
