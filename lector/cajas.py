# -*- coding: utf-8 -*-
"""
lector/cajas — los movimientos de CAJA (efectivo) de A y de AA, sacados de Tango.

PARA QUÉ (tarea 48, pedido de NAVAR 06/10/2026)
    El cash mostraba solo la caja de AA, y sin decir de qué caja de Tango salía cada peso ni para
    qué. Este lector arma, para las DOS empresas, un renglón de Movimientos por cada pase por una
    caja, con dos columnas nuevas: "Cuenta Tango" (de qué caja sale) y "Leyenda" (lo que escribió
    quien cargó la orden). El cash suma cada empresa en su caja: "Caja A" y "Caja AA".

DE DÓNDE SALE
    El "Detalle de comprobantes" de tesorería de Tango (consultas 21 de A y 22 de AA): un renglón
    por cuenta imputada. Un pago por caja típico tiene dos renglones: el gasto (Debe) y la caja
    (Haber). A la caja le corresponde SOLO su renglón: si un recibo entra mitad en caja y mitad en
    el banco, acá va la mitad de caja (la otra la trae el extracto del banco).
    Qué cuentas son caja de cada empresa lo dice clientes/<cliente>/perfil.json → "cajas".

LO QUE NO PUEDE PASAR: CONTAR DOS VECES
    El cash suma los bancos y las dos cajas. La plata que pasa de una caja a otra, o a una cuenta
    propia de pase (perfil → cajas.cuentas_de_pase: Mercado Pago, la otra empresa), es
    "Transferencia Interna": no es cobro ni pago. El DEPÓSITO de la caja en el banco es especial:
    el extracto del banco lo trae como cobro ("Depósito en efectivo" → Cobranza Facturas), y ese
    cobro ya se contó cuando la plata entró a la caja. Por eso en la caja va como Cobranza Facturas
    EN NEGATIVO: los dos se cancelan y el cobro queda una sola vez. Los saldos no cambian por esto.

CATEGORÍAS (mirando hacia dónde va la plata: las cuentas del lado contrario a la caja;
            las retenciones y lo que se descuenta en el mismo pago no deciden nada)
    pagos entre las dos empresas (NAVAR S.A.)      → Transferencia Interna
    deudores (clientes)                            → A: Cobranza Facturas        · AA: Cobranza AA
    sueldos, cargas sociales, aportes sindicales   → Sueldos y Jornales
    honorarios, cuenta particular de un director   → Honorarios y Dividendos (si entra plata: Otros)
    préstamo                                       → Prestamo
    acreedores (proveedores)                       → A: Proveedores MP y Logist. · AA: Proveedores AA
    de la caja al banco (depósito)                 → Cobranza Facturas en negativo (ver arriba)
    entre la caja y la cartera de cheques          → Cheques (cheques cambiados por efectivo)
    entre cajas, del banco a la caja, pases        → Transferencia Interna
    el resto (gastos varios, viáticos...)          → Otros
    Las reversiones (REV) quedan con su signo: restan.

SALIDA
    para_pegar_cajas_<hoy>.xlsx   (solapa Movimientos, con "Cuenta Tango" y "Leyenda" al final)
    resumen_cajas_<hoy>.md
    Origen: AA → "Tango AA · cajas · …" (el cash reconoce lo real de AA por "Tango AA").
            A  → "Tango caja A · cajas · …" (NO puede empezar con "Tango A": ese patrón agarraría AA).

USO
    python lector/cajas.py --a "<carpeta>/A tesoreria detalle 2026-10-07.xlsx" \
                           --aa "<carpeta>/AA tesoreria detalle 2026-10-07.xlsx" --cliente navar
"""

import os
import sys
import json
import argparse
import datetime
import unicodedata
from collections import OrderedDict, defaultdict

import openpyxl

BASE_REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

ENC_MOV = ["ID", "Fecha", "Empresa", "Tipo", "Categoria", "Concepto / Detalle", "Importe", "Medio de Pago",
           "Banco / Cuenta", "Origen", "Estado", "Referencia", "Semana (lunes)", "Observaciones",
           "Cuenta Tango", "Leyenda"]
COLS_CON_FORMULA = {"Semana (lunes)"}
MARCAS = {"A": "Tango caja A", "AA": "Tango AA"}

# Categorías por empresa: tienen que ser las que ya existen en la Sheet (validación de Categoria).
COBRANZA = {"A": "Cobranza Facturas", "AA": "Cobranza AA"}
PROVEEDORES = {"A": "Proveedores MP y Logist.", "AA": "Proveedores AA"}
DEPOSITO = "deposito"   # de la caja al banco: en la planilla va como cobranza en negativo (ver categoria)


def _norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(s.upper().split())


def _num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = str(v or "").strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(s[:10], formato).date()
        except ValueError:
            pass
    return None


def config_cajas(cliente):
    with open(os.path.join(BASE_REPO, "clientes", cliente, "perfil.json"), encoding="utf-8") as f:
        perfil = json.load(f)
    cfg = perfil.get("cajas")
    if not cfg:
        raise ValueError("perfil.json no tiene la sección 'cajas'")
    for empresa in ("A", "AA"):
        if not cfg.get(empresa, {}).get("cuentas") or not cfg[empresa].get("nombre_en_el_cash"):
            raise ValueError("perfil.json → cajas.%s necesita 'cuentas' y 'nombre_en_el_cash'" % empresa)
    return cfg


def _es_de(desc, cuentas):
    """¿La cuenta empieza con alguna de las de la lista? (sin importar mayúsculas ni tildes)"""
    d = _norm(desc)
    return any(d.startswith(_norm(c)) for c in cuentas)


def leer_detalle(ruta):
    """Lee el detalle por ENCABEZADO (no por posición). Devuelve una lista de dicts."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        filas = wb.active.iter_rows(values_only=True)
        enc = [str(c).strip() if c is not None else "" for c in next(filas, ())]
        claves = {_norm(e): i for i, e in enumerate(enc) if e}
        obligatorias = ["NRO. INTERNO", "COD. COMPROBANTE", "DESC. CUENTA",
                        "DEBE (CTE) (RENGLON)", "HABER (CTE) (RENGLON)"]
        faltan = [c for c in obligatorias if c not in claves]
        if faltan or ("FECHA" not in claves and "FECHA DE EMISION" not in claves):
            raise ValueError("%s: faltan columnas del detalle de tesorería (%s)"
                             % (os.path.basename(ruta), ", ".join(faltan) or "Fecha"))

        def col(r, nombre):
            i = claves.get(_norm(nombre))
            return r[i] if i is not None and i < len(r) else None

        salida = []
        for r in filas:
            if not r or all(v in (None, "") for v in r):
                continue
            salida.append({
                "fecha": _fecha(col(r, "Fecha")) or _fecha(col(r, "Fecha de emisión")),
                "tipo": str(col(r, "Cód. comprobante") or "").strip().upper(),
                "comprobante": str(col(r, "Comprobante") or "").strip(),
                "interno": str(col(r, "Nro. interno") or "").strip(),
                "cuenta": str(col(r, "Desc. cuenta") or col(r, "Desc. contable") or "").strip(),
                "tipo_cuenta": str(col(r, "COD_TIPO_CUENTA") or "").strip().upper(),
                "debe": _num(col(r, "Debe (cte) (renglón)")),
                "haber": _num(col(r, "Haber (cte) (renglón)")),
                "razon": str(col(r, "Razón social (encab.)") or col(r, "Proveedor (encab.)") or "").strip(),
                "leyenda": str(col(r, "Leyenda") or "").strip(),
            })
        return salida
    finally:
        wb.close()


def _es_banco(r):
    # COD_TIPO_CUENTA = B es banco (incluye tarjetas). Si la columna no viniera, por el nombre.
    if r["tipo_cuenta"]:
        return r["tipo_cuenta"] == "B"
    n = _norm(r["cuenta"])
    return n.startswith("BANCO") or n.startswith("BCO")


def categoria(empresa, cajas_renglones, otros, cfg):
    """La categoría de un comprobante de caja, mirando HACIA DÓNDE va (o de dónde viene) la plata.

    Solo cuentan las cuentas del lado contrario a la caja: en un pago, la caja está en el Haber y
    el destino (proveedor, sueldos, banco) en el Debe. Lo que está del mismo lado que la caja son
    otras fuentes del mismo pago (retenciones, el convenio gremial que se descuenta al productor,
    el banco que pone la otra parte): no dicen qué es el comprobante.

    Primero manda el TERCERO (cliente, proveedor, empleados). Si no hay tercero:
      - de la caja al BANCO (depósito): el extracto del banco lo trae como cobro ("Depósito en
        efectivo" → Cobranza). Ese cobro ya se contó cuando la plata entró a la caja; por eso acá
        va como COBRANZA EN NEGATIVO: se cancelan y el cobro queda contado una sola vez;
      - entre la caja y la CARTERA de cheques: la cartera está fuera del cash; un cheque cambiado
        por efectivo es plata que entra (Cheques, como cuando se deposita en el banco);
      - entre cajas, del banco a la caja o a una cuenta de pase (Mercado Pago, la otra empresa):
        plata que se mueve entre cuentas propias → Transferencia Interna.
    """
    razon = _norm(cajas_renglones[0]["razon"])
    if "NAVAR S.A" in razon or "NAVAR SA" in razon or "NAV007" in razon:
        return "Transferencia Interna"          # pagos entre las dos empresas (igual que tesoreria_aa)
    neto_caja = sum(r["debe"] - r["haber"] for r in cajas_renglones)
    entra = neto_caja > 0
    if cajas_renglones[0]["tipo"] == "REV":
        entra = not entra                       # una anulación tiene el signo al revés del original
    # Del lado contrario a la caja (si la caja quedó en cero, se miran todas).
    contra = [o for o in otros if neto_caja == 0 or (o["debe"] - o["haber"]) * neto_caja < 0]
    contra = [o for o in contra if "RETENC" not in _norm(o["cuenta"]) and not _norm(o["cuenta"]).startswith("RET.")]
    texto = " ".join(_norm(o["cuenta"]) for o in contra)
    leyenda = _norm(cajas_renglones[0]["leyenda"])
    if entra and "DEUDOR" in texto:
        return COBRANZA[empresa]
    if any(p in texto for p in ("SUELDO", "JORNAL", "CARGAS SOCIALES", "CONVENIO")) or \
            any(p in leyenda for p in ("SUELDO", "JORNAL", "HORAS EXTRA", "OSECAC", "FAECYS", "UATRE",
                                       "SINDICATO", "OBRA SOCIAL", "APORTES")):
        return "Sueldos y Jornales"
    if any(p in texto for p in ("HONOR", "CTA. PARTICULAR", "DIVIDEND")):
        # Plata que pone un socio no es un honorario: va a revisar (Otros).
        return "Otros" if entra else "Honorarios y Dividendos"
    if "PRESTAMO" in texto:
        return "Prestamo"
    if "ACREEDOR" in texto:
        return PROVEEDORES[empresa]
    if "DEUDOR" in texto:
        return COBRANZA[empresa]
    # Sin tercero: ¿va a otra cuenta propia?
    todas_las_cajas = cfg["A"]["cuentas"] + cfg["AA"]["cuentas"]
    pase = [_norm(p) for p in cfg.get("cuentas_de_pase", [])]
    if len({_norm(r["cuenta"]) for r in cajas_renglones}) > 1:
        return "Transferencia Interna"          # de una caja a otra, en el mismo comprobante
    if any(_es_de(o["cuenta"], todas_las_cajas) or any(p and p in _norm(o["cuenta"]) for p in pase) for o in contra):
        return "Transferencia Interna"
    if any(_es_banco(o) for o in contra):
        return DEPOSITO if not entra else "Transferencia Interna"
    if any(o["tipo_cuenta"] == "C" for o in contra):
        return "Cheques"
    return "Otros"


def movimientos(renglones, empresa, cfg, nombre_archivo):
    """Un renglón de Movimientos por cada pase por una caja de esa empresa."""
    cajas = cfg[empresa]["cuentas"]
    por_comprobante = OrderedDict()
    for r in renglones:
        por_comprobante.setdefault((r["tipo"], r["interno"] or r["comprobante"]), []).append(r)
    salida = []
    for (_, _), grupo in por_comprobante.items():
        de_caja = [r for r in grupo if _es_de(r["cuenta"], cajas)]
        if not de_caja:
            continue
        otros = [r for r in grupo if r not in de_caja]
        cat = categoria(empresa, de_caja, otros, cfg)
        deposito = cat == DEPOSITO
        if deposito:
            cat = "Cobranza Facturas"           # la misma categoría con que el extracto trae el depósito
        for r in de_caja:
            importe = round(r["debe"] - r["haber"], 2)
            if not importe or not r["fecha"]:
                continue
            # Tipo: lo que ES el movimiento, no el signo. Una cobranza es ingreso aunque reste (vuelto,
            # compensación, anulación): así se descuenta en la misma fila del cash. Un pago a proveedor
            # o un sueldo es egreso aunque vuelva plata. El resto, por el signo (una anulación REV, al revés).
            if cat.startswith("Cobranza") or cat == "Cheques":
                ingreso = True
            elif cat in ("Proveedores MP y Logist.", "Proveedores AA", "Sueldos y Jornales", "Honorarios y Dividendos"):
                ingreso = False
            else:
                ingreso = (importe > 0) != (r["tipo"] == "REV")
            quien = r["razon"]
            detalle = " · ".join(x for x in (quien, r["leyenda"]) if x) or r["tipo"]
            if deposito:
                detalle = "Depósito de la caja en el banco (el extracto lo trae como cobro)" + \
                    (" · " + r["leyenda"] if r["leyenda"] else "")
            salida.append(OrderedDict([
                ("ID", None), ("Fecha", r["fecha"]), ("Empresa", empresa),
                ("Tipo", "Ingreso" if ingreso else "Egreso"), ("Categoria", cat),
                ("Concepto / Detalle", detalle[:120]), ("Importe", importe), ("Medio de Pago", "Efectivo"),
                ("Banco / Cuenta", cfg[empresa]["nombre_en_el_cash"]),
                ("Origen", MARCAS[empresa] + " · cajas · " + nombre_archivo), ("Estado", "Real"),
                ("Referencia", (r["tipo"] + " " + r["comprobante"]).strip()), ("Semana (lunes)", None),
                ("Observaciones", "Tango · detalle de tesorería · caja de " + empresa +
                 (" · depósito: resta acá porque el banco ya lo suma como cobro" if deposito else "")),
                ("Cuenta Tango", r["cuenta"]), ("Leyenda", r["leyenda"]),
            ]))
    return salida


def escribir_para_pegar(filas, carpeta, hoy):
    ruta = os.path.join(carpeta, "para_pegar_cajas_%s.xlsx" % hoy.isoformat())
    temporal = ruta + ".parte.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    ws.append(ENC_MOV)
    for x in filas:
        ws.append([None if c in COLS_CON_FORMULA else x.get(c) for c in ENC_MOV])
    for fila in ws.iter_rows(min_row=2):
        fila[1].number_format = "dd/mm/yyyy"
        fila[6].number_format = "#,##0.00"
        for c in fila:
            if isinstance(c.value, str):
                c.data_type = "s"   # nombres y leyendas son texto, nunca fórmulas
    try:
        wb.save(temporal)
        os.replace(temporal, ruta)
    finally:
        if os.path.exists(temporal):
            os.unlink(temporal)
    return ruta


def _m(v):
    return ("-" if v < 0 else "") + "$" + format(int(round(abs(v))), ",d").replace(",", ".")


def resumen(filas, hoy, archivos):
    L = ["# Cajas de A y AA → Movimientos · %s" % hoy.strftime("%d/%m/%Y"), "",
         "Archivos: " + ", ".join("`%s`" % a for a in archivos), "",
         "| Empresa | Cuenta de Tango | Mes | Movimientos | Entra | Sale | Pases internos |",
         "|---|---|---|---:|---:|---:|---:|"]
    tabla = defaultdict(lambda: [0, 0.0, 0.0, 0])
    for x in filas:
        k = (x["Empresa"], x["Cuenta Tango"], x["Fecha"].strftime("%Y-%m"))
        t = tabla[k]
        t[0] += 1
        if x["Importe"] > 0:
            t[1] += x["Importe"]
        else:
            t[2] += -x["Importe"]
        if x["Categoria"] == "Transferencia Interna":
            t[3] += 1
    for k in sorted(tabla):
        t = tabla[k]
        L.append("| %s | %s | %s | %d | %s | %s | %d |" % (k[0], k[1], k[2], t[0], _m(t[1]), _m(t[2]), t[3]))
    porcat = defaultdict(float)
    for x in filas:
        porcat[(x["Empresa"], x["Categoria"])] += x["Importe"]
    L += ["", "## Neto por categoría", "", "| Empresa | Categoría | Neto |", "|---|---|---:|"]
    for k in sorted(porcat):
        L.append("| %s | %s | %s |" % (k[0], k[1], _m(porcat[k])))
    L += ["", "Las transferencias internas no son ni cobros ni pagos: mueven plata entre cuentas propias.",
          "El saldo de cada caja en el cash = último arqueo (Saldos Bancarios) + estos movimientos posteriores."]
    return "\n".join(L)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Movimientos de caja de A y AA desde el detalle de tesorería")
    ap.add_argument("--a", required=True, help="detalle de tesorería de A (.xlsx)")
    ap.add_argument("--aa", required=True, help="detalle de tesorería de AA (.xlsx)")
    ap.add_argument("--cliente", default="navar")
    ap.add_argument("--hoy", default=datetime.date.today().isoformat())
    ap.add_argument("--salidas", default=None, help="carpeta de salida (default: la del archivo de A)")
    a = ap.parse_args(argv)
    hoy = datetime.date.fromisoformat(a.hoy)
    # Las dos o ninguna: el importador pisa las filas de caja de las dos empresas juntas;
    # publicar una sola borraría la otra.
    for ruta in (a.a, a.aa):
        if not os.path.isfile(ruta):
            raise ValueError("falta el detalle de tesorería: %s; no se publica ninguna caja" % ruta)
    cfg = config_cajas(a.cliente)
    filas = []
    for empresa, ruta in (("A", a.a), ("AA", a.aa)):
        filas += movimientos(leer_detalle(ruta), empresa, cfg, os.path.basename(ruta))
    filas.sort(key=lambda x: (x["Fecha"], x["Empresa"], x["Referencia"]))
    for i, x in enumerate(filas, 1):
        x["ID"] = i
    carpeta = a.salidas or os.path.dirname(os.path.abspath(a.a))
    os.makedirs(carpeta, exist_ok=True)
    ruta = escribir_para_pegar(filas, carpeta, hoy)
    md = resumen(filas, hoy, [os.path.basename(a.a), os.path.basename(a.aa)])
    with open(os.path.join(carpeta, "resumen_cajas_%s.md" % hoy.isoformat()), "w", encoding="utf-8") as f:
        f.write(md + "\n")
    print(md)
    print("\n- `%s`: solapa Movimientos (%d filas)" % (ruta, len(filas)))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (OSError, ValueError) as ex:
        sys.exit(str(ex))
