# -*- coding: utf-8 -*-
"""
lector/tesoreria_aa — la operación en EFECTIVO de la empresa AA, sacada de Tango.

AA no tiene banco: cobra y paga en efectivo, y Tango lo registra en Tesorería (recibos, órdenes
de pago, gastos). Este lector toma el export de Tango Live "Movimientos de tesorería" de AA
(un renglón por comprobante) y lo convierte en filas de la solapa Movimientos, con Origen
"Tango AA" y Estado "Real". Así "Cobranza AA" y "Proveedores AA" del cash tienen parte real,
igual que los renglones de A la tienen por el extracto.

EL EXPORT (Tango Live → Tesorería → Comprobantes, empresa NAVAR SA Otros; consulta 11591)
    Tipo · Comprobante · Fecha · Concepto · Clase · Total (cte) · Cód. relacionado ·
    Desc. relacionado · Clasificación · Fecha de emisión
    Clase "Cobros" (REC)                      → Ingreso · Cobranza AA
    Clase "Pagos" (O/P, OPF, FPR, REV)        → Egreso · Proveedores AA
                                                (si el concepto habla de sueldos → Sueldos y Jornales;
                                                 si el relacionado es NAVAR S.A. → Transferencia Interna)
    Clase "Otros movimientos..." (EXT, REV)   → Egreso · Otros (viáticos, gastos varios y las salidas
                                                de caja que Tango llama "débito y gastos bancarios")
    Las anulaciones (REV) vienen con importe negativo y quedan con el signo al revés: restan.

    Se cargan solo los movimientos desde --desde (default 2026-06-01, como los extractos).

SALIDA
    para_pegar_tesoreria_aa_<hoy>.xlsx   (solapa Movimientos) → importación automática
    resumen_tesoreria_aa_<hoy>.md

USO
    python lector/tesoreria_aa.py --archivo "<carpeta>/AA movimientos tesoreria 2026-09-22.xlsx" --cliente navar
"""

import os
import io
import re
import sys
import argparse
import datetime
import unicodedata
from collections import OrderedDict, defaultdict

import openpyxl

ENC_MOV = ["ID", "Fecha", "Empresa", "Tipo", "Categoria", "Concepto / Detalle", "Importe", "Medio de Pago",
           "Banco / Cuenta", "Origen", "Estado", "Referencia", "Semana (lunes)", "Observaciones"]
COLS_CON_FORMULA = {"Semana (lunes)"}
MARCA = "Tango AA"
DESDE = datetime.date(2026, 6, 1)


def _norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(s.upper().split())


def _m(v):
    return ("-" if v < 0 else "") + "$" + format(int(round(abs(v))), ",d").replace(",", ".")


def clasificar(r):
    """(tipo, categoria) para un renglón del export."""
    clase, concepto, rel = _norm(r.get("Clase")), _norm(r.get("Concepto")), _norm(r.get("Desc. relacionado"))
    if clase.startswith("COBRO"):
        return "Ingreso", "Cobranza AA"
    if "NAVAR S.A" in rel or "NAV007" in rel:
        return "Egreso", "Transferencia Interna"
    if clase.startswith("PAGO"):
        if "SUELDO" in concepto or "HORAS EXTRA" in concepto or "JORNAL" in concepto:
            return "Egreso", "Sueldos y Jornales"
        return "Egreso", "Proveedores AA"
    # "Otros movimientos de bancos y carteras" (EXT): salidas de caja sin proveedor. Las que Tango
    # llama "DEBITO Y GASTOS BANCARIOS" son montos redondos ($20 M, $10 M, $6 M): no son
    # comisiones, es plata que sale de la caja de AA (¿a A? ¿retiros?). Van a Otros hasta que la
    # empresa diga qué son (pregunta abierta); no a Gastos Bancarios.
    return "Egreso", "Otros"


def leer(ruta, desde):
    ws = openpyxl.load_workbook(ruta, read_only=False).active
    enc = [c.value for c in ws[1]]
    filas = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(enc, r))
        f = d.get("Fecha") or d.get("Fecha de emisión")
        if not d.get("Tipo") or not hasattr(f, "date"):
            continue
        if f.date() < desde:
            continue
        total = float(d.get("Total (cte)") or 0)
        if not total:
            continue
        tipo, cat = clasificar(d)
        importe = total if tipo == "Ingreso" else -total       # egresos en negativo, como en el extracto
        quien = str(d.get("Desc. relacionado") or "").strip()
        detalle = (quien + " · " if quien else "") + str(d.get("Concepto") or "").strip()
        filas.append(OrderedDict([
            ("ID", None), ("Fecha", f.date()), ("Empresa", "AA"), ("Tipo", tipo), ("Categoria", cat),
            ("Concepto / Detalle", detalle[:120]), ("Importe", round(importe, 2)), ("Medio de Pago", "Efectivo"),
            ("Banco / Cuenta", "Caja AA"), ("Origen", MARCA + " · movimientos tesorería · " + os.path.basename(ruta)),
            ("Estado", "Real"), ("Referencia", (str(d.get("Tipo")) + " " + str(d.get("Comprobante") or "").strip()).strip()),
            ("Semana (lunes)", None), ("Observaciones", "Tango Live · Tesorería · comprobante de AA"),
        ]))
    filas.sort(key=lambda x: (x["Fecha"], x["Referencia"]))
    for i, x in enumerate(filas, 1):
        x["ID"] = i
    return filas


def escribir_para_pegar(filas, carpeta, hoy):
    ruta = os.path.join(carpeta, "para_pegar_tesoreria_aa_%s.xlsx" % hoy.isoformat())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    ws.append(ENC_MOV)
    for x in filas:
        ws.append([None if c in COLS_CON_FORMULA else x.get(c) for c in ENC_MOV])
    for fila in ws.iter_rows(min_row=2):
        fila[1].number_format = "dd/mm/yyyy"
        fila[6].number_format = "#,##0.00"
    wb.save(ruta)
    return ruta


def resumen(filas, ruta, hoy, desde):
    por_mes = defaultdict(lambda: defaultdict(float))
    for x in filas:
        por_mes[x["Fecha"].strftime("%Y-%m")][x["Categoria"]] += x["Importe"]
    L = ["# Tesorería AA (efectivo) → Movimientos · %s" % hoy.strftime("%d/%m/%Y"), "",
         "Movimientos cargados: %d (desde %s)" % (len(filas), desde.strftime("%d/%m/%Y")), "",
         "| Mes | Cobranza AA | Proveedores AA | Sueldos | Otros (salidas de caja) | Interno (a A) | Neto |", "|---|---|---|---|---|---|---|"]
    for mes in sorted(por_mes):
        c = por_mes[mes]
        neto = sum(v for k, v in c.items() if k != "Transferencia Interna")
        L.append("| %s | %s | %s | %s | %s | %s | %s |" % (mes, _m(c.get("Cobranza AA", 0)), _m(-c.get("Proveedores AA", 0)), _m(-c.get("Sueldos y Jornales", 0)),
                 _m(-c.get("Otros", 0)), _m(-c.get("Transferencia Interna", 0)), _m(neto)))
    L += ["", "Reglas: recibos → Cobranza AA · órdenes de pago → Proveedores AA (sueldos si el concepto lo dice; a NAVAR S.A. → interno) · otros movimientos → Otros (salidas de caja) · anulaciones restan.",
          "", "## Archivo generado", "- `%s`: solapa Movimientos, Origen '%s' (el importador pisa lo que cargó la vez anterior con esa marca)." % (ruta, MARCA)]
    return "\n".join(L)


def main():
    ap = argparse.ArgumentParser(description="Movimientos de tesorería de AA (Tango) -> solapa Movimientos")
    ap.add_argument("--archivo", required=True)
    ap.add_argument("--cliente", default="navar")
    ap.add_argument("--hoy", default=None)
    ap.add_argument("--desde", default=DESDE.isoformat())
    a = ap.parse_args()
    hoy = datetime.date.fromisoformat(a.hoy) if a.hoy else datetime.date.today()
    desde = datetime.date.fromisoformat(a.desde)
    filas = leer(a.archivo, desde)
    if not filas:
        sys.exit("el export no tiene movimientos desde %s" % desde)
    carpeta = os.path.dirname(os.path.abspath(a.archivo))
    ruta = escribir_para_pegar(filas, carpeta, hoy)
    md = resumen(filas, ruta, hoy, desde)
    with io.open(os.path.join(carpeta, "resumen_tesoreria_aa_%s.md" % hoy.isoformat()), "w", encoding="utf-8") as f:
        f.write(md + "\n")
    print(md)


if __name__ == "__main__":
    main()
