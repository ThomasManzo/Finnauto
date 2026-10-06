# -*- coding: utf-8 -*-
"""
lector/cruce_semanal — el cruce banco ↔ Tango de todos los miércoles, para la administración.

PARA QUÉ
    La administración no tiene acceso a los bancos: carga los débitos y créditos cuando le llegan los
    resúmenes, a fin de mes. Nosotros tenemos los extractos todos los días. Una vez por semana este
    programa corre el cruce (lector/cruce.py) del mes anterior y del mes en curso y deja en Drive:
      · un Excel simple para la administración (no el Excel de trabajo del cruce), y
      · un .json con lo que tiene que decir el mail (lo manda cruce_semanal.gs desde la Sheet).
    No toca la Sheet.

QUÉ DICE
    1. Por banco: hasta qué fecha llega su extracto y cuánto concilia cada mes. Un banco sin extracto
       nuevo se ve como "extracto hasta tal fecha", no como diferencias.
    2. Lo que el banco debitó o acreditó hace más de 7 días (--dias) y no está en Tango, de mayor a
       menor. Los gastos e impuestos bancarios chicos van como un total por banco.
    3. Posibles errores de carga: fechas imposibles y dígitos cambiados.
    El Excel trae además el DETALLE de todo lo que no cruza, de los dos lados, con lo necesario para
    buscarlo (fecha, banco, concepto, referencia, comprobante de Tango, CUIT, importe).

DE DÓNDE TOMA LOS DATOS (todo de NAVAR - Datos en el Drive montado)
    _para la Sheet/para_pegar_bancos_<fecha>.xlsx (el más nuevo) · Tesoreria A detalle/ · Cheques/

USO
    python lector/cruce_semanal.py --cliente navar                 (lo que corre la tarea del miércoles)
    python lector/cruce_semanal.py --cliente navar --hoy 2026-10-07 --destino <carpeta>   (para probar)
"""

import io
import os
import sys
import glob
import json
import argparse
import datetime

BASE_REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_REPO not in sys.path:
    sys.path.insert(0, BASE_REPO)

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from lector import cruce as cz
from lector import cuotas_prestamos

DIAS_PENDIENTE = 7          # lo más nuevo se espera: la administración carga con días de atraso
CUANTOS_EN_EL_MAIL = 15     # el resto va en el Excel


def _mes(fecha):
    return datetime.date(fecha.year, fecha.month, 1)


def _mes_anterior(fecha):
    primero = _mes(fecha)
    return _mes(primero - datetime.timedelta(days=1))


def _limites(primero):
    ultimo = datetime.date(primero.year + (primero.month == 12), primero.month % 12 + 1, 1) - datetime.timedelta(days=1)
    return primero, ultimo


def _mas_nuevo(patron):
    archivos = [p for p in glob.glob(patron) if not os.path.basename(p).startswith("~$")
                and not p.endswith(".parte.xlsx")]
    return max(archivos, key=os.path.getmtime) if archivos else None


def _nombre_banco(cuenta):
    """'GALICIA 0005459-5 070-1' → 'Galicia 0005459-5 070-1' (más fácil de leer en el mail)."""
    primera, _, resto = cuenta.partition(" ")
    nombres = {"GALICIA": "Galicia", "MACRO": "Macro", "BBVA": "BBVA", "NACION": "Nación", "CORRIENTES": "Corrientes"}
    return (nombres.get(primera, primera.capitalize()) + " " + resto).strip()


def armar(drive, cliente="navar", hoy=None, dias=DIAS_PENDIENTE):
    """Corre el cruce de los dos meses y devuelve todo lo que van a mostrar el Excel y el mail."""
    hoy = hoy or datetime.date.today()
    sheet = _mas_nuevo(os.path.join(drive, "_para la Sheet", "para_pegar_bancos_*.xlsx"))
    if not sheet:
        raise SystemExit("no encontré ningún para_pegar_bancos en %s" % os.path.join(drive, "_para la Sheet"))
    tango = os.path.join(drive, "Tesoreria A detalle")
    cheques = os.path.join(drive, "Cheques")
    cfg = cz.cargar_config(cliente)
    meses = [_mes_anterior(hoy), _mes(hoy)]
    desde = _limites(meses[0])[0] - datetime.timedelta(days=cz.MARGEN_CARGA)
    hasta = _limites(meses[1])[1] + datetime.timedelta(days=cz.MARGEN_CARGA)
    banco, ultima, sin_par = cz.leer_banco(sheet, cfg, desde, hasta)
    renglones, archivo_tango = cz.leer_tango(tango, cfg)
    lista_cheques, sin_cuenta = cz.leer_cheques(cheques if os.path.isdir(cheques) else None, renglones)

    informes = []
    for primero in meses:
        # cada mes trabaja con sus propias copias: el cruce les agrega marcas a los renglones
        b = [dict(x) for x in banco]
        t = [dict(x) for x in renglones]
        ch = [dict(x) for x in lista_cheques]
        d, h = _limites(primero)
        informes.append(cz.armar_informe(b, t, cfg, d, h, ultima, hoy, sin_par, ch, sin_cuenta))

    # --- 1. por banco: hasta dónde llega el extracto y cuánto concilia cada mes
    extracto_hasta = {}
    for x in banco:
        extracto_hasta[x["cuenta"]] = max(extracto_hasta.get(x["cuenta"], x["fecha"]), x["fecha"])
    cuentas = sorted({r["cuenta"] for inf in informes for r in inf["resumen"]} | set(extracto_hasta),
                     key=lambda c: list(cfg["cuentas"]).index(c) if c in cfg["cuentas"] else 99)
    bancos = []
    for cta in cuentas:
        fila = {"cuenta": _nombre_banco(cta), "extracto_hasta": None, "dias_sin_extracto": None, "conciliado": []}
        if cta in extracto_hasta:
            fila["extracto_hasta"] = extracto_hasta[cta].isoformat()
            fila["dias_sin_extracto"] = (hoy - extracto_hasta[cta]).days
        for inf in informes:
            r = next((r for r in inf["resumen"] if r["cuenta"] == cta), None)
            fila["conciliado"].append(round(r["pct"], 4) if r and r["n"] else None)
        bancos.append(fila)

    # --- 2. en el banco y no en Tango (hace más de `dias`), y gastos como total por banco
    limite = hoy - datetime.timedelta(days=dias)
    falta = [b for inf in informes for b in inf["solo_banco"] if b["fecha"] <= limite]
    falta.sort(key=lambda b: (-abs(b["c"]), b["fecha"]))
    # cuotas de préstamos: el desglose de la tabla del banco (lo que la administración necesita para
    # cargarlas en Tango: capital, interés, IVA, percepción)
    tabla = cuotas_prestamos.leer(os.path.join(drive, "Deuda bancaria", cuotas_prestamos.ARCHIVO))
    for b in falta:
        c = cuotas_prestamos.para_debito(tabla, b["cuenta"].split()[0], b["fecha"], b["c"] / 100.0) if b["c"] < 0 else None
        b["desglose"] = cuotas_prestamos.desglose(c) if c else ""
    gastos = {}
    for inf in informes:
        for b in inf["gastos"]:
            if b["fecha"] > limite:
                continue
            g = gastos.setdefault(b["cuenta"], [0, 0])
            g[0] += 1
            g[1] += b["c"]

    # --- 3. posibles errores de carga (sin repetir: las fechas imposibles salen en los dos meses)
    errores, vistos = [], set()
    for inf in informes:
        for r in inf["revisar"]:
            if r["lado"] == "Tango" and r["motivo"].startswith("fecha imposible"):
                clave = (r["tipo"], r["comprobante"], r["c"])
                if clave not in vistos:
                    vistos.add(clave)
                    errores.append({"texto": "En Tango, %s %s por %s: %s" % (
                        r["tipo"], r["comprobante"], _plata(r["c"]), r["motivo"].strip())})
        for b in inf["solo_banco"]:
            if "tipeo" in (b.get("pista") or ""):
                errores.append({"texto": "%s · %s · el banco movió %s; %s" % (
                    b["fecha"].strftime("%d/%m"), _nombre_banco(b["cuenta"]), _plata(b["c"]),
                    b["pista"].replace("¿error de tipeo? ", "¿error de tipeo? "))})

    # --- detalle completo de lo que no cruza (los dos lados), para buscar cada movimiento
    detalle = []
    for inf in informes:
        for b in inf["solo_banco"]:
            detalle.append({"lado": "En el banco, falta en Tango", "fecha": b["fecha"], "cuenta": _nombre_banco(b["cuenta"]),
                            "tipo": b["categoria"], "concepto": b["texto"], "referencia": b.get("ref") or "",
                            "comprobante": "", "contraparte": b.get("contraparte", ""), "cuit": cz._cuit_lindo(b["cuit"]),
                            "importe": b["c"], "nota": b.get("pista", "")})
        for b in inf["gastos"]:
            detalle.append({"lado": "En el banco, falta en Tango (gasto/impuesto bancario)", "fecha": b["fecha"],
                            "cuenta": _nombre_banco(b["cuenta"]), "tipo": b["categoria"], "concepto": b["texto"],
                            "referencia": b.get("ref") or "", "comprobante": "", "contraparte": "", "cuit": "",
                            "importe": b["c"], "nota": ""})
        for t in inf["solo_tango"]:
            if t.get("anulado"):
                continue
            detalle.append({"lado": "En Tango, no aparece en el banco", "fecha": t["fecha"], "cuenta": t["desc_cuenta"],
                            "tipo": t["tipo"], "concepto": t["texto"], "referencia": "",
                            "comprobante": ("%s %s" % (t["tipo"], t["comprobante"])).strip(),
                            "contraparte": t["contraparte"], "cuit": cz._cuit_lindo(t["cuit"]), "importe": t["c"],
                            "nota": t["estado"]})
    detalle.sort(key=lambda x: (x["lado"], x["cuenta"], x["fecha"]))

    return {
        "hoy": hoy, "meses": meses, "bancos": bancos, "falta": falta, "dias": dias,
        "gastos": {_nombre_banco(k): v for k, v in gastos.items()}, "errores": errores, "detalle": detalle,
        "control_ok": all(inf["control_ok"] for inf in informes),
        "fecha_de_emision": any(inf.get("fecha_de_emision") for inf in informes),
        "fuentes": [os.path.basename(sheet), os.path.basename(archivo_tango)],
    }


def _plata(centavos):
    v = centavos / 100.0
    s = format(abs(v), ",.2f").replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-$" if v < 0 else "$") + s


MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]


def escribir_excel(d, ruta):
    """El Excel para la administración: cuatro solapas simples, sin la jerga del cruce."""
    wb = openpyxl.Workbook()
    negrita, encabezado = Font(bold=True), PatternFill("solid", fgColor="35506B")

    def hoja(titulo, columnas, filas, anchos, formatos=None):
        ws = wb.create_sheet(titulo)
        ws.append(columnas)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = encabezado
        for f in filas:
            ws.append(f)
        for j, (col, ancho) in enumerate(zip(columnas, anchos), start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = ancho
            fmt = (formatos or {}).get(col)
            if fmt:
                for (celda,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    celda.number_format = fmt
        if filas:
            ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    nombres = [MESES[m.month - 1].capitalize() for m in d["meses"]]
    ws = hoja("Por banco", ["Cuenta", "Extracto hasta", "Días sin extracto nuevo"] + ["Concilia %s" % n for n in nombres],
              [[b["cuenta"], datetime.date.fromisoformat(b["extracto_hasta"]) if b["extracto_hasta"] else None,
                b["dias_sin_extracto"]] + b["conciliado"] for b in d["bancos"]],
              [30, 15, 12, 16, 16], {"Extracto hasta": "dd/mm/yyyy", "Concilia " + nombres[0]: "0%",
                                     "Concilia " + nombres[1]: "0%"})
    ws.append([])
    ws.append(["'Concilia' = parte de lo que se movió en esa cuenta que encontramos igual en el banco y en Tango."])

    hoja("Falta cargar en Tango", ["Fecha", "Banco", "Concepto (como lo escribe el banco)", "Referencia", "CUIT", "Importe",
                                   "Desglose (tabla del banco)", "Nota"],
         [[b["fecha"], _nombre_banco(b["cuenta"]), b["texto"], b.get("ref") or "", cz._cuit_lindo(b["cuit"]),
           b["c"] / 100.0, b.get("desglose", ""), b.get("pista", "")] for b in d["falta"]],
         [12, 26, 50, 16, 16, 18, 70, 50], {"Fecha": "dd/mm/yyyy", "Importe": "#,##0.00"})

    hoja("Posibles errores", ["Qué vemos"], [[e["texto"]] for e in d["errores"]], [120])

    hoja("Detalle de lo que no cruza", ["Dónde está", "Fecha", "Banco / cuenta", "Tipo", "Concepto / leyenda", "Referencia",
                                        "Comprobante Tango", "Contraparte", "CUIT", "Importe", "Nota"],
         [[x["lado"], x["fecha"], x["cuenta"], x["tipo"], x["concepto"], x["referencia"], x["comprobante"],
           x["contraparte"], x["cuit"], x["importe"] / 100.0, x["nota"]] for x in d["detalle"]],
         [34, 12, 26, 16, 45, 14, 22, 26, 15, 18, 40], {"Fecha": "dd/mm/yyyy", "Importe": "#,##0.00"})
    tmp = ruta + ".parte.xlsx"
    wb.save(tmp)
    os.replace(tmp, ruta)            # aparece entero en Drive
    return ruta


def para_el_mail(d, archivo_excel):
    """Lo que necesita cruce_semanal.gs para armar el mail (fechas como texto, plata en pesos)."""
    falta = d["falta"]
    return {
        "fecha": d["hoy"].isoformat(),
        "generado": datetime.datetime.now().astimezone().replace(microsecond=0).isoformat(),
        "meses": [MESES[m.month - 1] for m in d["meses"]],
        "dias": d["dias"],
        "bancos": d["bancos"],
        "falta_total": {"cantidad": len(falta), "importe": round(sum(abs(b["c"]) for b in falta) / 100.0, 2)},
        "falta": [{"fecha": b["fecha"].isoformat(), "banco": _nombre_banco(b["cuenta"]), "concepto": b["texto"][:70],
                   "importe": b["c"] / 100.0, "desglose": b.get("desglose", "")} for b in falta[:CUANTOS_EN_EL_MAIL]],
        "gastos": [{"banco": k, "cantidad": v[0], "importe": v[1] / 100.0} for k, v in sorted(d["gastos"].items())],
        "errores": [e["texto"] for e in d["errores"]],
        "archivo": os.path.basename(archivo_excel),
        "control_ok": d["control_ok"],
        "aviso": ("el detalle de Tango no trae la columna 'Fecha' (la del movimiento)" if d["fecha_de_emision"] else ""),
    }


def main(argv=None):
    ap = argparse.ArgumentParser(description="Cruce banco ↔ Tango semanal (Excel + datos del mail en Drive)")
    ap.add_argument("--cliente", default="navar")
    ap.add_argument("--hoy", default=None)
    ap.add_argument("--dias", type=int, default=DIAS_PENDIENTE)
    ap.add_argument("--drive", default=None, help="carpeta NAVAR - Datos (default: la del Drive montado)")
    ap.add_argument("--destino", default=None, help="dónde dejar Excel y .json (default: <drive>/Cruce)")
    a = ap.parse_args(argv)
    if a.drive:
        drive = a.drive
    else:
        from ingestas.drive_local import carpeta_datos
        drive = carpeta_datos()
    if not drive or not os.path.isdir(drive):
        raise SystemExit("no encuentro la carpeta de Drive; pasá --drive o fijá FINAUTO_DRIVE")
    hoy = datetime.date.fromisoformat(a.hoy) if a.hoy else datetime.date.today()
    destino = a.destino or os.path.join(drive, "Cruce")
    os.makedirs(destino, exist_ok=True)
    d = armar(drive, a.cliente, hoy, a.dias)
    excel = escribir_excel(d, os.path.join(destino, "Cruce banco-Tango %s.xlsx" % hoy.isoformat()))
    datos = para_el_mail(d, excel)
    ruta_json = os.path.join(destino, "cruce_semanal_%s.json" % hoy.isoformat())
    with io.open(ruta_json + ".tmp", "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=1)
    os.replace(ruta_json + ".tmp", ruta_json)
    print("%s  cruce semanal: %d para cargar en Tango (%s), %d posibles errores, control %s → %s" % (
        datetime.datetime.now().strftime("%H:%M:%S"), datos["falta_total"]["cantidad"],
        _plata(int(round(datos["falta_total"]["importe"] * 100))), len(datos["errores"]),
        "OK" if datos["control_ok"] else "NO DA", destino))
    return 0 if datos["control_ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
