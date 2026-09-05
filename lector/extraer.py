# -*- coding: utf-8 -*-
"""
lector.extraer — la planilla del cliente, convertida en contrato.

Ultimo eslabon del circuito de onboarding:

    planilla -> lector/planilla.py -> lector/borrador.py -> [reunion]
             -> mapeo.json + catalogo.json -> ESTE MODULO -> contrato.json
             -> simulador / proyeccion / consejo / auditoria

A partir de aca el cliente nuevo es indistinguible de MAGA para todo el resto
del sistema: las herramientas leen el contrato y no saben de que planilla salio.

NO INTERPRETA NADA
------------------
Este modulo no decide que columna es cual: eso ya lo decidio una persona en la
reunion y quedo escrito en mapeo.json. Aca solo se lee lo que ese archivo dice.
Si el mapeo esta mal, la salida esta mal — y por eso se avisa fuerte cuando algo
no cierra (fechas ilegibles, importes en cero, filas salteadas) en vez de
producir un contrato prolijo con numeros inventados.

Uso:
    python lector/extraer.py --archivo "Cash.xlsx" --cliente panaderia
    python lector/extraer.py --archivo "Cash.xlsx" --cliente panaderia --salida c.json
"""

import io
import os
import sys
import json
import argparse
import datetime

BASE_REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_REPO not in sys.path:
    sys.path.insert(0, BASE_REPO)

from lector.planilla import leer_hoja, _vacio
from lector.borrador import _num, _fmt, FALTA

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl. Instalalo con:  pip install openpyxl")
    sys.exit(1)


def cargar_mapeo(cliente):
    ruta = os.path.join(BASE_REPO, "clientes", cliente, "mapeo.json")
    if not os.path.exists(ruta):
        raise IOError("No existe %s.\nCorrer primero:\n"
                      "  python lector/borrador.py --archivo <planilla> --salida b.json\n"
                      "  (completar b.json con el cliente)\n"
                      "  python lector/borrador.py --aplicar b.json --cliente %s"
                      % (ruta, cliente))
    with io.open(ruta, encoding="utf-8") as f:
        return json.load(f)


def _a_iso(v):
    """Cualquier fecha razonable -> aaaa-mm-dd. None si no se puede leer."""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = str(v or "").strip()
    if not s:
        return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%m/%d/%Y"):
        try:
            return datetime.datetime.strptime(s[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _col_idx(encabezados, titulo):
    if not titulo or titulo.startswith(FALTA):
        return None
    t = titulo.strip().upper()
    for i, e in enumerate(encabezados):
        if str(e or "").strip().upper() == t:
            return i
    return None


def extraer(archivo, mapeo):
    hoja = mapeo["hoja"]
    fila_enc = int(mapeo["fila_encabezado"])
    cols = mapeo["columnas"]

    wb = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    if hoja not in wb.sheetnames:
        wb.close()
        raise IOError("La planilla no tiene la hoja \"%s\". Tiene: %s"
                      % (hoja, ", ".join(wb.sheetnames)))
    filas = leer_hoja(wb[hoja])
    wb.close()

    encabezados = filas[fila_enc - 1] if fila_enc - 1 < len(filas) else []
    idx = dict((k, _col_idx(encabezados, v)) for k, v in cols.items())

    faltan = [k for k in ("fecha", "importe") if idx.get(k) is None]
    if faltan:
        raise IOError("El mapeo apunta a columnas que no existen en la planilla: %s.\n"
                      "Encabezados de la fila %d: %s"
                      % (", ".join(faltan), fila_enc,
                         ", ".join(str(e) for e in encabezados if e)))

    signo = mapeo.get("_signo", {})
    modo = str(signo.get("modo", "")).split(":")[0].strip().strip("'\"")
    i_ing = _col_idx(encabezados, signo.get("columna_ingresos", ""))

    movs, avisos = [], {"sin_fecha": 0, "sin_importe": 0, "sin_tipo": 0, "leidas": 0}
    for f in filas[fila_enc:]:
        if not any(not _vacio(v) for v in f):
            continue
        avisos["leidas"] += 1

        fecha = _a_iso(f[idx["fecha"]] if idx["fecha"] < len(f) else None)
        if not fecha:
            avisos["sin_fecha"] += 1
            continue

        bruto = f[idx["importe"]] if idx["importe"] < len(f) else None
        imp = _num(bruto)
        ingreso = 0.0
        if modo == "columnas" and i_ing is not None and i_ing < len(f):
            ingreso = _num(f[i_ing])

        if modo == "signo":
            es_egreso = imp < 0
            imp = abs(imp)
        elif modo == "columnas":
            es_egreso = abs(imp) >= abs(ingreso)
            imp = abs(imp) if es_egreso else abs(ingreso)
        else:                       # 'todos_egresos' o sin definir
            es_egreso = True
            imp = abs(imp)

        if not imp:
            avisos["sin_importe"] += 1
            continue

        tipo = ""
        if idx.get("tipo") is not None and idx["tipo"] < len(f):
            tipo = str(f[idx["tipo"]] or "").strip()
        if not tipo:
            avisos["sin_tipo"] += 1

        concepto = ""
        if idx.get("concepto") is not None and idx["concepto"] < len(f):
            concepto = str(f[idx["concepto"]] or "").strip()

        movs.append({"fecha": fecha, "tipo": tipo, "concepto": concepto,
                     "importe": imp, "egreso": es_egreso})

    movs.sort(key=lambda m: m["fecha"])
    return movs, avisos


def contrato(movs, cliente, archivo):
    egr = [m for m in movs if m.get("egreso")]
    ing = [m for m in movs if not m.get("egreso")]
    return {
        "version": "1.2",
        "_origen": {"cliente": cliente, "archivo": os.path.basename(archivo),
                    "generado": datetime.datetime.now().isoformat(timespec="seconds"),
                    "_nota": "Generado por lector/extraer.py a partir del mapeo "
                             "acordado con el cliente. No se interpreto nada: si "
                             "el mapeo esta mal, estos numeros estan mal."},
        "caja_hoy": None,
        "movimientos": [{"fecha": m["fecha"], "tipo": m["tipo"],
                         "concepto": m["concepto"], "importe": m["importe"]}
                        for m in egr],
        "cobros_previstos": [{"fecha": m["fecha"], "concepto": m["concepto"],
                              "importe": m["importe"], "naturaleza": "VARIABLE"}
                             for m in ing],
    }


def imprimir(movs, avisos, c):
    L = 74
    print("=" * L)
    print("  PLANILLA -> CONTRATO")
    print("=" * L)
    egr = c["movimientos"]
    ing = c["cobros_previstos"]
    print("  Filas leidas          : %d" % avisos["leidas"])
    print("  Movimientos extraidos : %d  (%d egresos + %d ingresos)"
          % (len(egr) + len(ing), len(egr), len(ing)))
    if egr:
        print("  Rango de fechas       : %s al %s"
              % (min(m["fecha"] for m in egr), max(m["fecha"] for m in egr)))
        print("  Total de egresos      : %s" % _fmt(sum(m["importe"] for m in egr)))

    perdidas = avisos["sin_fecha"] + avisos["sin_importe"]
    if perdidas or avisos["sin_tipo"]:
        print("")
        print("  LO QUE NO ENTRO (revisar antes de confiar en el numero)")
        if avisos["sin_fecha"]:
            print("   . %d fila(s) sin fecha legible" % avisos["sin_fecha"])
        if avisos["sin_importe"]:
            print("   . %d fila(s) con importe vacio o en cero" % avisos["sin_importe"])
        if avisos["sin_tipo"]:
            print("   . %d movimiento(s) sin categoria: el motor no los va a poder "
                  "clasificar" % avisos["sin_tipo"])
        if avisos["leidas"]:
            print("   . se perdio el %.0f%% de las filas"
                  % (100.0 * perdidas / avisos["leidas"]))


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    ap = argparse.ArgumentParser(description="Planilla del cliente -> contrato.json")
    ap.add_argument("--archivo", required=True)
    ap.add_argument("--cliente", required=True)
    ap.add_argument("--salida")
    args = ap.parse_args()

    mapeo = cargar_mapeo(args.cliente)
    movs, avisos = extraer(args.archivo, mapeo)
    c = contrato(movs, args.cliente, args.archivo)
    imprimir(movs, avisos, c)

    salida = args.salida or os.path.join(
        BASE_REPO, "clientes", args.cliente,
        "contrato_%s.json" % datetime.date.today().isoformat())
    with io.open(salida, "w", encoding="utf-8") as f:
        f.write(json.dumps(c, ensure_ascii=False, indent=2))
        f.write(u"\n")
    print("\n  Contrato: %s" % salida)
    print("  Siguiente: python simulador/proyeccion.py --contrato %s --horizonte 45"
          % os.path.basename(salida))


if __name__ == "__main__":
    main()
