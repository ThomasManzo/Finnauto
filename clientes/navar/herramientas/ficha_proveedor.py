# -*- coding: utf-8 -*-
"""
ficha_proveedor — todo lo que sabemos de un proveedor (o cliente), junto y cruzado.

Junta lo que ya está bajado, sin pedirle nada nuevo a nadie:
  · facturas pendientes            (Cuentas a pagar / a cobrar de Tango)
  · órdenes de pago y recibos      (movimientos de tesorería de Tango, A y AA)
  · cheques propios entregados     (cartera de cheques de Tango)
  · débitos en el banco            (extracto: se busca el débito que corresponde a cada cheque,
                                    mismo importe dentro de los 7 días posteriores a la fecha del
                                    cheque; así se ve si un cheque que Tango muestra "Al Cobro"
                                    en realidad ya lo pagó el banco)

Sirve para contestar "¿cuánto le debemos de verdad?" cuando Tango no concilia los cheques.

USO
    python clientes/navar/herramientas/ficha_proveedor.py --buscar COPETEGLA WIELIKI
    (deja un xlsx y un resumen .md en privado/salidas/)
"""

import os
import io
import sys
import glob
import argparse
import datetime
import unicodedata
from collections import OrderedDict

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE_REPO = os.path.abspath(os.path.join(AQUI, "..", "..", ".."))
sys.path.insert(0, BASE_REPO)
from ingestas.drive_local import carpeta_datos

DIAS_CANJE = 7          # un cheque se debita entre 0 y 7 días después de su fecha (48 hs de canje)


def _n(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return "".join(ch for ch in s.upper() if ch.isalnum())


def _m(v):
    return ("-" if v < 0 else "") + "$" + format(int(round(abs(v or 0))), ",d").replace(",", ".")


def _filas(ruta, hoja=None):
    wb = openpyxl.load_workbook(ruta)
    ws = wb[hoja] if hoja else wb.active
    enc = [c.value for c in ws[1]]
    return [dict(zip(enc, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(x is not None for x in r)]


def _ultimo(patron):
    c = sorted(glob.glob(patron))
    return c[-1] if c else None


def coincide(fila, claves):
    txt = _n(" ".join(str(v) for v in fila.values() if isinstance(v, str)))
    return any(k in txt for k in claves)


def juntar(claves, drive):
    out = {"facturas": [], "pagos": [], "cheques": []}
    # facturas pendientes: el archivo más nuevo de cada empresa y lista (no alcanza con ordenar
    # por nombre: "AA pagos" ordena después de "A pagos" y se perdía la empresa A)
    ultimos = {}
    for carpeta in ("Cuentas a pagar", "Cuentas a cobrar"):
        for f in glob.glob(os.path.join(drive, carpeta, "*.xlsx")):
            if "para_pegar" in f:
                continue
            base = os.path.basename(f)
            clave = (carpeta, base.split()[0].upper())          # ("Cuentas a pagar", "A"/"AA")
            if clave not in ultimos or os.path.getmtime(f) > os.path.getmtime(ultimos[clave]):
                ultimos[clave] = f
    for f in ultimos.values():
        for r in _filas(f):
            if coincide(r, claves):
                out["facturas"].append(OrderedDict([
                    ("Empresa", "AA" if os.path.basename(f).upper().startswith("AA ") else "A"),
                    ("Tipo", "a pagar" if "pagar" in f else "a cobrar"),
                    ("Comprobante", "%s %s" % (r.get("Tipo de comprobante") or r.get("Tipo comprobante") or "", r.get("Nro. comprobante") or "")),
                    ("Emisión", r.get("Fecha de emisión")), ("Vencimiento", r.get("Fecha de vencimiento")),
                    ("Total", r.get("Total al vencimiento (CTE)") or r.get("Importe al Vencimiento (CTE)")),
                    ("Pendiente", r.get("Total pendiente (CTE)") or r.get("Importe Pendiente (CTE)")),
                ]))
    # pagos y cobros (tesorería de Tango: A y AA)
    for f, empresa in [(_ultimo(os.path.join(BASE_REPO, "clientes", "navar", "privado", "tango", "*", "A movimientos consulta.xlsx")), "A"),
                       (_ultimo(os.path.join(drive, "Tesoreria AA", "AA Movimientos*.xlsx")), "AA")]:
        if not f:
            continue
        for r in _filas(f):
            if not coincide(r, claves):
                continue
            out["pagos"].append(OrderedDict([
                ("Empresa", empresa), ("Fecha", r.get("Fecha")), ("Tipo", r.get("Tipo")),
                ("Comprobante", (r.get("Comprobante") or "").strip()), ("Clase", r.get("Clase")),
                ("Importe", r.get("Total (cte)") if r.get("Total (cte)") is not None else r.get("Total (ext)")),
                ("Concepto", r.get("Concepto")), ("Quién", r.get("Desc. relacionado")),
            ]))
    # cheques propios entregados
    f = _ultimo(os.path.join(drive, "Cheques", "* cheques propios *.xlsx"))
    if f:
        for r in _filas(f):
            if coincide(r, claves):
                out["cheques"].append(OrderedDict([
                    ("Nro. cheque", r.get("Nro. de cheque")), ("Banco", r.get("Banco") or r.get("Nombre de banco")),
                    ("Emisión", r.get("Fecha de emisión")), ("Fecha del cheque", r.get("Fecha del cheque")),
                    ("Importe", r.get("Importe mon. cta.")), ("Estado en Tango", r.get("Estado")),
                ]))
    return out


def cruzar_con_banco(cheques, drive):
    """Para cada cheque, busca en el extracto un débito del mismo importe dentro de los 7 días."""
    f = _ultimo(os.path.join(drive, "_para la Sheet", "para_pegar_bancos_*.xlsx"))
    if not f:
        return "sin extracto para cruzar"
    movs = [r for r in _filas(f, "Movimientos") if r.get("Importe") and float(r["Importe"]) < 0]
    desde = min((r["Fecha"] for r in movs if r.get("Fecha")), default=None)
    for ch in cheques:
        fch, imp = ch["Fecha del cheque"], float(ch["Importe"] or 0)
        if not hasattr(fch, "date") or not imp:
            continue
        if desde and fch < desde:
            ch["Débito en el banco"] = "el extracto arranca el %s: no se puede ver" % desde.strftime("%d/%m/%Y")
            continue
        halla = [m for m in movs if m.get("Fecha") and 0 <= (m["Fecha"] - fch).days <= DIAS_CANJE
                 and abs(abs(float(m["Importe"])) - imp) <= max(1000.0, imp * 0.001)]
        ch["Débito en el banco"] = ("%s · %s · %s" % (halla[0]["Fecha"].strftime("%d/%m/%Y"), halla[0]["Banco / Cuenta"],
                                                     str(halla[0]["Concepto / Detalle"])[:40])) if halla else "NO aparece"
    return "cruzado contra %s" % os.path.basename(f)


def escribir(fichas, salida, hoy):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, d in fichas.items():
        for hoja, filas in d.items():
            ws = wb.create_sheet(("%s %s" % (nombre[:14], hoja))[:31])
            if not filas:
                ws.append(["(sin datos)"])
                continue
            ws.append(list(filas[0].keys()))
            for r in filas:
                ws.append(list(r.values()))
            for fila in ws.iter_rows(min_row=2):
                for c in fila:
                    if isinstance(c.value, datetime.datetime):
                        c.number_format = "dd/mm/yyyy"
                    elif isinstance(c.value, (int, float)):
                        c.number_format = "#,##0.00"
    wb.save(salida)
    return salida


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--buscar", nargs="+", required=True, help="parte del nombre o el código (ej. COPETEGLA COO002)")
    ap.add_argument("--hoy", default=None)
    a = ap.parse_args()
    hoy = datetime.date.fromisoformat(a.hoy) if a.hoy else datetime.date.today()
    drive = carpeta_datos()
    if not drive:
        sys.exit("no encuentro la carpeta de Drive montada")

    fichas, L = OrderedDict(), ["# Ficha de proveedor · %s" % hoy.strftime("%d/%m/%Y"), ""]
    for quien in a.buscar:
        d = juntar([_n(quien)], drive)
        nota = cruzar_con_banco(d["cheques"], drive)
        fichas[quien] = d
        pend = sum(float(x["Pendiente"] or 0) for x in d["facturas"] if x["Tipo"] == "a pagar")
        cobrar = sum(float(x["Pendiente"] or 0) for x in d["facturas"] if x["Tipo"] == "a cobrar")
        pagos26 = [x for x in d["pagos"] if hasattr(x["Fecha"], "year") and x["Fecha"].year >= 2026]
        chq = sum(float(x["Importe"] or 0) for x in d["cheques"])
        visibles = [x for x in d["cheques"] if not str(x.get("Débito en el banco", "")).startswith("el extracto")]
        debitados = [x for x in visibles if str(x.get("Débito en el banco", "")) not in ("", "NO aparece")]
        suma = lambda xs: sum(float(x["Importe"] or 0) for x in xs)
        L += ["## %s" % quien, "",
              "- Facturas pendientes de pago: **%s** (%d)" % (_m(pend), len([x for x in d["facturas"] if x["Tipo"] == "a pagar"])),
              "- Pendientes de cobro (es también cliente): %s" % _m(cobrar),
              "- Pagos registrados en Tango: %d en total, %d en 2026 por %s" % (len(d["pagos"]), len(pagos26), _m(sum(float(x["Importe"] or 0) for x in pagos26))),
              "- Cheques propios entregados que Tango muestra pendientes: **%s** (%d) · %s" % (_m(chq), len(d["cheques"]), nota),
              "  - con fecha dentro del período del extracto: %s (%d) → **con débito encontrado: %s (%d)** · sin débito: %s (%d)" % (
                  _m(suma(visibles)), len(visibles), _m(suma(debitados)), len(debitados),
                  _m(suma(visibles) - suma(debitados)), len(visibles) - len(debitados)),
              "  - anteriores al extracto (no se pueden verificar todavía): %s (%d)" % (_m(chq - suma(visibles)), len(d["cheques"]) - len(visibles)), ""]
        if pagos26:
            L += ["| Fecha | Comprobante | Importe | Concepto |", "|---|---|---|---|"]
            for x in sorted(pagos26, key=lambda x: x["Fecha"]):
                L.append("| %s | %s %s | %s | %s |" % (x["Fecha"].strftime("%d/%m/%Y"), x["Tipo"], x["Comprobante"], _m(x["Importe"]), x["Concepto"]))
            L.append("")
    salida = os.path.join(BASE_REPO, "clientes", "navar", "privado", "salidas", "ficha_proveedores_%s.xlsx" % hoy.isoformat())
    escribir(fichas, salida, hoy)
    L += ["## Archivo", "- `%s`: una solapa por proveedor y tipo (facturas · pagos · cheques)." % salida]
    md = "\n".join(L)
    with io.open(salida.replace(".xlsx", ".md"), "w", encoding="utf-8") as f:
        f.write(md + "\n")
    print(md)


if __name__ == "__main__":
    main()
