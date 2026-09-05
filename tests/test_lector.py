# -*- coding: utf-8 -*-
"""
tests/test_lector.py — el lector de planillas, contra planillas que no existen.

Se generan al vuelo, a proposito. Si los tests usaran el Cash de MAGA, estarian
midiendo si el lector entiende ESA planilla, que es justo lo contrario de lo que
tiene que hacer: entender una que ve por primera vez.

Cada planilla de prueba tiene una macana distinta de las que aparecen en el
mundo real: titulo suelto arriba, importes cargados como texto, una columna a
medio llenar, una hoja vieja vacia.

    python tests/test_lector.py
"""

import os
import sys
import tempfile
import datetime

BASE_REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_REPO not in sys.path:
    sys.path.insert(0, BASE_REPO)

from openpyxl import Workbook
from lector.planilla import analizar, preguntas, perfil_columna, _fila_encabezado

_fallos = []


def ok(cond, nombre, detalle=""):
    if cond:
        print("  [OK]   %s" % nombre)
    else:
        print("  [FALLA] %s   %s" % (nombre, detalle))
        _fallos.append(nombre)


def _guardar(wb):
    ruta = os.path.join(tempfile.mkdtemp(), "p.xlsx")
    wb.save(ruta)
    return ruta


def _hoja(a, nombre):
    return [x for x in a if x["hoja"] == nombre][0]


def _col(h, titulo):
    return [c for c in h["columnas"] if c["titulo"] == titulo][0]


# ------------------------------------------------------------ clasificacion
def test_perfil_columna():
    print("\n== Reconocer que hay en una columna ==")

    ok(perfil_columna(["01/06/2026", "15/07/2026", "3/8/2026"])["clase"] == "fecha",
       "fechas dd/mm/aaaa")
    ok(perfil_columna([datetime.date(2026, 6, 1)] * 5)["clase"] == "fecha",
       "fechas como fecha de verdad, no texto")

    ok(perfil_columna([1000, 2500.5, -300])["clase"] == "numero", "numeros")
    # La macana mas comun: importes cargados como texto con formato argentino.
    ok(perfil_columna(["$862.393,02", "$53.895,17", "$1.200,00"])["clase"] == "numero",
       "importes escritos como texto igual se reconocen")

    # Una CATEGORIA se define porque los valores SE REPITEN.
    p = perfil_columna(["Sueldos", "Proveedores", "Sueldos", "Alquiler"] * 10)
    ok(p["clase"] == "categoria", "pocos valores repetidos = categoria")
    ok(p["distintos"] == 3, "y cuenta cuantos son")
    ok([v for v, _ in p["valores"]], "guarda los valores: de ahi sale el catalogo")

    # ERROR REAL: 23 valores TODOS distintos pasaban como categoria solo por ser
    # menos de 30. Pocos no alcanza; tienen que repetirse.
    p = perfil_columna(["ROCIO ZABALA", "JAVIER PEREZ", "ANA GOMEZ"] + ["X%d" % i for i in range(20)])
    ok(p["clase"] != "categoria", "muchos valores unicos NO son una categoria",
       p["clase"])

    ok(perfil_columna([None, "", "  "])["clase"] == "vacia", "columna vacia")
    ok(perfil_columna([])["clase"] == "vacia", "columna sin filas")


def test_encabezado_no_esta_en_la_primera_fila():
    print("\n== Encontrar la tabla aunque no arranque arriba ==")
    filas = [
        ["PANADERIA LA ESPIGA SRL", None, None],
        ["Movimientos del mes", None, None],
        [None, None, None],
        ["Fecha", "Detalle", "Importe"],
        ["01/06/2026", "Molino Sur", 1000],
        ["02/06/2026", "Luz", 2000],
        ["03/06/2026", "Alquiler", 3000],
    ]
    ok(_fila_encabezado(filas) == 3, "salta el titulo suelto y encuentra la fila 4",
       str(_fila_encabezado(filas)))

    # Sin nada abajo no hay tabla: un encabezado solo no es una tabla.
    ok(_fila_encabezado([["Fecha", "Detalle", "Importe"]]) is None,
       "un encabezado sin datos debajo no cuenta como tabla")


# ------------------------------------------------------------ planilla entera
def _planilla_pyme():
    wb = Workbook()
    ws = wb.active
    ws.title = "Flujo"
    ws["A1"] = "PANADERIA LA ESPIGA S.R.L."      # titulo suelto
    for i, e in enumerate(["Fecha", "Detalle", "Rubro", "Egreso", "Ingreso", "Obs"], 1):
        ws.cell(row=3, column=i, value=e)
    rubros = ["Proveedores", "Sueldos", "Alquiler"]
    for r in range(4, 34):
        ws.cell(row=r, column=1, value="%02d/06/2026" % ((r % 28) + 1))
        ws.cell(row=r, column=2, value="detalle %d" % r)
        ws.cell(row=r, column=3, value=rubros[r % 3])
        ws.cell(row=r, column=4, value=1000 + r)
        if r % 4 == 0:
            ws.cell(row=r, column=6, value="revisar")
    wb.create_sheet("Hoja vieja")                 # vacia
    return _guardar(wb)


def test_planilla_completa():
    print("\n== Radiografia de una planilla que nunca vio ==")
    a = analizar(_planilla_pyme())
    ok(len(a) == 2, "recorre todas las hojas", str(len(a)))

    vieja = _hoja(a, "Hoja vieja")
    ok(vieja.get("vacia") is True, "detecta la hoja vacia")

    h = _hoja(a, "Flujo")
    ok(h["fila_encabezado"] == 3, "ubica el encabezado con el titulo arriba",
       str(h.get("fila_encabezado")))
    ok(h["filas_datos"] == 30, "cuenta las filas de datos", str(h["filas_datos"]))
    ok(_col(h, "Fecha")["clase"] == "fecha", "reconoce la columna de fecha")
    ok(_col(h, "Egreso")["clase"] == "numero", "reconoce la de importe")
    ok(_col(h, "Rubro")["clase"] == "categoria", "reconoce la categoria")
    ok(_col(h, "Ingreso")["clase"] == "vacia", "y la que quedo sin usar")


def test_preguntas():
    print("\n== Las preguntas para el cliente ==")
    a = analizar(_planilla_pyme())
    qs = [q for _, q in preguntas(a)]
    txt = " || ".join(qs)

    # Esta es LA pregunta: los valores de la categoria son el catalogo de tipos
    # del cliente, y lo que se puede o no postergar solo lo sabe una persona.
    ok(any("Rubro" in q and "postergar" in q for q in qs),
       "pregunta que significa cada valor de la categoria", txt[:90])
    ok(any("Proveedores" in q for q in qs),
       "y le muestra los valores concretos que encontro")

    ok(any("Ingreso" in q and "ni un dato" in q for q in qs),
       "pregunta por la columna que existe pero esta vacia")
    ok(any("Obs" in q and "%" in q for q in qs),
       "y por la que quedo a medio llenar")

    # Una hoja vacia no genera preguntas: seria ruido.
    ok(not any("Hoja vieja" in h for h, _ in preguntas(a)),
       "la hoja vacia no genera preguntas")


def test_dos_columnas_de_importe():
    print("\n== Cuando hay mas de una columna de numeros ==")
    wb = Workbook()
    ws = wb.active
    for i, e in enumerate(["Fecha", "Detalle", "Importe", "Saldo"], 1):
        ws.cell(row=1, column=i, value=e)
    for r in range(2, 22):
        ws.cell(row=r, column=1, value="%02d/06/2026" % r)
        ws.cell(row=r, column=2, value="d%d" % r)
        ws.cell(row=r, column=3, value=100 * r)
        ws.cell(row=r, column=4, value=100000 - r)
    qs = [q for _, q in preguntas(analizar(_guardar(wb)))]
    # No puede adivinar cual es el importe del movimiento y cual el saldo
    # acumulado: confundirlos daria una caja completamente inventada.
    ok(any("Importe" in q and "Saldo" in q for q in qs),
       "avisa que hay dos columnas de numeros y pregunta cual es cual", str(qs))


if __name__ == "__main__":
    print("=" * 62)
    print("  LECTOR DE PLANILLAS  ·  contra planillas desconocidas")
    print("=" * 62)
    test_perfil_columna()
    test_encabezado_no_esta_en_la_primera_fila()
    test_planilla_completa()
    test_preguntas()
    test_dos_columnas_de_importe()
    print("\n" + "=" * 62)
    if _fallos:
        print("  %d TEST(S) FALLARON: %s" % (len(_fallos), ", ".join(_fallos)))
        sys.exit(1)
    print("  TODOS LOS TESTS PASARON")
