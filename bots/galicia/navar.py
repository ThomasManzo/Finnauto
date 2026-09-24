# -*- coding: utf-8 -*-
"""Controles de NAVAR: no publicar otra cuenta ni un CSV disfrazado de Excel."""
import datetime
import os
import re
import shutil
import tempfile
from pathlib import Path

from bots.galicia.bot import BotGalicia
from nucleo.log import captura
from nucleo.utilidades import _click_robusto


def _empresa(texto):
    return re.sub(r"[\s.]", "", texto or "").upper()


class BotGaliciaNavar(BotGalicia):
    def __init__(self, cfg):
        super().__init__()
        self.cuenta = cfg["cuenta"]
        self.empresas = {_empresa(n) for n in cfg["solo_estas_empresas"]}
        self.prefijo = cfg["prefijo_archivo"]
        self.rango = None

    def capturar_empresa_activa(self, page):
        nombre = super().capturar_empresa_activa(page)
        return "NAVAR SA" if _empresa(nombre) in self.empresas else nombre

    def descubrir_empresas(self, page, timeout, excluir, solo=None):
        nombres = super().descubrir_empresas(page, timeout, excluir, solo=None)
        return [n for n in nombres if _empresa(n) in self.empresas]

    def ir_a_cuenta(self, page, timeout):
        # El motor permite seguir si no pudo leer la empresa. Acá eso no alcanza:
        # el lector atribuye todo a una cuenta fija, así que ante dudas frenamos.
        if _empresa(self.capturar_empresa_activa(page)) not in self.empresas:
            raise RuntimeError("No pude confirmar la empresa de NAVAR; revisar captura.")
        page.get_by_text("Cuentas", exact=True).click(timeout=timeout)
        page.wait_for_timeout(3000)
        captura(page, "listado_cuentas")
        # Este enlace debe confirmarse en la notebook. No recurrir al primer monto.
        cuenta = page.get_by_text(self.cuenta, exact=True)
        if cuenta.count() != 1:
            raise RuntimeError("La cuenta pedida no aparece una sola vez; revisar listado_cuentas.")
        cuenta.click(timeout=timeout)
        page.wait_for_timeout(3500)
        captura(page, "movimientos")
        if not page.get_by_text(self.cuenta, exact=True).is_visible():
            raise RuntimeError("No pude confirmar la cuenta abierta.")

    def aplicar_filtro_fechas(self, page, desde, hasta, timeout):
        self.rango = None
        # El motor no mira el False del bot original: frenamos antes de descargar.
        if not super().aplicar_filtro_fechas(page, desde, hasta, timeout):
            raise RuntimeError("No pude confirmar las fechas elegidas; no publico el extracto.")
        self.rango = (desde, hasta)
        return True

    def validar_excel(self, ruta):
        from openpyxl import load_workbook
        from lector.extractos import leer_planilla_galicia

        with open(ruta, "rb") as archivo:
            libro = load_workbook(archivo, read_only=True, data_only=True)
            try:
                texto = " ".join(str(c or "") for fila in libro.active.iter_rows(values_only=True) for c in fila)
            finally:
                libro.close()
        # Pedimos evidencia en el archivo también; si no viene, hay que revisar
        # el export real antes de decidir cómo verificarlo. No adivinamos cuentas.
        patron = r"(?<!\d)" + r"[\s-]*".join(re.escape(c) for c in re.sub(r"\D", "", self.cuenta)) + r"(?!\d)"
        if not re.search(patron, texto):
            raise RuntimeError("El Excel no permite confirmar la cuenta pedida.")
        datos = leer_planilla_galicia(ruta)
        if not datos["movimientos"]:
            raise RuntimeError("Excel sin movimientos legibles; revisar antes de marcarlo al día.")
        desde, hasta = self.rango
        for mov in datos["movimientos"]:
            fecha = datetime.date.fromisoformat(str(mov["fecha"])[:10])
            if not desde <= fecha <= hasta:
                raise RuntimeError("El Excel trae movimientos fuera del rango pedido.")
        return datos

    def descargar_csv(self, page, carpeta_destino, nombre_empresa, timeout, ctx):
        # El nombre del método viene del contrato compartido; NAVAR necesita Excel.
        if self.rango is None or _empresa(nombre_empresa) not in self.empresas:
            raise RuntimeError("Falta confirmar empresa o fechas antes de descargar.")
        _click_robusto(page, [
            lambda: page.get_by_role("button", name=re.compile(r"descarg", re.I)).first,
            lambda: page.locator("button[aria-label*='descarg' i]").first,
            lambda: page.get_by_role("button", name=re.compile(r"filtros", re.I)).locator("xpath=following::button[1]"),
        ], "botón de descarga")
        page.wait_for_timeout(1200)
        captura(page, "menu_descarga")
        with page.expect_download(timeout=timeout) as info:
            page.get_by_text(re.compile(r"^\.?xlsx$", re.I)).click(timeout=timeout)
        descarga = info.value
        if Path(descarga.suggested_filename).suffix.lower() != ".xlsx":
            raise RuntimeError("El banco no entregó XLSX; no se renombra un CSV como Excel.")
        temporal = os.path.join(ctx.descargas_dir, "galicia_por_validar.xlsx")
        descarga.save_as(temporal)
        self.validar_excel(temporal)
        nombre = "%s %s.xlsx" % (self.prefijo, datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S_%f"))
        destino = os.path.join(carpeta_destino, nombre)
        # El vigilante sólo ve el Excel cuando terminó de copiarse y fue validado.
        fd, parcial = tempfile.mkstemp(suffix=".part", dir=carpeta_destino)
        os.close(fd)
        try:
            shutil.copyfile(temporal, parcial)
            os.replace(parcial, destino)
        finally:
            if os.path.exists(parcial):
                os.remove(parcial)
        return destino

    def extraer_cuenta_id(self, archivo_csv):
        return self.cuenta
